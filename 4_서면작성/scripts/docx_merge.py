# -*- coding: utf-8 -*-
"""
docx_merge.py — 서면초안.md 를 서면_프레임.docx 의 "다음" 이후에 병합하여 서면 docx 로 저장

사용법
  python scripts/docx_merge.py <프레임.docx> <초안.md> <출력.docx>
      출력 이름은 4_서면작성/CLAUDE.md 3-1절을 따른다. 예: "<출력폴더>/(GGM) 답변서_초안.docx"
  python scripts/docx_merge.py <출력 폴더>
      호환용. 폴더 안의 서면_프레임.docx + 서면초안.md → 최종서면.docx(3-1절 이름이 아니다)

처리 순서
  0. 초안을 먼저 점검하고, 맞지 않으면 아무 파일도 만들거나 바꾸지 않고 멈춘다.
     - md 는 BOM 이 있어도 읽는다(utf-8-sig). 본문은 첫 줄이 아닌 '---'(대시 3개 이상) 줄 앞까지다
       (style_check·citation_check 와 같은 기준). 그 위에 하단 목록(① [확인 필요 사항] 등)이나
       ※ 고지문이 있으면 멈춘다.
     - 초안에 '입 증 방 법' 줄이 없거나 프레임에 입증방법 제목('별지제목')이 없으면 멈춘다.
     - 한 줄 안에서 '[[각주:' 와 ']]' 짝이 맞지 않으면 멈춘다(']' 하나로 닫았거나 줄바꿈으로 끊긴 표시).
     - '|' 로 시작하는 md 표 줄이 있으면 멈춘다(docx 표로 바꾸지 않는다).
  1. 프레임의 '다음' 문단과 '귀중' 앞 빈 문단 사이를 비우고, md 본문을 채운다.
     - "N. 제목" 줄이 프레임 대제목('1.번호매기기')과 같으면(공백 무시) 프레임 요소를 그대로 재사용.
       프레임 목차에 없으면 새 대제목 문단에 '1.번호매기기' 스타일을 이름으로 입히고 글자 번호는 지운다
       (자동 번호). 스타일이 없으면 Normal 에 볼드로 넣는다(공통/문서변환.md 가. 서식).
       짝을 못 찾은 초안 대제목과 쓰이지 않고 지워진 프레임 목차는 경고로 알리고 종료 코드 1을 낸다
       (프레임 목차의 항목명을 그대로 따른다: 4_서면작성/CLAUDE.md 5절).
       번호가 직전 대제목 다음 번호가 아니거나 '…지급하라.'처럼 해라체로 끝나는 줄은 본문으로 두고 알린다.
     - "가. …" 줄 → 소제목(프레임의 같은 줄을 재사용, 없으면 복제) + 볼드
     - "[갑 제N호증 …]" 한 줄 → 증거 캡션('표의첫행제목', 가운데, keepNext·keepLines)
     - 들여쓴 "을나 제N호증  이름" 줄 → 입증방법 항목('(가)번호매기기_내용')
     - 그 밖의 줄 → 본문 문단(원문자는 본문 한글 글꼴, 가운뎃점 낱말은 noProof)
     - 요소 사이에는 빈 문단 하나
     - 발췌 그림: 캡션 바로 뒤(빈 줄 없이), 폭 15cm, keepNext·keepLines(문서변환.md 나. 증거 이미지).
       캡션 다음 줄의 ![](파일) 지정이 우선하고, 없으면 라운드N/발췌/ 에서 호증키 이름으로 찾는다.
       가지번호는 '의'(을나 제2호증의 2 → 을나2의2.png), 같은 캡션의 두 번째 이후 발췌는 '-2'·'-3'
       (을나2-2.png, 을나2의2-2.png). 순번은 ![]() 로 지정한 캡션도 센다.
       가지번호를 '-2' 로 적은 옛 이름(을나2-2.png)은 같은 초안에 원 호증의 두 번째 캡션이 없을 때만 읽고,
       있으면 어느 쪽 발췌인지 가릴 수 없어 멈춘다. 같은 그림 파일이 두 캡션에 쓰이게 되어도 멈춘다.
  2. docx_restyle  : 본문 문단에 제목 계층별 '내용' 스타일
  3. docx_footnotes: [[각주: …]] 표시를 실제 각주로
  4. docx_normalize: 원문자 글꼴 통일
  5. 세 점검(--check)을 다시 돌려 출력하고, 남은 것이 있으면 종료 코드 1을 낸다.
  1~5 는 출력 폴더의 임시 파일에서 하고, 모두 끝난 뒤에만 출력 파일 자리로 옮긴다. 중간에 실패하면
  출력 파일은 그대로 두고 임시 파일을 지운다.
  출력 파일이 이미 있으면 <출력>.bak_merge_<연월일_시분초>.docx 로 보관한 뒤 바꾼다. 같은 이름의
  보관본이 있으면 _2, _3 … 을 붙여 앞선 보관본(사람이 고친 원본일 수 있다)을 덮어쓰지 않는다.
"""
import re, copy, sys, os, shutil, datetime, uuid
from collections import Counter
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_restyle, docx_footnotes, docx_normalize  # noqa: E402

CAPTION_RE = re.compile(r"^\[(?:(?:갑|을가|을나|을|사|노|증)\s*제\d+호증[^\]]*|[^\]]*(?:준비서면|참고서면|답변서|이유서|판결)[^\]]*)\]$")
KEY_RE = re.compile(r"^\[(갑|을가|을나|을|사|노|증)\s*제(\d+)호증(?:의\s*(\d+))?")
IMG_RE = re.compile(r"^!\[[^\]]*\]\(([^)]+)\)$")
EV_RE = re.compile(r'^(갑|을가|을나|을|사|노|증) 제\d+호증')
SUB_RE = re.compile(r'^[가나다라마바사아자차카타파하]\. ')
# 본문과 하단 목록의 경계. style_check.RULE_LINE·citation_check.body_text 와 같은 기준이며 첫 줄은 보지 않는다.
RULE_LINE = re.compile(r'^\s*-{3,}\s*$')
# 하단 세 목록 머리와 AI 고지문(SKILL.md '작업 방식'·'제한 및 주의사항'). 본문 범위에 있으면 멈춘다.
BOTTOM_RE = re.compile(r'^\s*(?:[①②③]\s*(?:⚠\ufe0f?|📎)?\s*\[(?:확인 필요|취약 논리|추가 확보|추천 증거)'
                       r'|※\s*본 서면 초안은 AI)')
TABLE_RE = re.compile(r'^\s*\|')


# 발췌 그림 삽입 규칙(2026. 9. 8. 박준선 사건 변호사 지시)
#  - 테두리는 이미지에 그리지 않고 Word '그림 테두리'(a:ln)로 준다. 나중에 자르기를 해도 테두리가 남는다.
#  - 폭 15cm 로 넣었을 때 세로가 MAX_PIC_CM 를 넘으면 캡션과 다른 면으로 밀리므로,
#    같은 그림을 두 번 넣고 Word 자르기(a:srcRect)로 위/아래를 나눠 자연스럽게 이어 보이게 한다.
PIC_W_CM = 15.0
MAX_PIC_CM = 14.0
PIC_OVERLAP = 0.02  # 분할 경계에서 글자가 잘려 보이지 않도록 2% 겹친다


def _pic_of(run):
    for el in run._r.iter():
        if el.tag == qn('pic:pic'):
            return el
    return None


def _style_picture(run, crop_t=0.0, crop_b=0.0, line_pt=0.75):
    """그림 테두리와 자르기(srcRect)를 도형 서식으로 넣는다."""
    from docx.oxml import OxmlElement
    pic = _pic_of(run)
    if pic is None:
        return
    if crop_t or crop_b:
        blip_fill = pic.find(qn('pic:blipFill'))
        if blip_fill is not None:
            src = OxmlElement('a:srcRect')
            if crop_t:
                src.set('t', str(int(round(crop_t * 100000))))
            if crop_b:
                src.set('b', str(int(round(crop_b * 100000))))
            blip = blip_fill.find(qn('a:blip'))
            (blip.addnext(src) if blip is not None else blip_fill.insert(0, src))
    sp_pr = pic.find(qn('pic:spPr'))
    if sp_pr is None:
        return
    for old in sp_pr.findall(qn('a:ln')):
        sp_pr.remove(old)
    ln = OxmlElement('a:ln')
    ln.set('w', str(int(round(line_pt * 12700))))
    fill = OxmlElement('a:solidFill')
    clr = OxmlElement('a:srgbClr')
    clr.set('val', '000000')
    fill.append(clr)
    ln.append(fill)
    sp_pr.append(ln)


def _insert_picture(anchor, blank_tpl, parent, styles, img):
    """캡션 뒤에 그림 문단을 넣는다. 너무 크면 두 문단으로 나눠 각각 자른다."""
    from docx.shared import Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from PIL import Image
    with Image.open(img) as im:
        w, h = im.size
    full_cm = PIC_W_CM * h / w
    parts = [(0.0, 0.0)] if full_cm <= MAX_PIC_CM else [
        (0.0, 0.5 - PIC_OVERLAP / 2), (0.5 - PIC_OVERLAP / 2, 0.0)]
    made = []
    for crop_t, crop_b in parts:
        el = copy.deepcopy(blank_tpl)
        anchor.addprevious(el)
        par = Paragraph(el, parent)
        par.style = styles['Normal']
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        # 캡션·그림 붙임(공통/문서변환.md 나., 문체-용례.md: 두 장으로 나눈 경우 두 번째 그림 문단에도)
        par.paragraph_format.keep_with_next = True
        par.paragraph_format.keep_together = True
        run = par.add_run()
        run.add_picture(img, width=Cm(PIC_W_CM),
                        height=Cm(full_cm * (1 - crop_t - crop_b)))
        _style_picture(run, crop_t, crop_b)
        made.append(el)
    return made, full_cm, len(parts)


def _key(text):
    """대제목 비교용: 공백을 모두 뺀 글자."""
    return re.sub(r'\s+', '', text)


def _excerpt_names(m, n):
    """호증키 캡션의 n번째 발췌 이미지 이름 (새 이름, 옛 이름 또는 None).
    가지번호는 '의', 순번은 '-n' 으로 적어 '을나 제2호증의 2'(을나2의2.png)와
    '을나 제2호증'의 두 번째 발췌(을나2-2.png)가 겹치지 않게 한다. 옛 이름은 가지번호를 '-k' 로 적은 것이다."""
    stem, branch = m.group(1) + m.group(2), m.group(3)
    tail = '' if n == 1 else '-%d' % n
    if not branch:
        return stem + tail + '.png', None
    return stem + '의' + branch + tail + '.png', stem + '-' + branch + tail + '.png'


def _backup_name(out):
    """앞선 보관본을 덮어쓰지 않는 보관 이름(초 단위, 겹치면 _2·_3 …)."""
    stem = os.path.splitext(out)[0] + '.bak_merge_' + datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
    bak, k = stem + '.docx', 2
    while os.path.exists(bak):
        bak, k = '%s_%d.docx' % (stem, k), k + 1
    return bak


def read_draft(md):
    """초안 md 의 제출용 본문 줄. 병합 전에 형식을 점검하고, 맞지 않으면 파일을 건드리지 않고 멈춘다."""
    lines = open(md, encoding='utf-8-sig').read().split('\n')
    body = []
    for i, ln in enumerate(lines):
        if i and RULE_LINE.match(ln):
            break
        body.append((i + 1, ln.rstrip()))
    problems = []
    for no, ln in body:
        s = ln.strip()
        if BOTTOM_RE.match(s):
            problems.append(f"{no}행: 하단 목록·고지문이 본문 범위에 있습니다. 입증방법 뒤, 하단 목록 앞에 '---' 한 줄을 둡니다: {s[:40]}")
        elif TABLE_RE.match(s):
            problems.append(f"{no}행: md 표는 docx 표로 바꾸지 않습니다. 문장으로 쓰거나 병합 뒤 Word 에서 표를 넣습니다: {s[:40]}")
        elif s.count(docx_footnotes.OPEN) != len(docx_footnotes.MARK.findall(s)):
            problems.append(f"{no}행: '[[각주:' 와 ']]' 짝이 맞지 않습니다(']' 하나로 닫았거나 줄바꿈으로 끊김): {s[:40]}")
    if not any(_key(ln) == '입증방법' for _, ln in body):
        problems.append("초안 본문('---' 위)에 '입 증 방 법' 줄이 없습니다. 새로 내는 호증이 없어도 제목 줄은 둡니다"
                        "(SKILL.md 14: 서면 말미는 '입 증 방 법'으로 끝냅니다).")
    if problems:
        raise SystemExit("병합하지 않았습니다(파일을 만들거나 바꾸지 않음).\n  " + "\n  ".join(problems))
    return [ln for _, ln in body]


def _final_check(path):
    """병합 뒤 세 점검(4_서면작성/CLAUDE.md 6절)을 돌려 출력하고, 남은 항목이 있는 점검 수를 돌려준다."""
    print("점검:")
    left = docx_restyle.main(path, check=True)
    left += docx_footnotes.convert(path, check_only=True)
    font, bad = docx_normalize.normalize(path, check_only=True)
    print(f"원문자 글꼴 비정합 {bad}건")
    return left + (1 if bad else 0)


def merge(frame, md, out):
    d = Document(frame)
    paras = d.paragraphs

    def find(pred):
        for i, p in enumerate(paras):
            if pred(p):
                return i
        raise SystemExit("프레임에서 기준 문단을 찾지 못했습니다")
    i_daum = find(lambda p: p.text.replace('\t', '').replace(' ', '') == '다음')
    i_gwi = max(i for i, p in enumerate(paras) if p.text.strip().endswith('귀중'))
    i_end = i_gwi - 1
    p_end = paras[i_end]
    if p_end.text.strip():
        raise SystemExit("'귀중' 바로 앞에 빈 문단이 없습니다")
    # 본문 run 서식 표본: '다음' 앞의 첫 본문 문단
    body_font_rpr = None
    for p in paras[:i_daum]:
        if p.runs and p.style.name == 'Normal' and p.text.strip():
            body_font_rpr = p.runs[0]._r.find(qn('w:rPr'))
    blank_tpl = copy.deepcopy(p_end._p)

    h1 = []; sub = {}; ev_title = None; ev_tpl = None; sub_tpl = None
    for p in paras[i_daum + 1:i_end]:
        t = p.text.strip()
        if p.style.name == '1.번호매기기' and t:
            h1.append((t, p._p))
        elif p.style.name == '별지제목' and t.replace(' ', '') == '입증방법':
            ev_title = p._p
        elif p.style.name == '(가)번호매기기_내용' and '호증' in t:
            ev_tpl = p._p
        elif SUB_RE.match(t):
            sub[t] = p._p
            if sub_tpl is None:
                sub_tpl = p._p
    for p in paras[i_daum + 1:i_end]:
        p._p.getparent().remove(p._p)

    st = d.styles['Normal'].element.find(qn('w:rPr'))
    rfs = st.find(qn('w:rFonts')) if st is not None else None
    EA = (rfs.get(qn('w:eastAsia')) if rfs is not None else None) or '바탕체'

    def set_text(elem, text, bold=None, plain=False):
        """plain=True 면 본문 run 서식 표본을 쓰지 않는다(제목 스타일의 글꼴·볼드가 그대로 보이게)."""
        for r in elem.findall(qn('w:r')):
            elem.remove(r)
        p = Paragraph(elem, p_end._parent)
        for seg in re.split(r"([①-⑳]+|\S*·\S*)", text):
            if not seg:
                continue
            run = p.add_run(seg)
            rpr = copy.deepcopy(body_font_rpr) if body_font_rpr is not None and not plain else None
            if rpr is not None:
                for b in rpr.findall(qn('w:b')) + rpr.findall(qn('w:bCs')):
                    rpr.remove(b)
            if re.fullmatch(r"[①-⑳]+", seg):
                if rpr is None:
                    rpr = blank_tpl.makeelement(qn('w:rPr'), {})
                for old in rpr.findall(qn('w:rFonts')):
                    rpr.remove(old)
                f = rpr.makeelement(qn('w:rFonts'), {})
                for k in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
                    f.set(qn('w:' + k), EA)
                f.set(qn('w:hint'), 'eastAsia'); rpr.insert(0, f)
            elif '·' in seg:
                if rpr is None:
                    rpr = blank_tpl.makeelement(qn('w:rPr'), {})
                rpr.append(rpr.makeelement(qn('w:noProof'), {}))
            if rpr is not None:
                run._r.insert(0, rpr)
            if bold is not None:
                run.bold = bold
        return elem

    def new_body(text):
        return set_text(copy.deepcopy(blank_tpl), text, bold=False)

    cap_style = next((s.name for s in d.styles if s.name in ('표의첫행제목', 'Caption')), None)

    def new_caption(text):
        e = copy.deepcopy(blank_tpl); set_text(e, text)
        p = Paragraph(e, p_end._parent)
        if cap_style:
            p.style = d.styles[cap_style]
        p.alignment = 1
        # 캡션만 앞 면 끝에 남지 않게 다음 문단(그림)과 붙인다(공통/문서변환.md 나. 증거 이미지)
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.keep_together = True
        return e

    h1_style = docx_restyle.S_H1 if any(s.name == docx_restyle.S_H1 for s in d.styles) else None

    def new_h1(title, line):
        """프레임 목차에 없는 대제목. 스타일이 있으면 이름으로 입히고 글자 번호는 지운다(자동 번호),
        없으면 Normal 에 볼드로 글자 번호째 넣는다(공통/문서변환.md 가. 서식)."""
        e = copy.deepcopy(blank_tpl)
        ppr = e.find(qn('w:pPr'))
        if ppr is not None:
            for tag in ('w:numPr', 'w:ind', 'w:jc'):
                for old in ppr.findall(qn(tag)):
                    ppr.remove(old)
        if h1_style:
            set_text(e, title, plain=True)
            Paragraph(e, p_end._parent).style = d.styles[h1_style]
        else:
            set_text(e, line, bold=True)
        return e

    body_lines = read_draft(md)
    if ev_title is None:
        raise SystemExit("병합하지 않았습니다(파일을 만들거나 바꾸지 않음). 프레임의 '다음'과 '귀중' 사이에 "
                         "'입 증 방 법'(별지제목 스타일) 문단이 없습니다. 프레임에 입증방법 제목을 두고 다시 병합합니다.")

    out_elems = []; used_sub = set(); img_map = {}
    h1_by_key = {}
    for t, e in h1:
        h1_by_key.setdefault(_key(t), e)
    used_h1 = set(); last_no = 0; new_heads = []; kept_body = []
    for ln in body_lines:
        s = ln.strip()
        if not s:
            continue
        mi = IMG_RE.match(s)
        if mi:
            if out_elems:
                img_map[len(out_elems) - 1] = os.path.basename(mi.group(1))
            continue
        m1 = docx_restyle.H1_RE.match(s)
        if m1:
            k = _key(m1.group(2))
            if k in h1_by_key and k not in used_h1:
                used_h1.add(k); last_no = int(m1.group(1))
                out_elems.append(h1_by_key[k]); continue
            cand = docx_restyle.h1_candidate(s)
            if cand and cand[0] == last_no + 1:
                last_no = cand[0]; new_heads.append(s)
                out_elems.append(new_h1(cand[1], s)); continue
            kept_body.append(s)
        if _key(s) == '입증방법' and ev_title is not None:
            out_elems.append(ev_title); continue
        if ln.startswith('    ') and EV_RE.match(s) and ev_tpl is not None:
            parts = re.split(r'\s{2,}', s, maxsplit=1)
            e = copy.deepcopy(ev_tpl)
            set_text(e, parts[0] + '\t\t' + (parts[1] if len(parts) > 1 else ''), bold=False)
            out_elems.append(e); continue
        if SUB_RE.match(s):
            if s in sub and s not in used_sub:
                e = sub[s]; used_sub.add(s); set_text(e, s, bold=True)
            elif sub_tpl is not None:
                e = copy.deepcopy(sub_tpl); set_text(e, s, bold=True)
            else:
                e = copy.deepcopy(blank_tpl); set_text(e, s, bold=True)
            out_elems.append(e); continue
        if CAPTION_RE.match(s):
            out_elems.append(new_caption(s)); continue
        out_elems.append(new_body(s))

    # 대제목이 조용히 본문으로 떨어지거나 프레임 목차가 조용히 사라지지 않게 알린다.
    # 프레임에 목차가 있으면 그 항목명을 그대로 따라야 하므로(4_서면작성/CLAUDE.md 5절) 어긋남은 종료 코드 1로 낸다.
    toc_mismatch = 0
    if h1:
        for s in new_heads:
            print("경고: 프레임 목차에 없는 대제목을 새 대제목으로 넣었습니다(항목명이 프레임과 같은지 확인):", s)
        for t, _ in h1:
            if _key(t) not in used_h1:
                print("경고: 초안에 없어 빠진 프레임 목차:", t)
                toc_mismatch += 1
        toc_mismatch += len(new_heads)
    elif new_heads:
        print(f"프레임에 목차가 없어 초안의 대제목 {len(new_heads)}개를 "
              + (f"'{h1_style}' 스타일로" if h1_style else "볼드 본문으로") + " 넣었습니다")
    for s in kept_body:
        print("경고: 번호로 시작하지만 대제목 순서·형식이 아니어서 본문으로 둔 줄:", s[:60])

    # 발췌 이미지: 캡션마다 넣을 그림을 먼저 정한다. md 의 ![](파일) 지정이 우선하고, 없으면
    # 호증키 이름(가지번호 '의', 같은 캡션 반복 시 -2, -3)으로 라운드N/발췌/ 에서 찾는다.
    # 순번은 ![]() 로 지정한 캡션도 세어, 지정과 자동 찾기가 섞여도 첫 그림을 다시 쓰지 않는다.
    base = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(out))), '발췌')
    caps = []
    for idx, e in enumerate(out_elems):
        t = "".join(x.text or '' for x in e.iter(qn('w:t'))).strip()
        if t.startswith('[') and CAPTION_RE.match(t):
            caps.append((idx, KEY_RE.match(t), t))
    parents = Counter(m.group(1) + m.group(2) for _, m, _ in caps if m and not m.group(3))
    pictures, seen, used = {}, Counter(), {}
    for idx, m, t in caps:
        n = 0
        if m:
            seen[m.groups()] += 1; n = seen[m.groups()]
        if img_map.get(idx):
            img = os.path.join(base, img_map[idx])
        elif m:
            name, old = _excerpt_names(m, n)
            img = os.path.join(base, name)
            if old and not os.path.exists(img) and os.path.exists(os.path.join(base, old)):
                k = int(m.group(3))
                if n == 1 and k >= 2 and parents[m.group(1) + m.group(2)] >= k:
                    raise SystemExit(
                        f"병합하지 않았습니다(파일을 만들거나 바꾸지 않음). 발췌 이미지 {old} 가 "
                        f"'{m.group(1)} 제{m.group(2)}호증의 {k}' 발췌인지 '{m.group(1)} 제{m.group(2)}호증'의 "
                        f"{k}번째 발췌인지 가릴 수 없습니다. 가지번호 발췌라면 파일 이름을 바꾸거나(→ {name}), "
                        "캡션 다음 줄에 ![](파일) 로 지정한 뒤 다시 병합합니다.")
                print(f"참고: {t} 에 옛 이름의 발췌 이미지({old})를 넣었습니다. 가지번호 발췌 이름은 {name} 입니다.")
                img = os.path.join(base, old)
        else:
            continue
        if not os.path.exists(img):
            print("발췌 이미지 없음:", os.path.basename(img)); continue
        key = os.path.normcase(os.path.abspath(img))
        if key in used:
            raise SystemExit(
                f"병합하지 않았습니다(파일을 만들거나 바꾸지 않음). 같은 발췌 이미지 {os.path.basename(img)} 가 "
                f"두 캡션에 들어가게 됩니다: {used[key]} / {t}. 대목이 다르면 다른 이름으로 저장하고 "
                "캡션 다음 줄에 ![](파일) 로 지정합니다(같은 대목은 한 번만 넣는다: 공통/문서변환.md 나.).")
        used[key] = t
        pictures[idx] = img

    # 삽입 기준점: '다음' 문단과 '귀중' 앞 빈 문단 사이에 표(날짜·대리인 마무리 블록)가 있으면
    # 그 표 앞에 본문을 넣는다. 그렇지 않으면 마무리 표가 본문 앞으로 밀려 나온다(2026. 9. 8. 변호사 지적).
    anchor = p_end._p
    node = paras[i_daum]._p.getnext()
    while node is not None and node is not p_end._p:
        if node.tag == qn('w:tbl'):
            anchor = node
            break
        node = node.getnext()
    # '다음' 제목과 본문 사이 빈 문단 하나
    anchor.addprevious(copy.deepcopy(blank_tpl))
    def _is_ev(el):
        return el is not None and el is not ev_title and ev_tpl is not None and             "".join(t.text or '' for t in el.iter(qn('w:t'))).strip().startswith(tuple('갑을사노증'))            and EV_RE.match("".join(t.text or '' for t in el.iter(qn('w:t'))).strip()) is not None             and el.find(qn('w:pPr')) is not None and el.find(qn('w:pPr')).find(qn('w:pStyle')) is not None             and el.find(qn('w:pPr')).find(qn('w:pStyle')).get(qn('w:val')) == (ev_tpl.find(qn('w:pPr')).find(qn('w:pStyle')).get(qn('w:val')) if ev_tpl.find(qn('w:pPr')) is not None and ev_tpl.find(qn('w:pPr')).find(qn('w:pStyle')) is not None else None)
    prev = None
    for idx, e in enumerate(out_elems):
        if e is ev_title:
            # 입증방법은 새 면에서 시작(Ctrl+Enter)
            br_r = e.makeelement(qn('w:r'), {}); br = br_r.makeelement(qn('w:br'), {qn('w:type'): 'page'}); br_r.append(br)
            ppr = e.find(qn('w:pPr'))
            (ppr.addnext(br_r) if ppr is not None else e.insert(0, br_r))
        if prev is not None and not (_is_ev(prev) and _is_ev(e)) and not getattr(e, '_no_blank_before', False):
            anchor.addprevious(copy.deepcopy(blank_tpl))
        anchor.addprevious(e)
        prev = e
        # 발췌 이미지: 캡션 바로 뒤(빈 줄 없이) 삽입
        img = pictures.get(idx)
        if img is None:
            continue
        made, full_cm, n = _insert_picture(anchor, blank_tpl, p_end._parent, d.styles, img)
        prev = made[-1]
        print("발췌 이미지 삽입: %s (%.1fcm%s)" % (
            os.path.basename(img), full_cm, ", 2장 분할" if n > 1 else ""))
    if anchor is not p_end._p:
        anchor.addprevious(copy.deepcopy(blank_tpl))

    # 후처리까지 임시 파일에서 마친 뒤에만 출력 자리로 옮긴다. 도중에 멈추면 기존 출력(사람이 고친 판일 수
    # 있다)을 반쯤 처리된 파일로 바꾸거나 보관본으로 밀어내지 않는다.
    tmp = os.path.join(os.path.dirname(os.path.abspath(out)), '.docx_merge_' + uuid.uuid4().hex + '.docx')
    try:
        d.save(tmp)
        docx_restyle.main(tmp, backup=False)
        docx_footnotes.convert(tmp)
        docx_normalize.normalize(tmp)
        left = _final_check(tmp)
    except BaseException:
        try:
            os.remove(tmp)
        except OSError:
            pass
        print("병합을 마치지 못해 출력 파일은 바꾸지 않았습니다:", out)
        raise
    try:
        if os.path.exists(out):
            bak = _backup_name(out)
            shutil.copy2(out, bak)
            print("기존 파일 보관 →", os.path.basename(bak))
        os.replace(tmp, out)
    except OSError as err:
        try:
            os.remove(tmp)
        except OSError:
            pass
        raise SystemExit(f"출력 파일을 바꾸지 못했습니다(Word 에서 열려 있는지 확인): {out} ({err})")
    print("병합 저장", out, "요소", len(out_elems))
    if left or toc_mismatch:
        print("점검에서 남은 항목이 있습니다(프레임 목차와 어긋남 %d건 포함). 위 내용을 확인합니다." % toc_mismatch)
    return 1 if left or toc_mismatch else 0


if __name__ == '__main__':
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, 'reconfigure'):
            stream.reconfigure(encoding='utf-8')
    a = sys.argv[1:]
    if len(a) == 1 and os.path.isdir(a[0]):
        frame, md, out = (os.path.join(a[0], n) for n in ('서면_프레임.docx', '서면초안.md', '최종서면.docx'))
        print("참고: 폴더 형태는 호환용이라 최종서면.docx 로 저장합니다. 서면 이름 규칙(4_서면작성/CLAUDE.md 3-1절)대로 "
              "저장하려면 <프레임.docx> <초안.md> <출력.docx> 세 인자로 부릅니다.")
    elif len(a) == 3:
        frame, md, out = a
    else:
        print(__doc__); sys.exit(2)
    sys.exit(merge(frame, md, out))
