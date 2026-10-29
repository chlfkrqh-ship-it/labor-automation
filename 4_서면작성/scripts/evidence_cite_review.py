"""서면 본문의 증거 인용을 호증 원문과 맞대는 대조표를 만든다(`.claude/commands/서면.md` 8-1단계 ①).

    python -B 4_서면작성/scripts/evidence_cite_review.py scan "<라운드N>/출력/서면초안.md" [--pdf <서면 pdf>] [--no-ocr]
           [--ocr-dir <폴더> ...] [--ocr-map 을31=<폴더> ...]
    python -B 4_서면작성/scripts/evidence_cite_review.py doc  "<라운드N>/출력/서면초안.md" [--no-pdf]

scan  본문('---' 위)의 증거 인용 괄호마다 그 괄호가 받치는 구절과, 인용한 호증 면에서 그 구절에 가장 가까운 대목을
      찾아 `라운드N/작업/증거인용검토/` 에 review.json·review.txt·crops/ 를 만든다.
      - 호증 파일은 사건 폴더의 우리증거·상대증거·제출/증거·노동위원회_기록에서 파일 이름의 호증번호로 찾는다
        ('갑 제3호증 …', '갑3-1_…', 전자소송 내려받기 '…(을104-2_…)…'). 같은 라운드의 원본을 먼저 쓴다.
      - 글자층이 없거나 깨진 면은 Windows 문자인식(`ocr_win.ps1`)으로 읽는다. 결과는 증거인용검토/ocr/ 에 남겨 다시 쓰고,
        글자가 깨져 나오면 90도·270도로 돌려 한 번 더 읽는다. 이미 만든 결과는 --ocr-dir(호증키 하위 폴더나
        '<호증키>_p01.png' 묶음)·--ocr-map(호증키=폴더)으로 넘긴다.
      - 비교는 날짜(2018. 2. 8. ↔ 2018-02-08)·금액(2,750만 원 ↔ 27,500,000)을 같은 꼴로 바꾸고 기호·띄어쓰기를 지운 뒤
        두 글자 묶음이 겹치는 정도로 한다. 구절의 날짜·금액·따옴표 말이 인용한 면에 없으면 '못찾음'으로 적는다.
doc   review.json 과 판정 파일 `증거인용검토/notes.json` 으로 `증거인용 대조표(확대).docx` 를 만들고, Word 가 있으면 PDF 로도 낸다.
      인용마다 서면 문장, 확인한 원문 문언, 원본 해당 대목을 키운 그림을 싣고, 글자층이 있는 표(계좌내역 등)는
      그 줄을 원본 글자 그대로 옮겨 적는다.
      notes.json 꼴: {"_요약": ["…"], "12": ["일치", "확인한 원문 문언"], "20": ["고침", "…"],
                      "crops": {"30": [["을104-2", 16, ["20180309"]], ["img", "발췌/을104의2.png", "서면에 넣은 발췌 그림"]]}}
      crops 는 자동으로 잡은 대목이 맞지 않을 때만 적는다. [호증키, 면, [찾을 말…]] 은 그 면에서 찾을 말이 든 줄만 자르고,
      ["img", 경로, 설명] 은 그 그림을 그대로 싣는다(경로는 라운드 폴더 기준).

판정은 작성 세션이 원문을 읽고 notes.json 에 적는다. 이 도구는 대목을 찾아 보여 줄 뿐 일치 여부를 판정하지 않는다.
2026. 10. 20. 한전 태양광 준비서면(3) 증거 인용 재검토(99건)에서 만든 도구를 옮긴 것이다.
"""
from __future__ import annotations

import argparse
import importlib.util
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import unicodedata
from pathlib import Path

import pymupdf
from PIL import Image, ImageDraw, ImageFont

pymupdf.TOOLS.mupdf_display_errors(False)
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

ROOT = Path(__file__).resolve().parents[2]
_spec = importlib.util.spec_from_file_location('workflow_system', ROOT / '공통/scripts/system.py')
_system = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(_system)
PREFIX = _system.EVIDENCE_PREFIX

HEADER = re.compile(r'개인정보유출|제출자\s*:|제출일시|다운로드일시|등재일시|원본과 상위 없음')
REF = re.compile(r'(' + PREFIX + r')\s*제\s*(\d+)\s*호증(?:\s*의\s*(\d+))?')
REF_LIST = re.compile(r'(' + PREFIX + r')\s*제\s*(\d+(?:\s*[·ㆍ,]\s*\d+)+)\s*호증')
GROUP = re.compile(r'[\(\[]([^()\[\]]*?(?:' + PREFIX + r')\s*제\s*\d+[^()\[\]]*)[\)\]]')
PAGES = re.compile(r'(\d+(?:\s*~\s*\d+)?)\s*(?:면|쪽)')
EVIDENCE_DIRS = ('우리증거', '상대증거', '증거', '노동위원회_기록')
FILE_COURT = re.compile(r'\((' + PREFIX + r')(\d+)(?:-(\d+))?_')
FILE_LEAD = re.compile(r'^(' + PREFIX + r')\s*(?:제\s*)?(\d+)\s*(?:호증)?\s*(?:(?:의|-)\s*(\d+))?(?=[\s_.()\-]|호증|$)')
FONT_PATH = 'C:/Windows/Fonts/malgun.ttf'
OUT_NAME = '증거인용 대조표(확대)'


# ---------------------------------------------------------------- 글 비교
def canon(s: str) -> str:
    """날짜·금액을 같은 꼴로 바꾸고 글자·숫자만 남긴다."""
    s = unicodedata.normalize('NFC', s or '')
    s = re.sub(r'(\d{4})\s*[.\-/년]\s*(\d{1,2})\s*[.\-/월]\s*(\d{1,2})\s*(?:일|\.)?',
               lambda m: f'{m.group(1)}{int(m.group(2)):02d}{int(m.group(3)):02d}', s)
    s = re.sub(r'(\d{4})\s*[.년]\s*(\d{1,2})\s*(?:월|\.)(?!\d)', lambda m: f'{m.group(1)}{int(m.group(2)):02d}', s)

    def money(m):
        n = [int(g.replace(',', '')) if g else 0 for g in m.groups()]
        return str(n[0] * 100000000 + n[1] * 10000 + n[2])
    s = re.sub(r'(?:(\d[\d,]*)\s*억\s*)?(?:(\d[\d,]*)\s*만\s*)?(\d[\d,]*)?\s*원',
               lambda m: money(m) if any(m.groups()) else m.group(0), s)
    s = re.sub(r'(\d),(?=\d{3})', r'\1', s)
    return re.sub(r'[^0-9A-Za-z가-힣]', '', s)


def keys_of(span: str) -> list:
    """구절에서 원문에 나와야 할 날짜·금액·따옴표 말."""
    ks = []
    for m in re.finditer(r'(\d{4})\.\s*(\d{1,2})\.\s*(\d{1,2})\.', span):
        ks.append(('날짜', m.group(0), f'{m.group(1)}{int(m.group(2)):02d}{int(m.group(3)):02d}'))
    for m in re.finditer(r'(?:\d[\d,]*\s*억\s*)?(?:\d[\d,]*\s*만\s*)?(?:\d[\d,]*)?\s*원', span):
        tok = m.group(0).strip()
        if re.search(r'\d', tok):
            ks.append(('금액', tok, canon(tok)))
    for m in re.finditer(r"['\"‘“]([^'\"’”]{2,40})['\"’”]", span):
        ks.append(('따옴표', m.group(1), canon(m.group(1))))
    return ks


def _bigrams(s: str) -> set:
    return {s[i:i + 2] for i in range(len(s) - 1)}


def best_window(span: str, units: list, key_tokens: list):
    """units 안에서 span 과 가장 많이 겹치는 연속 구간 (점수, 시작, 끝)."""
    cu = [canon(u[0]) for u in units]
    sp = canon(span)
    target = max(80, min(int(len(sp) * 1.6), 700))
    sb = _bigrams(sp)
    best = (-1.0, 0, 0)
    j = total = 0
    for i in range(len(cu)):
        if i:
            total -= len(cu[i - 1])
        while j < len(cu) and total < target:
            total += len(cu[j])
            j += 1
        w = ''.join(cu[i:j])
        if not w:
            continue
        score = len(sb & _bigrams(w)) / (len(sb) or 1) + 0.25 * sum(1 for k in key_tokens if k and k in w)
        if score > best[0]:
            best = (score, i, j)
    return best


def quality(text: str) -> float:
    """문자인식·글자층이 읽을 만한지: 한글·숫자 비율."""
    t = re.sub(r'\s', '', text or '')
    if not t:
        return 0.0
    good = sum(1 for c in t if '가' <= c <= '힣' or c.isdigit())
    return good / len(t)


# ---------------------------------------------------------------- 서면 인용
def parse_citations(text: str) -> list:
    """본문('---' 위)의 증거 인용. 괄호 하나에 호증이 여럿이면 하나씩."""
    lines = text.split('\n')
    end = next((i for i, s in enumerate(lines) if i and re.match(r'^\s*-{3,}\s*$', s)), len(lines))
    rows = []
    for ln, s in enumerate(lines[:end], 1):
        t = re.sub(r'\[\[각주:.*?\]\]', '', s)
        caption = t.strip().startswith('[') and t.strip().endswith(']')
        prev_end = 0
        for g in GROUP.finditer(t):
            gs, ge = g.start(), g.end()
            starts = [m.end() for m in re.finditer(r'(?:다|\))\.\s+', t[:gs])]
            sent_start = starts[-1] if starts else 0
            span = t[max(sent_start, prev_end):gs].strip(' ,')
            m_end = re.search(r'(?:다|\))\.(?:\s|$)', t[ge - 1:])
            sentence = t[sent_start:(ge - 1 + m_end.end()) if m_end else len(t)].strip()
            for part in re.split(r',\s*(?=(?:' + PREFIX + r')\s*제)', g.group(1)):
                pages = [re.sub(r'\s', '', p) for p in PAGES.findall(part)]
                refs = []
                ml = REF_LIST.search(part)
                if ml:
                    refs = [(ml.group(1), int(n), None) for n in re.split(r'\s*[·ㆍ,]\s*', ml.group(2))]
                else:
                    m = REF.search(part)
                    if m:
                        refs = [(m.group(1), int(m.group(2)), int(m.group(3)) if m.group(3) else None)]
                for pre, no, sub in refs:
                    key = f'{pre}{no}' + (f'-{sub}' if sub else '')
                    rows.append({'line': ln, 'key': key, 'raw': part.strip(), 'pages': pages if len(refs) == 1 else [],
                                 'span': span, 'sentence': sentence, 'caption': caption})
            prev_end = ge
    return rows


def headings(text: str) -> list:
    lines = text.split('\n')
    end = next((i for i, s in enumerate(lines) if i and re.match(r'^\s*-{3,}\s*$', s)), len(lines))
    return [(i, s.strip()) for i, s in enumerate(lines[:end], 1)
            if re.match(r'^\d+\.\s+[가-힣]', s) or re.match(r'^[가-하]\.\s', s)]


# ---------------------------------------------------------------- 호증 파일
def file_key(name: str):
    n = unicodedata.normalize('NFC', name)
    for pat in (FILE_COURT, FILE_LEAD):
        m = pat.search(n) if pat is FILE_COURT else pat.match(n)
        if m:
            return f'{m.group(1)}{int(m.group(2))}' + (f'-{int(m.group(3))}' if m.group(3) else '')
    return None


def find_exhibits(case_dir: Path, round_dir: Path) -> dict:
    """호증키 → 파일 목록(같은 라운드의 원본 먼저, 제출 묶음 사본은 뒤)."""
    found = {}
    for dp, dn, fn in os.walk(case_dir):
        parts = Path(dp).relative_to(case_dir).parts
        if '백업' in parts or not any(p in EVIDENCE_DIRS for p in parts):
            continue
        if '원고별 정리' in parts:
            continue
        for f in fn:
            if not f.lower().endswith('.pdf'):
                continue
            k = file_key(f)
            if not k:
                continue
            p = Path(dp) / f
            rank = (0 if round_dir in p.parents else 1, 1 if '제출' in parts else 0, len(str(p)))
            found.setdefault(k, []).append((rank, p))
    return {k: [p for _, p in sorted(v)] for k, v in found.items()}


def lookup(exhibits: dict, key: str):
    if key in exhibits:
        return exhibits[key][0]
    if '-' not in key and f'{key}-1' in exhibits:
        return exhibits[f'{key}-1'][0]
    return None


# ---------------------------------------------------------------- 면 글
def pdf_rows(page) -> list:
    """글자층 낱말을 줄(같은 높이)로 묶는다. [(글, (x0,y0,x1,y1))]"""
    words = sorted(page.get_text('words'), key=lambda w: ((w[1] + w[3]) / 2, w[0]))
    rows = []
    for w in words:
        yc, h = (w[1] + w[3]) / 2, w[3] - w[1]
        if rows and abs(yc - rows[-1][0]) <= max(2.5, h * 0.45):
            rows[-1][1].append(w)
        else:
            rows.append([yc, [w]])
    out = []
    for _, ws in rows:
        ws.sort(key=lambda w: w[0])
        t = ' '.join(w[4] for w in ws)
        if not HEADER.search(t):
            out.append((t, (min(w[0] for w in ws), min(w[1] for w in ws), max(w[2] for w in ws), max(w[3] for w in ws))))
    return out


def _ocr_lines(entry) -> list:
    ls = entry.get('lines') or []
    ls = [ls] if isinstance(ls, dict) else ls
    out = [(l['t'], (l['x'], l['y'], l['x'] + l['w'], l['y'] + l['h'])) for l in ls if not HEADER.search(l.get('t', ''))]
    return sorted(out, key=lambda u: (round(u[1][1] / 12), u[1][0]))


def _read_ocr_json(path: Path) -> list:
    data = json.loads(path.read_text(encoding='utf-8-sig'))
    return [data] if isinstance(data, dict) else data


class Exhibit:
    """호증 하나의 면별 글. 글자층을 쓰고, 없거나 깨진 면은 문자인식."""

    def __init__(self, key, path, cache: Path, ocr_dirs=(), ocr_map=None, use_ocr=True):
        self.key, self.path = key, path
        self.doc = pymupdf.open(path)
        self.cache = cache / 'ocr' / key
        self.ocr_dirs, self.ocr_map, self.use_ocr = list(ocr_dirs), dict(ocr_map or {}), use_ocr
        self.pages = {}

    def __len__(self):
        return len(self.doc)

    def prepare(self, wanted):
        need = []
        for no in wanted:
            if no in self.pages or not 1 <= no <= len(self.doc):
                continue
            rows = pdf_rows(self.doc[no - 1])
            text = ' '.join(t for t, _ in rows)
            if len(re.sub(r'\s', '', text)) >= 80 and quality(text) >= 0.3:
                self.pages[no] = {'src': '글자층', 'units': rows, 'png': None}
            else:
                need.append(no)
        if need:
            got = self._existing_ocr(need)
            rest = [n for n in need if n not in got]
            if rest and self.use_ocr:
                got.update(self._run_ocr(rest))
            for no in need:
                if no in got:
                    self.pages[no] = got[no]
                else:                                 # 문자인식을 못 하면 적게라도 있는 글자층을 쓴다
                    rows = pdf_rows(self.doc[no - 1])
                    self.pages[no] = {'src': '글자층(적음)' if rows else '글자 없음', 'units': rows, 'png': None}

    def close(self):
        self.doc.close()

    def printed_pages(self) -> dict:
        """보고서처럼 면마다 '- 220 -' 꼴 쪽 번호가 인쇄된 문서의 {인쇄 쪽: PDF 면}."""
        if not hasattr(self, '_printed'):
            self._printed = {}
            for i, p in enumerate(self.doc, 1):
                for line in p.get_text().split('\n'):
                    m = re.fullmatch(r'\s*[-–]\s*(\d{1,4})\s*[-–]\s*', line)
                    if m:
                        self._printed.setdefault(int(m.group(1)), i)
        return self._printed

    def _existing_ocr(self, need) -> dict:
        got = {}
        cands = []
        if self.key in self.ocr_map:
            cands.append(Path(self.ocr_map[self.key]))
        cands += [Path(d) / self.key for d in self.ocr_dirs]
        cands.append(self.cache)
        for c in cands:
            j = c / 'pages.json' if (c / 'pages.json').is_file() else c / 'ocr.json'
            if not j.is_file():
                continue
            if j.name == 'pages.json':
                for no, e in json.loads(j.read_text(encoding='utf-8')).items():
                    if int(no) in need and int(no) not in got:
                        got[int(no)] = {'src': '문자인식', 'units': [tuple((u[0], tuple(u[1]))) for u in e['units']],
                                        'png': str(c / e['png'])}
                continue
            for k2, e in enumerate(_read_ocr_json(j), 1):
                m = re.findall(r'\d+', e.get('file', ''))
                no = int(m[-1]) if m else k2
                if no in need and no not in got:
                    got[no] = {'src': '문자인식', 'units': _ocr_lines(e), 'png': str(c / e['file'])}
        for d in self.ocr_dirs:                       # '<호증키>_p01.png' 묶음
            j = Path(d) / 'ocr.json'
            if not j.is_file():
                continue
            for e in _read_ocr_json(j):
                m = re.match(re.escape(self.key) + r'_p(\d+)\.png$', e.get('file', ''))
                if m and int(m.group(1)) in need and int(m.group(1)) not in got:
                    got[int(m.group(1))] = {'src': '문자인식', 'units': _ocr_lines(e), 'png': str(Path(d) / e['file'])}
        return got

    def _render(self, no, folder: Path, rot=0) -> Path:
        p = self.doc[no - 1]
        z = 2300 / max(p.rect.width, p.rect.height)
        pix = p.get_pixmap(matrix=pymupdf.Matrix(z, z).prerotate(rot))
        out = folder / (f'p{no:03d}.png' if rot == 0 else f'p{no:03d}_r{rot}.png')
        pix.save(str(out))
        return out

    def _run_ocr(self, need) -> dict:
        self.cache.mkdir(parents=True, exist_ok=True)
        got = {}
        work = Path(tempfile.mkdtemp(prefix='eocr_'))
        try:
            for no in need:
                self._render(no, work)
            res = run_windows_ocr(work)
            if res is None:
                return {}
            by = {e['file']: e for e in res}
            low = []
            for no in need:
                e = by.get(f'p{no:03d}.png')
                if not e:
                    continue
                units = _ocr_lines(e)
                q = quality(' '.join(t for t, _ in units))
                got[no] = {'src': '문자인식', 'units': units, 'png': f'p{no:03d}.png', 'q': q}
                if q < 0.55:
                    low.append(no)
            if low:
                rot_dir = work / 'rot'
                rot_dir.mkdir()
                for no in low:
                    for r in (90, 270):
                        self._render(no, rot_dir, r)
                res2 = run_windows_ocr(rot_dir) or []
                by2 = {e['file']: e for e in res2}
                for no in low:
                    for r in (90, 270):
                        e = by2.get(f'p{no:03d}_r{r}.png')
                        if not e:
                            continue
                        units = _ocr_lines(e)
                        q = quality(' '.join(t for t, _ in units))
                        if q > got[no]['q'] + 0.1:
                            shutil.copy2(rot_dir / e['file'], work / e['file'])
                            got[no] = {'src': f'문자인식({r}도 돌림)', 'units': units, 'png': e['file'], 'q': q}
            for no, e in got.items():
                shutil.copy2(work / e['png'], self.cache / e['png'])
            old = json.loads((self.cache / 'pages.json').read_text(encoding='utf-8')) if (self.cache / 'pages.json').is_file() else {}
            old.update({str(no): {'png': e['png'], 'units': e['units'], 'src': e['src']} for no, e in got.items()})
            (self.cache / 'pages.json').write_text(json.dumps(old, ensure_ascii=False), encoding='utf-8')
            for no in got:
                got[no]['png'] = str(self.cache / got[no]['png'])
            return got
        finally:
            shutil.rmtree(work, ignore_errors=True)


def run_windows_ocr(folder: Path):
    ps1 = Path(__file__).with_name('ocr_win.ps1')
    try:
        r = subprocess.run(['powershell', '-NoProfile', '-ExecutionPolicy', 'Bypass', '-File', str(ps1),
                            '-Dir', str(folder), '-Filter', '*.png', '-Out', 'ocr.json'],
                           capture_output=True, timeout=3600)
    except (OSError, subprocess.TimeoutExpired):
        return None
    j = folder / 'ocr.json'
    if r.returncode != 0 or not j.is_file():
        return None
    return _read_ocr_json(j)


def count_exhibits(report) -> str:
    keys = {r['key'] for r in report}
    main = {k.split('-')[0] for k in keys}
    return f'{len(main)}개' + (f', 가지번호까지 {len(keys)}개' if len(keys) != len(main) else '')


def expand_pages(ps, n) -> list:
    out = []
    for p in ps:
        if '~' in p:
            a, b = map(int, p.split('~'))
            out += list(range(a, b + 1))
        else:
            out.append(int(p))
    return out or list(range(1, n + 1))


# ---------------------------------------------------------------- 그림
def band_image(page_entry, doc, no, y0, y1, hl) -> Image.Image:
    if page_entry.get('png'):
        with Image.open(page_entry['png']) as f:
            im0 = f.convert('RGB')
        W, H = im0.size
        top, bottom = max(0, int(y0)), min(H, int(y1))
        im = im0.crop((0, top, W, bottom))
        sc, oy = 1.0, top
    else:
        p = doc[no - 1]
        clip = pymupdf.Rect(0, max(0, y0), p.rect.width, min(p.rect.height, y1))
        z = 2.6
        pix = p.get_pixmap(matrix=pymupdf.Matrix(z, z), clip=clip)
        im = Image.frombytes('RGB', (pix.width, pix.height), pix.samples)
        sc, oy = z, clip.y0
    if hl:
        ov = Image.new('RGBA', im.size, (0, 0, 0, 0))
        dr = ImageDraw.Draw(ov)
        for r in hl:
            dr.rectangle([r[0] * sc - 3, (r[1] - oy) * sc - 3, r[2] * sc + 3, (r[3] - oy) * sc + 3],
                         fill=(255, 230, 0, 80), outline=(230, 140, 0, 255), width=2)
        im = Image.alpha_composite(im.convert('RGBA'), ov).convert('RGB')
    return im


def stack(pieces, out_png: Path, width=1700):
    try:
        font = ImageFont.truetype(FONT_PATH, 26)
    except OSError:
        font = ImageFont.load_default()
    rows, last = [], None
    for lab, im in pieces:
        if im.width != width:
            im = im.resize((width, max(1, int(im.height * width / im.width))), Image.LANCZOS)
        if lab and lab != last:
            l = Image.new('RGB', (width, 36), (238, 242, 248))
            ImageDraw.Draw(l).text((8, 3), lab, fill=(40, 70, 140), font=font)
            rows.append(l)
            last = lab
        rows += [im, Image.new('RGB', (width, 6), 'white')]
    sheet = Image.new('RGB', (width, sum(r.height for r in rows)), 'white')
    y = 0
    for r in rows:
        sheet.paste(r, (0, y))
        y += r.height
    sheet.save(out_png, quality=85)
    return out_png


def rows_matching(units, terms) -> list:
    """찾을 말이 든 줄(두 줄에 걸친 말도)."""
    tc = [canon(t) for t in terms if canon(t)]
    cl = [canon(t) for t, _ in units]
    hit = []
    for i, (t, b) in enumerate(units):
        if any(x in cl[i] for x in tc):
            hit.append(i)
        elif i + 1 < len(units) and any(x in cl[i] + cl[i + 1] and x not in cl[i + 1] for x in tc):
            hit += [i, i + 1]
    return sorted(set(hit))


def group_bands(units, idx):
    if not idx:
        return []
    boxes = [units[i][1] for i in idx]
    hh = max(b[3] - b[1] for b in boxes) or 10
    bands = []
    for i, b in sorted(zip(idx, boxes), key=lambda x: x[1][1]):
        if bands and b[1] - bands[-1][1] < hh * 2.5:
            bands[-1][1] = max(bands[-1][1], b[3])
            bands[-1][2].append(i)
        else:
            bands.append([b[1], b[3], [i]])
    return [(y0 - hh * 0.6, y1 + hh * 0.6, ids) for y0, y1, ids in bands]


# ---------------------------------------------------------------- scan
def round_of(md: Path) -> Path:
    return md.resolve().parent.parent


def brief_pdf_of(md: Path):
    out = md.resolve().parent
    pdfs = sorted((p for p in out.glob('*.pdf') if '초안' in p.name or '수정안' in p.name or re.search(r'_Rev_\d+', p.name)),
                  key=lambda p: p.stat().st_mtime, reverse=True)
    return pdfs[0] if pdfs else None


def brief_page_numbers(brief: Path) -> list:
    """서면 PDF 면마다 인쇄된 쪽 번호(머리·바닥 세 줄 안에서 숫자만 있는 줄, '- 3 -' 꼴 포함). 없는 면은 None.
    표지 있는 서면은 표지가 None 이고 본문이 1부터, 표지 없는 소장 청구원인은 첫 면부터 1이다. 2026. 10. 21. 김경민
    청구원인에서 '표지가 1면'만 가정하여 인쇄 면이 한 면씩 적게(첫 면은 0면으로) 나왔다."""
    nums = []
    for p in pymupdf.open(brief):
        lines = [s.strip() for s in p.get_text().split('\n') if s.strip()]
        num = None
        for s in lines[:3] + lines[-3:]:
            m = re.fullmatch(r'[-−–]?\s*(\d{1,4})\s*[-−–]?', s)
            if m:
                num = int(m.group(1))
                break
        nums.append(num)
    return nums


def cmd_scan(args) -> int:
    md = Path(args.md)
    text = md.read_text(encoding='utf-8-sig')
    rnd = round_of(md)
    case_dir = rnd.parent
    out = rnd / '작업' / '증거인용검토'
    (out / 'crops').mkdir(parents=True, exist_ok=True)
    cites = parse_citations(text)
    md_lines = text.split('\n')
    exhibits = find_exhibits(case_dir, rnd)
    ocr_map = dict(x.split('=', 1) for x in (args.ocr_map or []))
    brief = Path(args.pdf) if args.pdf else brief_pdf_of(md)
    brief_txt = [canon(p.get_text()) for p in pymupdf.open(brief)] if brief and brief.is_file() else []
    brief_nums = brief_page_numbers(brief) if brief_txt else []
    opened = {}
    report = []
    for n, c in enumerate(cites, 1):
        path = lookup(exhibits, c['key'])
        r = {'n': n, 'line': c['line'], 'key': c['key'], 'raw': c['raw'], 'pages': c['pages'], 'span': c['span'],
             'sentence': c['sentence'], 'caption': c['caption'], 'file': str(path) if path else None,
             'npages': 0, 'bad_pages': [], 'found': [], 'miss': [], 'windows': [], 'crop': None, 'rows': [],
             'brief_print_page': None}
        if brief_txt:
            cs = canon(c['span'] or c['sentence'])
            probes = [cs[k:k + 18] for k in range(0, max(1, len(cs) - 18), 9)] or [cs]
            for pi, t in enumerate(brief_txt):
                if sum(1 for pr in probes if pr and pr in t) >= max(1, len(probes) // 2):
                    # 그 면에 인쇄된 쪽 번호. 번호를 읽지 못하면 표지가 1면인 것으로 보아 인쇄 면 = pdf 면 - 1
                    r['brief_print_page'] = brief_nums[pi] if brief_nums[pi] is not None else pi
                    break
        if c['caption']:                       # 발췌 그림 제목: 바로 아래 그림을 싣는다
            nxt = next((s for s in md_lines[c['line']:c['line'] + 3] if s.strip()), '')
            m = re.match(r'^!\[[^\]]*\]\(([^)]+)\)', nxt.strip())
            if m and (rnd / '발췌' / m.group(1)).is_file():
                with Image.open(rnd / '발췌' / m.group(1)) as f:
                    im = f.convert('RGB')
                r['crop'] = str(stack([('서면에 넣은 발췌 그림', im)], out / 'crops' / f'{n:03d}.jpg'))
        if not path:
            report.append(r)
            continue
        if c['key'] not in opened:
            opened[c['key']] = Exhibit(c['key'], path, out, args.ocr_dir or [], ocr_map, not args.no_ocr)
        ex = opened[c['key']]
        r['npages'] = len(ex)
        pages = expand_pages(c['pages'], len(ex))
        r['bad_pages'] = [p for p in pages if p > len(ex)]
        pages = [p for p in pages if p <= len(ex)]
        ex.prepare(pages)
        ks = keys_of(c['span'])
        ktok = [k[2] for k in ks]
        alltxt = ''.join(canon(u[0]) for p in pages for u in ex.pages[p]['units'])
        for kind, tok, cv in ks:
            (r['found'] if cv and cv in alltxt else r['miss']).append(f'{kind}:{tok}')
        cands = []
        for p in pages:
            units = ex.pages[p]['units']
            if units and c['span']:
                sc, i, j = best_window(c['span'], units, ktok)
                cands.append((sc, p, i, j))
        cands.sort(reverse=True)
        if c['pages'] and c['span'] and (not cands or cands[0][0] < 0.25):
            # 인용한 면 번호가 문서에 인쇄된 쪽 번호일 수 있다(감사보고서 등). 그 쪽이 인쇄된 PDF 면도 본다.
            printed = ex.printed_pages()
            for p in pages:
                q = printed.get(p)
                if q and q not in pages:
                    ex.prepare([q])
                    if ex.pages[q]['units']:
                        sc, i, j = best_window(c['span'], ex.pages[q]['units'], ktok)
                        cands.append((sc, q, i, j))
                        r.setdefault('printed', {})[str(q)] = p
            cands.sort(reverse=True)
        for rank, (sc, p, i, j) in enumerate(cands[:2]):
            if rank == 1 and sc < cands[0][0] * 0.8:
                break
            e = ex.pages[p]
            src = e['src'] + (f"(인쇄 쪽 {r['printed'][str(p)]})" if str(p) in r.get('printed', {}) else '')
            r['windows'].append({'page': p, 'src': src, 'score': round(sc, 2),
                                 'text': ' '.join(u[0] for u in e['units'][i:j])})
            if rank == 0:
                units = e['units'][i:j]
                hl = [u[1] for u in units if any(k and k in canon(u[0]) for k in ktok)]
                y0 = min(u[1][1] for u in units)
                y1 = max(u[1][3] for u in units)
                pad = 14 if e.get('png') else 6
                im = band_image(e, ex.doc, p, y0 - pad, y1 + pad, hl)
                r['crop'] = str(stack([(f'{c["key"]} {p}면', im)], out / 'crops' / f'{n:03d}.jpg'))
                if e['src'].startswith('글자층'):
                    r['rows'] = [f'{c["key"]} {p}면 | {u[0]}' for u in units if any(k and k in canon(u[0]) for k in ktok)][:12]
        report.append(r)
    for ex in opened.values():
        ex.close()
    (out / 'review.json').write_text(json.dumps(report, ensure_ascii=False, indent=1), encoding='utf-8')
    with open(out / 'review.txt', 'w', encoding='utf-8') as f:
        for r in report:
            head = f"#{r['n']} 줄{r['line']} 인쇄{r['brief_print_page']}면 {r['raw']} (전체 {r['npages']}면)"
            if not r['file']:
                head += ' 파일 못 찾음'
            if r['bad_pages']:
                head += f" 쪽없음{r['bad_pages']}"
            if r['caption']:
                head += ' [발췌 그림 제목]'
            f.write(head + '\n')
            f.write(f"  본문: {r['span']}\n  찾음 {r['found']} 못찾음 {r['miss']}\n")
            for w in r['windows']:
                f.write(f"  원문 {w['page']}면[{w['src']} {w['score']}]: {w['text'][:900]}\n")
            f.write('\n')
    nofile = [r['key'] for r in report if not r['file']]
    print(f"인용 {len(report)}건(호증 {count_exhibits(report)}) | 파일 못 찾음 {len(nofile)} {sorted(set(nofile))[:10]} | "
          f"쪽 없음 {sum(1 for r in report if r['bad_pages'])} | 날짜·금액·따옴표 못찾음 {sum(1 for r in report if r['miss'])}")
    print('결과:', out / 'review.txt')
    print('다음: review.txt 를 한 건씩 원문과 맞대어 판정을 notes.json 에 적고 doc 을 돌린다. 문자인식이 깨진 면은 그림(crops/)을 직접 본다.')
    return 1 if nofile else 0


# ---------------------------------------------------------------- doc
def crop_from_spec(spec, n, rnd: Path, out: Path, ocr_dirs, ocr_map, use_ocr, cache: dict):
    pieces, texts = [], []
    exhibits = cache.setdefault('_exhibits', find_exhibits(rnd.parent, rnd))
    for item in spec:
        if item[0] == 'img':
            p = Path(item[1])
            p = p if p.is_absolute() else rnd / p
            if p.is_file():
                with Image.open(p) as f:
                    pieces.append((item[2] if len(item) > 2 else '', f.convert('RGB')))
            else:
                print(f'  [{n}] 그림 없음: {p}')
            continue
        key, page, terms = item
        path = lookup(exhibits, key)
        if not path:
            print(f'  [{n}] 호증 파일 없음: {key}')
            continue
        if key not in cache:
            cache[key] = Exhibit(key, path, out, ocr_dirs, ocr_map, use_ocr)
        ex = cache[key]
        ex.prepare([page])
        e = ex.pages.get(page)
        if not e:
            continue
        idx = rows_matching(e['units'], terms)
        if not idx:
            print(f'  [{n}] 찾을 말 없음: {key} {page}면 {terms}')
            continue
        for y0, y1, ids in group_bands(e['units'], idx):
            pieces.append((f'{key} {page}면', band_image(e, ex.doc, page, y0, y1, [e['units'][i][1] for i in ids])))
        if e['src'].startswith('글자층'):
            texts += [f'{key} {page}면 | {e["units"][i][0]}' for i in idx]
    if not pieces:
        return None, []
    return str(stack(pieces, out / 'crops' / f'{n:03d}_지정.jpg')), texts[:20]


def export_pdf(docx: Path) -> Path | None:
    try:
        import win32com.client as win32
    except ImportError:
        return None
    work = Path(tempfile.mkdtemp(prefix='ecr_pdf_'))
    tmp = work / 'copy.docx'
    shutil.copy2(docx, tmp)
    pdf = docx.with_suffix('.pdf')
    word = win32.DispatchEx('Word.Application')
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        d = word.Documents.Open(str(tmp), ReadOnly=True, AddToRecentFiles=False)
        d.ExportAsFixedFormat(str(work / 'out.pdf'), 17)
        d.Close(False)
    finally:
        if word.Documents.Count == 0:
            word.Quit()
    for _ in range(20):
        if (work / 'out.pdf').is_file():
            break
        time.sleep(0.5)
    if not (work / 'out.pdf').is_file():
        return None
    shutil.copy2(work / 'out.pdf', pdf)
    shutil.rmtree(work, ignore_errors=True)
    return pdf


def cmd_doc(args) -> int:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    md = Path(args.md)
    rnd = round_of(md)
    out = rnd / '작업' / '증거인용검토'
    report = json.loads((out / 'review.json').read_text(encoding='utf-8'))
    notes_path = out / 'notes.json'
    notes = json.loads(notes_path.read_text(encoding='utf-8')) if notes_path.is_file() else {}
    crops = notes.get('crops', {})
    heads = headings(md.read_text(encoding='utf-8-sig'))
    ocr_map = dict(x.split('=', 1) for x in (args.ocr_map or []))
    cache = {}

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.6)
    sec.top_margin = sec.bottom_margin = Cm(1.5)
    st = doc.styles['Normal']
    st.font.name = '맑은 고딕'
    st.element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
    st.font.size = Pt(11)
    st.paragraph_format.space_after = Pt(2)

    def para(t, size=11, bold=False, color=None, after=2, keep=False):
        p = doc.add_paragraph()
        r = p.add_run(t)
        r.font.size, r.bold = Pt(size), bold
        if color:
            r.font.color.rgb = RGBColor(*color)
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.keep_with_next = keep
        return p

    verdicts = [(notes.get(str(r['n'])) or ['확인 전'])[0] for r in report]
    count = {v: verdicts.count(v) for v in dict.fromkeys(verdicts)}
    brief = brief_pdf_of(md)
    para(f'{brief.stem if brief else md.stem} 증거 인용 대조표', 17, True, after=4)
    para(f"인용 {len(report)}건(호증 {count_exhibits(report)}) · 판정 " + ', '.join(f'{k} {v}건' for k, v in count.items()),
         11, color=(80, 80, 80))
    for s in notes.get('_요약', []):
        para(s, 11.5)
    para("읽는 법: [번호] 서면 인쇄 면 · 인용 호증과 면 — 판정. '서면' 줄은 그 괄호가 받치는 본문 구간이고, '확인' 줄은 원본에서 "
         "맞댄 문언입니다. 그림은 원본 해당 면에서 그 대목만 잘라 붙인 것이고(노란 칠은 서면의 날짜·금액·따옴표 말이 나온 줄), "
         "고해상도라 PDF 를 확대하면 선명하게 보입니다. 글자층이 있는 표는 그림 아래에 해당 줄을 원본 글자 그대로 옮겨 적었습니다.",
         10.5, color=(90, 90, 90), after=8)

    last_top = last_sub = None
    for r, verdict in zip(report, verdicts):
        top = sub = ''
        for i, h in heads:
            if i > r['line']:
                break
            if re.match(r'^\d+\.', h):
                top, sub = h, ''
            else:
                sub = h
        if top and top != last_top:
            para(top, 14, True, (20, 50, 120), 4)
            last_top, last_sub = top, None
        if sub and sub != last_sub:
            para(sub, 12.5, True, (40, 70, 140), 3)
            last_sub = sub
        memo = (notes.get(str(r['n'])) or ['', ''])[1] if len(notes.get(str(r['n'])) or []) > 1 else ''
        col = (0, 120, 60) if verdict == '일치' else (200, 90, 0) if verdict == '고침' else (150, 150, 150)
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.keep_with_next = True
        a = p.add_run(f"[{r['n']}] 인쇄 {r['brief_print_page']}면 · {r['raw']}" + (' (발췌 그림 제목)' if r['caption'] else '') + '  ')
        a.bold, a.font.size = True, Pt(11.5)
        b = p.add_run(f'— {verdict}')
        b.bold, b.font.size = True, Pt(11.5)
        b.font.color.rgb = RGBColor(*col)
        if r['span']:
            para('서면: ' + r['span'] + '.', 11, keep=True)
        if memo:
            para('확인: ' + memo, 11, color=(60, 60, 60), keep=True)
        img, rows = r['crop'], r['rows']
        if str(r['n']) in crops:
            img2, rows2 = crop_from_spec(crops[str(r['n'])], r['n'], rnd, out, args.ocr_dir or [], ocr_map, not args.no_ocr, cache)
            if img2:
                img, rows = img2, rows2
        if img and Path(img).is_file():
            with Image.open(img) as im:
                iw, ih = im.size
            w = 17.8
            if w * ih / iw > 22:
                w = 22 * iw / ih
            doc.add_picture(img, width=Cm(w))
        elif not r['file']:
            para('호증 파일을 찾지 못하였습니다.', 10.5, color=(200, 0, 0))
        if rows:
            para('원본 글자 그대로(해당 줄):', 10, True, (40, 70, 140), 1)
            for row in rows:
                para(row, 10, after=0)
            para('', 4, after=2)
    for v in cache.values():
        if isinstance(v, Exhibit):
            v.close()
    docx_path = out / f'{OUT_NAME}.docx'
    doc.save(docx_path)
    print('저장', docx_path)
    if not args.no_pdf:
        pdf = export_pdf(docx_path)
        if pdf:
            print('PDF', pdf, f'{len(pymupdf.open(pdf))}면')
        else:
            print('PDF 를 만들지 못하였습니다(Word 없음). docx 를 넘긴다.')
    undecided = verdicts.count('확인 전')
    if undecided:
        print(f'판정 없는 인용 {undecided}건 — notes.json 에 판정을 적는다.')
    return 0


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sp = ap.add_subparsers(dest='cmd', required=True)
    for name in ('scan', 'doc'):
        s = sp.add_parser(name)
        s.add_argument('md', help='라운드N/출력/서면초안.md')
        s.add_argument('--ocr-dir', action='append', help='이미 만든 문자인식 결과 폴더(호증키 하위 폴더 또는 <호증키>_p01.png 묶음)')
        s.add_argument('--ocr-map', action='append', help='호증키=폴더 (예: 을31=작업/빌드/문답서OCR/구교웅)')
        s.add_argument('--no-ocr', action='store_true', help='Windows 문자인식을 돌리지 않는다')
    sp.choices['scan'].add_argument('--pdf', help='서면 PDF(인쇄 면 표시). 없으면 출력/ 의 _초안 PDF')
    sp.choices['doc'].add_argument('--no-pdf', action='store_true', help='PDF 로 내보내지 않는다')
    args = ap.parse_args(argv)
    return cmd_scan(args) if args.cmd == 'scan' else cmd_doc(args)


if __name__ == '__main__':
    sys.exit(main())
