# -*- coding: utf-8 -*-
"""
docx_normalize.py — 서면 docx 마무리 점검·정리 도구

사용법
  python scripts/docx_normalize.py <파일.docx>            # 원문자 글꼴을 본문 한글 글꼴로 통일하여 같은 파일에 저장
  python scripts/docx_normalize.py <파일.docx> --check    # 고치지 않고 검출 건수만 보고 (0건이면 종료코드 0)
  금지 낱말 검사는 이 스크립트가 하지 않는다: python 4_서면작성/scripts/style_check.py <파일>

원문자(①~⑳)를 담은 run은 프레임 Normal 스타일의 eastAsia 글꼴(통상 바탕체)로
ascii·hAnsi·eastAsia·cs를 모두 지정하고 hint="eastAsia"를 붙인다.
(2026. 9. 7. 이케아 사건 변호사 지적: 원문자가 본문과 다른 글꼴로 찍히는 문제)
run 은 원문자 조각과 나머지 조각으로만 나누고, 탭·줄바꿈 등 글자 아닌 자식은 원래 자리에 한 번만 둔다.
정리 전후 문단 글자(탭·줄바꿈 포함)가 다르면 저장하지 않고 멈춘다.
"""
import sys, re, copy
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

CIRC = re.compile(r"[①-⑳]")
CIRC_RUN = re.compile(r"[①-⑳]+")


def body_font(doc):
    rpr = doc.styles['Normal'].element.find(qn('w:rPr'))
    rf = rpr.find(qn('w:rFonts')) if rpr is not None else None
    return (rf.get(qn('w:eastAsia')) if rf is not None else None) or '바탕체'


def run_ok(r, font):
    rpr = r._r.find(qn('w:rPr'))
    rf = rpr.find(qn('w:rFonts')) if rpr is not None else None
    if rf is None:
        return False
    return all(rf.get(qn('w:' + k)) == font for k in ('ascii', 'hAnsi', 'eastAsia')) and rf.get(qn('w:hint')) == 'eastAsia'


def iter_paragraphs(doc):
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    yield p


def _visible(p_el):
    """문단 글자(탭 '\\t', 줄바꿈 '\\n' 포함). 탭 위치 정의(w:tabs 안의 w:tab)는 빼고 run 안의 것만 본다."""
    out = []
    for el in p_el.iter(qn('w:t'), qn('w:tab'), qn('w:br'), qn('w:cr')):
        if el.getparent().tag != qn('w:r'):
            continue
        if el.tag == qn('w:t'):
            out.append(el.text or '')
        else:
            out.append('\t' if el.tag == qn('w:tab') else '\n')
    return ''.join(out)


def _circled_font(rpr, font):
    for old in rpr.findall(qn('w:rFonts')):
        rpr.remove(old)
    f = OxmlElement('w:rFonts')
    for k in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        f.set(qn('w:' + k), font)
    f.set(qn('w:hint'), 'eastAsia')
    rs = rpr.find(qn('w:rStyle'))          # 스키마 순서: rStyle 다음이 rFonts
    (rs.addnext(f) if rs is not None else rpr.insert(0, f))


def _split_run(r_el, font):
    """run 을 원문자 조각과 나머지 조각으로 나눈다. w:t 는 글자 단위로 쪼개고,
    탭·줄바꿈 등 다른 자식은 원래 순서 그대로 나머지 조각 run 에 한 번만 옮긴다."""
    groups = []                                   # [원문자 조각인가, [자식]]
    for ch in list(r_el):
        if ch.tag == qn('w:rPr'):
            continue
        if ch.tag == qn('w:t'):
            for seg in re.split(r"([①-⑳]+)", ch.text or ''):
                if not seg:
                    continue
                t = OxmlElement('w:t'); t.text = seg; t.set(qn('xml:space'), 'preserve')
                circ = bool(CIRC_RUN.fullmatch(seg))
                if not circ and groups and not groups[-1][0]:
                    groups[-1][1].append(t)
                else:
                    groups.append([circ, [t]])
        elif groups and not groups[-1][0]:
            groups[-1][1].append(ch)
        else:
            groups.append([False, [ch]])
    base = r_el.find(qn('w:rPr'))
    prev = r_el
    for circ, kids in groups:
        new_r = r_el.makeelement(qn('w:r'), dict(r_el.attrib))
        rpr = copy.deepcopy(base) if base is not None else (OxmlElement('w:rPr') if circ else None)
        if circ:
            _circled_font(rpr, font)
        if rpr is not None:
            new_r.append(rpr)
        for k in kids:
            new_r.append(k)
        prev.addnext(new_r); prev = new_r
    r_el.getparent().remove(r_el)


def normalize(path, check_only=False):
    doc = Document(path)
    font = body_font(doc)
    bad = 0
    for p in iter_paragraphs(doc):
        before = None
        for r in list(p.runs):
            if not CIRC.search(r.text):
                continue
            if CIRC.fullmatch(r.text) and run_ok(r, font):
                continue
            bad += 1
            if check_only:
                continue
            if before is None:
                before = _visible(p._p)
            _split_run(r._r, font)
        if before is not None and _visible(p._p) != before:
            raise SystemExit("원문자 정리 중 문단 글자가 달라져 저장하지 않았습니다(원문 보존). 문단 앞부분: "
                             + repr(before[:60]))
    if not check_only and bad:
        doc.save(path)
    return font, bad


if __name__ == '__main__':
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, 'reconfigure'):
            stream.reconfigure(encoding='utf-8')
    if len(sys.argv) > 1 and sys.argv[1] in ('-h', '--help'):
        print(__doc__); sys.exit(0)
    if len(sys.argv) < 2:
        print(__doc__); sys.exit(2)
    path = sys.argv[1]
    if '--words' in sys.argv:
        # 금지 낱말 목록은 style_check.py 한 곳에만 둔다(두 벌이면 한쪽만 고쳐져 이미 갈라졌다).
        print("금지 낱말 검사는 style_check.py 로 합니다: python 4_서면작성/scripts/style_check.py " + path)
        sys.exit(2)
    check = '--check' in sys.argv
    font, bad = normalize(path, check_only=check)
    print(f"본문 한글 글꼴: {font} / 원문자 run 비정합 {bad}건" + ("" if check else " → 정리 완료" if bad else ""))
    sys.exit(1 if (check and bad) else 0)
