"""노동 금액 계산표 엑셀.

    요약 / 통상임금 / 시간외수당 / 연차휴가수당 / 평균임금 / 퇴직금 / 해고기간 임금 /
    지연손해금 / 소멸시효 / 경고 / 근거

신체손해 계산표(engine/excel.py)와 달리 대법원 프로그램 양식이 없으므로, 모듈 결과 행을 그대로 펼친다.
값만 쓰고 수식은 넣지 않는다. '근거' 시트에 계산한 절의 옵션 값(기본값 여부)과 규칙 상태(불명확 등)가 찍혀,
서면·검토의견에서 금액을 단정해도 되는지 여기서 판단한다. 해고기간 임금 시트 끝에는 장래분 월 금액 행
(구분 future, 원금 합계 불산입)을 붙인다. 끝수처리 전 중간값(통상시급·1일 통상임금 등)은 소수까지 보인다.
"""

from __future__ import annotations

import dataclasses
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import calculate as calc_mod

TITLE = Font(name="맑은 고딕", size=14, bold=True)
HEAD = Font(name="맑은 고딕", size=9, bold=True)
BODY = Font(name="맑은 고딕", size=9)
HEAD_FILL = PatternFill("solid", fgColor="EDEDED")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
MONEY = "#,##0"

LABELS = {
    # 공통
    "start": "시작", "end": "끝", "note": "비고", "kind": "구분", "amount": "금액", "days": "일수",
    "period_start": "기간 시작", "period_end": "기간 끝", "due_date": "지급기일", "paid": "기지급",
    "diff": "차액", "claim_amount": "청구액", "label": "항목", "rate": "비율/이율",
    # 통상임금
    "regime": "적용 법리", "items": "항목별 산입액", "monthly_sum": "월 산입액 합", "monthly_ordinary": "월 통상임금",
    "weekly_base_hours": "주 기준시간", "monthly_hours": "월 기준시간", "hourly_raw": "통상시급(끝수 전)",
    "hourly": "통상시급", "daily_hours": "1일 소정근로시간", "daily_ordinary": "1일 통상임금", "agreed_hourly": "약정 시급",
    # 시간외수당
    "key": "임금산정기간", "overtime_hours": "연장", "in_law_hours": "법내 초과", "part_time_hours": "단시간 초과",
    "night_hours": "야간", "holiday_le8_hours": "휴일 8h 이내", "holiday_gt8_hours": "휴일 8h 초과",
    "overtime_night_hours": "연장+야간", "holiday_night_hours": "휴일+야간", "unproven_hours": "미증명 시간",
    "lines": "산정 내역", "contract_lines": "약정 산정 내역", "legal_by_item": "법정 항목별", "contract_by_item": "약정 항목별",
    "legal_amount": "법정 수당", "contract_amount": "약정 수당", "amount_by_item": "항목별 수당", "paid_total": "기지급 합",
    # 연차
    "period_no": "년차", "accrual_date": "발생일", "source": "발생 사유", "attendance_rate": "출근율",
    "accrued_days": "발생일수", "accrued_hours": "발생시간", "unit": "단위", "used": "사용", "extinguished": "촉진 소멸",
    "unused": "미사용", "use_end": "사용기간 끝", "wage_ref_date": "기준임금 시점", "daily_wage": "1일 통상임금",
    "statutory_amount": "법정 수당", "agreed_amount": "약정 수당", "paid_amount": "기지급", "claim_arises": "청구권 발생일",
    "pay_due_date": "지급기일", "settlement": "퇴직 청산", "prescription_date": "시효 완성일",
    "promotion_lawful": "사용촉진 적법", "claim_pending": "청구권 미발생",
    # 평균임금·퇴직금
    "calc": "산정", "counted": "산입", "segment_no": "구간", "calendar_days": "역일수", "excluded_days": "제외일수",
    "counted_days": "계속근로일수", "years": "년", "remaining_days": "잔여일", "ratio": "지급률",
    "average_daily_wage": "1일 평균임금", "amount_raw": "금액(끝수 전)",
    # 해고기간 임금
    "period_days": "산정기간 일수", "covered_days": "해당 일수", "proration": "일할 방식", "wage": "임금",
    "deduction_base": "공제 대상 임금", "small_business": "상시 4명 이하", "allowance": "휴업수당 상당액",
    "limit": "공제 한도", "income": "중간수입(한도 적용)", "income_uncapped": "중간수입(한도 없는 자인분)",
    "deduction": "중간수입 공제", "pay": "세전 지급액", "pay_amount": "청구 금액", "pay_date": "정기지급일",
    # 지연손해금
    "claim": "원금 항목", "category": "분류", "principal": "원금", "fraction": "기간(년)", "basis": "이율 근거",
    "rule": "규칙", "interest": "지연손해금", "expiry": "시효 만료일", "suspect": "시효 완성 의심",
}
# 같은 필드 이름이 절마다 뜻이 다를 때
SECTION_LABELS = {
    "dismissal": {"items": "항목별 임금", "rate": "휴업수당 율", "label": "기간 표시"},
}

# 원 단위로 확정한 금액(끝수가 남아 있으면 소수도 보인다)
MONEY_FIELDS = {"amount", "paid", "diff", "claim_amount", "legal_amount", "contract_amount", "paid_total",
                "statutory_amount", "agreed_amount", "paid_amount", "principal", "interest",
                "wage", "deduction_base", "allowance", "limit", "income", "income_uncapped", "deduction", "pay",
                "pay_amount"}
# 끝수처리 전 중간값 — 엔진은 이 값을 그대로 곱하므로 원 단위로 반올림해 보이지 않는다
FRACTION_FIELDS = {"monthly_sum", "monthly_ordinary", "daily_ordinary", "daily_wage", "amount_raw", "hourly_raw",
                   "hourly", "agreed_hourly", "average_daily_wage"}
FRACTION = "#,##0.00##"


def _cell_value(v):
    if isinstance(v, bool):
        return "예" if v else "아니오"
    if isinstance(v, Decimal):
        return float(v) if v != v.to_integral_value() else int(v)
    if isinstance(v, (date, int, float, str)) or v is None:
        return v
    if isinstance(v, dict):
        return ", ".join(f"{k}: {_short(x)}" for k, x in v.items())
    if isinstance(v, (list, tuple)):
        return "; ".join(_short(x) for x in v)
    if dataclasses.is_dataclass(v):
        return ", ".join(f"{f.name}={_short(getattr(v, f.name))}" for f in dataclasses.fields(v))
    return str(v)


def _short(x):
    if isinstance(x, Decimal):
        return f"{x:,}"
    if isinstance(x, date):
        return f"{x:%Y.%m.%d}"
    if dataclasses.is_dataclass(x):
        return "(" + ", ".join(f"{f.name}={_short(getattr(x, f.name))}" for f in dataclasses.fields(x)) + ")"
    return str(x)


def _put(ws, r, c, v, font=BODY, fill=None, fmt=None, wrap=False):
    cell = ws.cell(row=r, column=c, value=_cell_value(v))
    cell.font = font
    if fill:
        cell.fill = fill
    if isinstance(v, date):
        cell.number_format = "yyyy.mm.dd"
    elif fmt:
        cell.number_format = fmt
    if wrap:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    return cell


def _number_format(name, value):
    if name in FRACTION_FIELDS:
        return FRACTION
    if name in MONEY_FIELDS:
        if isinstance(value, Decimal) and value != value.to_integral_value():
            return FRACTION
        return MONEY
    return None


def _rows_sheet(wb, title, rows, labels=None):
    ws = wb.create_sheet(title)
    _put(ws, 1, 1, title, font=TITLE)
    if not rows:
        _put(ws, 3, 1, "행 없음")
        return ws
    labels = {**LABELS, **(labels or {})}
    names = [f.name for f in dataclasses.fields(rows[0])]
    for j, n in enumerate(names, 1):
        _put(ws, 3, j, labels.get(n, n), font=HEAD, fill=HEAD_FILL)
        ws.column_dimensions[get_column_letter(j)].width = 14 if n not in ("note", "lines", "items") else 40
    for i, row in enumerate(rows, 4):
        for j, n in enumerate(names, 1):
            v = getattr(row, n)
            _put(ws, i, j, v, fmt=_number_format(n, v), wrap=n in ("note", "lines"))
    ws.freeze_panes = "A4"
    return ws


def _section_rows(section, part) -> list:
    rows = list(part.rows)
    future = getattr(part, "future_row", None) if section == "dismissal" else None
    if future is not None:
        note = "; ".join(x for x in (future.note, "장래분 월 금액(DW-19) — 원금 합계·지연손해금 불산입") if x)
        rows.append(dataclasses.replace(future, note=note))
    return rows


def write_labor_workbook(res: calc_mod.LaborResult, path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "요약"
    c = res.case
    _put(ws, 1, 1, "노동 금액 계산표", font=TITLE)
    r = 3
    for k, v in (("사건번호", c.case_no), ("근로자", c.name), ("사용자", c.employer)):
        _put(ws, r, 1, k, font=HEAD)
        _put(ws, r, 2, v)
        r += 1
    r += 1
    for line in calc_mod.summary_lines(res):
        _put(ws, r, 1, line)
        r += 1
    it = res.parts.get("interest")
    if it is not None and it.schedules:
        r += 1
        _put(ws, r, 1, "[청구취지·주문형 지연손해금 문구]", font=HEAD)
        r += 1
        for s in it.schedules:
            _put(ws, r, 1, f"{s.claim.category} / {s.claim.label}")
            _put(ws, r, 3, s.order_text)
            r += 1
        for name, (_, scheds) in it.alternatives.items():
            r += 1
            _put(ws, r, 1, f"[대안: {name}]", font=HEAD)
            r += 1
            for s in scheds:
                _put(ws, r, 1, f"{s.claim.category} / {s.claim.label}")
                _put(ws, r, 3, s.order_text)
                r += 1
    r += 1
    _put(ws, r, 1, "법원 제출 전: 노동 금액은 대법원 계산프로그램이 없으므로 담당자가 산식·입력값을 검산한다. "
                   "'근거' 시트의 불명확 옵션은 대안 금액과 함께 검토한다.", fill=WARN_FILL)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["C"].width = 90

    names = {s: n for s, n, _ in calc_mod.MODULES}
    for section, part in res.parts.items():
        if section == "interest":
            _rows_sheet(wb, "지연손해금", part.rows)
            _rows_sheet(wb, "소멸시효", part.limitation)
        else:
            _rows_sheet(wb, names[section], _section_rows(section, part), SECTION_LABELS.get(section))

    ws = wb.create_sheet("경고")
    _put(ws, 1, 1, "사람이 확인할 점", font=TITLE)
    for i, w in enumerate(res.warnings, 3):
        _put(ws, i, 1, w, wrap=True, fill=WARN_FILL)
    ws.column_dimensions["A"].width = 140

    ws = wb.create_sheet("근거")
    _put(ws, 1, 1, "옵션", font=TITLE)
    _put(ws, 2, 1, "계산한 절의 옵션만 싣는다. 노란 행은 규칙 상태가 불명확인 옵션이다.")
    for j, h in enumerate(("키", "값", "기본값", "규칙", "설명"), 1):
        _put(ws, 3, j, h, font=HEAD, fill=HEAD_FILL)
    # 옵션 키 -> (절, 모듈). 규칙 ID 는 모듈마다 따로 붙였으므로(M1~M6, RS-04 가 겹친다) 상태도 그 모듈의 RULES 에서 찾는다
    owners = {}
    for section, _, mod in calc_mod.MODULES:
        opts = mod.OPTIONS
        for spec in (opts.values() if isinstance(opts, dict) else opts):
            owners[spec.key] = (section, mod)
    i = 4
    for spec, value, is_default in c.option_rows:
        section, mod = owners.get(spec.key, (None, None))
        if section not in res.parts:
            continue
        status = getattr(mod, "RULES", {}).get(spec.rule, ("", ""))[1]
        fill = WARN_FILL if "불명확" in status else None
        for j, v in enumerate((spec.key, str(value), "예" if is_default else "사건 지정",
                               f"{spec.rule} {status}", spec.choices.get(value, spec.description)), 1):
            _put(ws, i, j, v, fill=fill)
        i += 1
    i += 1
    _put(ws, i, 1, "계산 근거 기록", font=TITLE)
    i += 2
    for j, h in enumerate(("모듈", "규칙", "내용", "값", "비고"), 1):
        _put(ws, i, j, h, font=HEAD, fill=HEAD_FILL)
    for mod_name, t in res.trace:
        i += 1
        for j, v in enumerate((mod_name, t.rule, t.label, t.value, t.note), 1):
            _put(ws, i, j, v)
    for col, width in zip("ABCDE", (22, 18, 50, 30, 60)):
        ws.column_dimensions[col].width = width

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
