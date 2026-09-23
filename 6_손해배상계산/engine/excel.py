"""손해배상액 계산표 엑셀 생성.

대법원 프로그램 `엑셀표저장` 의 시트 구성과 항목 배치를 그대로 따른다.

    종합 / 치료비 / 개호비 / 보조구 / 퇴직금(일반)

블록 순서와 라벨 문구는 신 프로그램(SBGCalc) 출력에서 그대로 옮겼다.
설명서 39쪽대로 대법원 엑셀에는 값만 있고 함수가 없으므로, 여기서도 값만 쓴다.

프로그램에 없는 것은 셋이다. 사람이 확인할 점이 있으면 맨 뒤에 '경고' 시트를 두고 종합 시트
3행에 알린다. [일실수입]·[향후 개호비] 제목 옆에 노임·개호 단가 기준을 적는다. 사망 사건
장례비는 합계에 넣지 않으므로 [적극손해] 밖의 [장례비] 블록에 따로 적는다.
"""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .calculate import Result
from .constants import LIVING_COST_DEATH

TITLE_FONT = Font(name="맑은 고딕", size=14, bold=True)
HEAD_FONT = Font(name="맑은 고딕", size=9, bold=True)
BODY_FONT = Font(name="맑은 고딕", size=9)
BLOCK_FONT = Font(name="맑은 고딕", size=9, bold=True)
HEAD_FILL = PatternFill("solid", fgColor="EDEDED")
CAP_FILL = PatternFill("solid", fgColor="FFE0E0")   # 상한에 걸린 칸
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WARN_FONT = Font(name="맑은 고딕", size=9, bold=True, color="C00000")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")

MONEY = '#,##0" 원"'
NUM = "#,##0"
RATE = "0.####"


def _put(ws, row, col, value, *, font=BODY_FONT, fmt=None, box=False, fill=None, align=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font = font
    if fmt:
        c.number_format = fmt
    if box:
        c.border = BOX
    if fill:
        c.fill = fill
    if align:
        c.alignment = Alignment(horizontal=align, vertical="center")
    return c


def _block(ws, row, label):
    _put(ws, row, 2, label, font=BLOCK_FONT)
    return row + 1


def _headers(ws, row, labels, start_col=2):
    for i, text in enumerate(labels):
        if text is None:
            continue
        _put(ws, row, start_col + i, text, font=HEAD_FONT, box=True,
             fill=HEAD_FILL, align="center")
    return row + 1


def _d(v) -> str:
    return f"{v:%Y.%m.%d}" if isinstance(v, date) else (v or "")


def _basic_block(ws, r: Result, row: int, columns: str = "full") -> int:
    """[기초사항] — 종합 시트와 부속 시트가 같은 문구를 쓴다."""
    c = r.case
    age = f"{r.age[0]}세 {r.age[1]}개월 {r.age[2]}일"
    row = _block(ws, row, "[기초사항]")
    pairs = [
        ("사건번호", c.case_no, "사건종류", c.case_type, "사고당시연령", age),
        ("성 명", c.name, "사건유형", c.injury_type, "기대여명",
         f"{c.life_expectancy}년" if c.life_expectancy else ""),
        ("성 별", "남" if c.sex == "M" else "여", "판결/조정", c.judgment_type, "", ""),
        ("생년월일", _d(c.birth), "여명 종료일", _d(c.life_end), "", ""),
        ("사고일자", _d(c.accident), "가동 종료일", _d(c.resolved_work_end()), "", ""),
        ("가동연한(년)", f"{c.work_limit_years}년", "변론 종결일", _d(c.argument_end), "", ""),
        ("입원치료 종료일", _d(c.cure_end), "법정이율(%)", f"{c.legal_rate}%", "", ""),
    ]
    for i, (a, b, cc, d, e, f) in enumerate(pairs):
        rr = row + i
        if a:
            _put(ws, rr, 2, a, font=HEAD_FONT)
            _put(ws, rr, 4, b)
        if cc:
            _put(ws, rr, 6, cc, font=HEAD_FONT)
            _put(ws, rr, 8, d)
        if e:
            _put(ws, rr, 10, e, font=HEAD_FONT)
            _put(ws, rr, 12, f)
    return row + len(pairs) + 1


def _summary_sheet(wb: Workbook, r: Result) -> None:
    c = r.case
    ws = wb.create_sheet("종합")
    _put(ws, 2, 2, "손해배상액 계산표", font=TITLE_FONT)
    if r.warnings:
        _put(ws, 3, 2, f"확인할 점 {len(r.warnings)}건 — '경고' 시트를 먼저 보십시오", font=WARN_FONT)
    death = c.injury_type == "사망"
    row = _basic_block(ws, r, 5)

    # ---------------------------------------------------------- 노동능력 상실률
    row = _block(ws, row, "[노동능력 상실률]")
    _put(ws, row, 2, "[영구장해]", font=BLOCK_FONT)
    _put(ws, row, 4, f"전체 후유장해 100 % 기준 기왕증 기여도 : {r.prior_contribution}%")
    row += 1
    row = _headers(ws, row, ["진료과", None, "개별수치(%)", None, "기왕증(%)",
                             "중복장해(%)", None, "단순중복장해(%)", None, "기왕증 기여도(%)"])
    permanent = [i for i in c.impairments if not i.years]
    for imp in permanent:
        _put(ws, row, 2, imp.dept, box=True)
        _put(ws, row, 4, float(imp.rate), box=True, fmt=RATE)
        _put(ws, row, 6, float(imp.prior) or None, box=True, fmt=RATE)
        row += 1
    if not permanent:
        row += 1      # 영구장해 행이 없으면 빈 행에 적는다. 머리글을 덮어쓰지 않는다.
    _put(ws, row - 1, 7, float(r.combined_rate), box=True, fmt=RATE)
    _put(ws, row - 1, 11, float(r.prior_contribution), box=True, fmt=RATE)
    row += 1

    if any(i.years for i in c.impairments):
        _put(ws, row, 2, "[한시장해]", font=BLOCK_FONT)
        row += 1
        row = _headers(ws, row, ["진료과", None, "개별수치(%)", None, "기왕증(%)",
                                 "시작일", "년수", "종료일", "중복장해(%)"])
        for imp in [i for i in c.impairments if i.years]:
            _put(ws, row, 2, imp.dept, box=True)
            _put(ws, row, 4, float(imp.rate), box=True, fmt=RATE)
            _put(ws, row, 6, float(imp.prior) or None, box=True, fmt=RATE)
            _put(ws, row, 8, float(imp.years), box=True)
            row += 1
        row += 1

    # ---------------------------------------------------------- 일실수입
    if r.wage_basis:
        _put(ws, row, 4, f"노임 기준 : {r.wage_basis}")
    row = _block(ws, row, "[일실수입]")
    head = ["순번", "기간초일", "기간말일", "노임단가", "일수", "월소득",
            "상실률(%)" if c.injury_type != "사망" else "생계비",
            "M1", "호프만1", "M2", "호프만2", "M1-M2", "적용호프만", "기간 일실수입"]
    row = _headers(ws, row, head)
    _put(ws, row - 1, 17, "비고(추가검토)", font=HEAD_FONT, box=True, fill=HEAD_FILL, align="center")
    _put(ws, row - 1, 19, "안분일수", font=HEAD_FONT, box=True, fill=HEAD_FILL, align="center")
    _put(ws, row - 1, 20, "안분액(30일기준)", font=HEAD_FONT, box=True, fill=HEAD_FILL, align="center")

    for r_ in r.income_rows:
        vals = [
            r_.step, _d(r_.start), _d(r_.end), int(r_.wage), r_.days, int(r_.salary),
            float(r_.loss_rate) if not death else LIVING_COST_DEATH,
            r_.m1, float(r_.raw_factor + 0), r_.m2, None,
            r_.m1 - r_.m2, float(r_.factor), int(r_.amount),
        ]
        from .hoffman import cumulative

        vals[8] = float(cumulative(r_.m1, float(c.legal_rate) / 100))
        vals[10] = float(cumulative(r_.m2, float(c.legal_rate) / 100))
        for i, v in enumerate(vals):
            fmt = NUM if i in (3, 5, 13) else (RATE if i in (8, 10, 12) else None)
            _put(ws, row, 2 + i, v, box=True, fmt=fmt,
                 fill=CAP_FILL if (i == 12 and r_.capped) else None)
        if r_.prorated_days:
            _put(ws, row, 17, "일자별 안분액 검토필요", box=True)
            _put(ws, row, 19, r_.prorated_days, box=True)
            _put(ws, row, 20, int(r_.prorated_amount), box=True, fmt=NUM)
        row += 1

    _put(ws, row, 17, "일실수입 전체합계 : ", font=HEAD_FONT)
    _put(ws, row, 20, int(r.income_total), font=HEAD_FONT, fmt=MONEY)
    row += 2

    # ---------------------------------------------------------- 손해 합계
    def money_line(label, value, bold=False):
        nonlocal row
        _put(ws, row, 2, label, font=HEAD_FONT if bold else BODY_FONT)
        _put(ws, row, 5, int(value), fmt=MONEY, font=HEAD_FONT if bold else BODY_FONT)
        row += 1

    row = _block(ws, row, "[소극손해]")
    money_line("일실수입", r.income_total)
    money_line("일실 퇴직금", c.severance)
    row += 1
    _put(ws, row, 2, "소극손해 합계(원) : ", font=HEAD_FONT)
    _put(ws, row, 14, int(r.passive_total), font=HEAD_FONT, fmt=MONEY)
    row += 2

    row = _block(ws, row, "[적극손해]")
    money_line("기왕 치료비", c.past_treatment)
    money_line("향후 치료비", r.treatment_total)
    money_line("기왕 개호비", r.past_caregiving_total)
    money_line("향후 개호비", r.future_caregiving_total)
    money_line("기왕 보조구", 0)
    money_line("향후 보조구", r.orthosis_total)
    row += 1
    _put(ws, row, 2, "적극손해 합계 : ", font=HEAD_FONT)
    _put(ws, row, 14, int(r.active_total), font=HEAD_FONT, fmt=MONEY)
    row += 2

    if death:
        # 장례비는 재산적 손해·합계에 넣지 않는다(calculate.py). [적극손해] 안에 두면 항목의 합과
        # '적극손해 합계'가 어긋나므로 블록을 따로 둔다.
        _put(ws, row, 2, "[장례비]", font=BLOCK_FONT)
        _put(ws, row, 4, "재산적 손해·과실상계·합계에 넣지 않음")
        row += 1
        money_line("장례비", c.funeral_cost)
        row += 1

    _put(ws, row, 2, "재산적 손해(소극손해 + 적극손해) : ", font=HEAD_FONT)
    _put(ws, row, 14, int(r.property_damage), font=HEAD_FONT, fmt=MONEY)
    row += 2

    st = r.settlement
    row = _block(ws, row, "[과실상계 전 공제]")
    _put(ws, row, 2, "과실상계 전 공제액 합계 : ", font=HEAD_FONT)
    _put(ws, row, 14, int(c.pre_offset_deduction), font=HEAD_FONT, fmt=MONEY)
    row += 2
    _put(ws, row, 2, "선공제 후 재산적 손해 : ", font=HEAD_FONT)
    _put(ws, row, 14, int(st["선공제_후_재산적손해"]), font=HEAD_FONT, fmt=MONEY)
    row += 2

    _put(ws, row, 2, "[과실상계]", font=BLOCK_FONT)
    _put(ws, row, 5, f"{c.fault_rate} %")
    row += 2
    _put(ws, row, 2, "과실상계 후 재산적 손해 : ", font=HEAD_FONT)
    _put(ws, row, 14, int(st["과실상계_후_표시값"]), font=HEAD_FONT, fmt=MONEY)
    row += 2

    row = _block(ws, row, "[공제]")
    _put(ws, row, 2, "기왕증 기여도")
    _put(ws, row, 5, f"{r.prior_contribution} %")
    row += 1
    _put(ws, row, 2, "지급 치료비")
    _put(ws, row, 5, f"{int(c.paid_cure):,} 원 중")
    row += 1
    _put(ws, row, 2, "기왕증 과실분")
    _put(ws, row, 5, int(st["기왕증_과실분"]), fmt=MONEY)
    row += 1
    _put(ws, row, 2, "손해배상 선급")
    _put(ws, row, 5, int(c.advance), fmt=MONEY)
    row += 1
    _put(ws, row, 2, "비율 공제액")
    _put(ws, row, 5, int(c.ratio_deduction), fmt=MONEY)
    row += 1
    _put(ws, row, 2, "전액 공제액")
    _put(ws, row, 5, int(c.full_deduction), fmt=MONEY)
    row += 2
    _put(ws, row, 2, "공제액 합계 : ", font=HEAD_FONT)
    _put(ws, row, 14, int(st["공제액_합계"]), font=HEAD_FONT, fmt=MONEY)
    row += 2
    _put(ws, row, 2, "공제 후 재산적 손해 : ", font=HEAD_FONT)
    _put(ws, row, 14, int(st["재산상손해_합계"]), font=HEAD_FONT, fmt=MONEY)
    row += 3

    row = _block(ws, row, "[위자료]")
    row = _headers(ws, row, ["순번", "원고", "위자료(자동계산)", None, "적용 위자료",
                             None, "재산상 손해", None, "재산손해 + 위자료"])
    _put(ws, row, 2, 1, box=True)
    _put(ws, row, 3, c.name, box=True)
    # 사망 사건은 자동계산 기준(노동능력상실률)이 없어 비워 둔다. 재산상 손해는 과실상계·공제 뒤 값이다.
    _put(ws, row, 4, None if death else int(r.solatium_auto), box=True, fmt=NUM)
    _put(ws, row, 6, int(c.solatium), box=True, fmt=NUM)
    _put(ws, row, 8, int(st["재산상손해_합계"]), box=True, fmt=NUM)
    _put(ws, row, 10, int(st["합계"]), box=True, fmt=NUM)
    row += 3

    for label, value in [
        ("소극손해 합계", r.passive_total),
        ("적극손해 합계", r.active_total),
        ("재산적 손해", r.property_damage),
        ("과실상계 전 공제액", c.pre_offset_deduction),
        ("원고측 과실비율액", st["원고측_과실비율액"]),
        ("공제액", st["공제액_합계"]),
    ]:
        money_line(label, value)
    row += 1
    _put(ws, row, 2, "재산상손해 합계 : ", font=HEAD_FONT)
    _put(ws, row, 14, int(st["재산상손해_합계"]), font=HEAD_FONT, fmt=MONEY)
    row += 2
    _put(ws, row, 2, "위자료 합계 : ", font=HEAD_FONT)
    _put(ws, row, 14, int(c.solatium), font=HEAD_FONT, fmt=MONEY)
    row += 2
    _put(ws, row, 2, "합계 : ", font=HEAD_FONT)
    _put(ws, row, 14, int(st["합계"]), font=HEAD_FONT, fmt=MONEY)

    _widths(ws)


def _detail_sheet(wb: Workbook, r: Result, title: str, build) -> None:
    ws = wb.create_sheet(title)
    _put(ws, 2, 2, title, font=TITLE_FONT)
    row = _basic_block(ws, r, 5)
    build(ws, row)
    _widths(ws)


def _treatment_sheet(wb, r):
    def build(ws, row):
        c = r.case
        _put(ws, row, 2, "[기왕 치료비]", font=BLOCK_FONT)
        row += 1
        row = _headers(ws, row, ["종류", None, "비용", None, "기왕증 기여도(%)", None, "비용 총액"])
        _put(ws, row + 1, 2, "기왕 치료비 합계 : ", font=HEAD_FONT)
        _put(ws, row + 1, 11, int(c.past_treatment), font=HEAD_FONT, fmt=MONEY)
        row += 3
        _put(ws, row, 2, "[향후 치료비]", font=BLOCK_FONT)
        row += 1
        row = _headers(ws, row, ["종류", "구분", "비용", "최초 필요일", "필요 최종일",
                                 "수명(월)", "기왕증 기여도(%)", None, "수치 합계", "비용 총액"])
        for src, cost in zip(c.treatments, r.treatment_rows):
            for i, v in enumerate([
                src.name, "반복" if src.repeating else "1회", int(src.cost),
                _d(src.first), _d(src.last), src.duration_month, float(src.prior), None,
                float(cost.factor_sum()), int(cost.total()),
            ]):
                _put(ws, row, 2 + i, v, box=True,
                     fmt=NUM if i in (2, 9) else (RATE if i == 8 else None),
                     fill=CAP_FILL if (i == 8 and cost.is_capped) else None)
            row += 1
        _put(ws, row + 1, 2, "향후 치료비 합계 : ", font=HEAD_FONT)
        _put(ws, row + 1, 11, int(r.treatment_total), font=HEAD_FONT, fmt=MONEY)

    _detail_sheet(wb, r, "치료비", build)


def _caregiving_sheet(wb, r):
    def build(ws, row):
        c = r.case
        _put(ws, row, 2, "[기왕 개호비]", font=BLOCK_FONT)
        row += 1
        row = _headers(ws, row, ["총일수", None, "개호비 단가", None, "기왕증(%)", None,
                                 "계산 개호비", None, "실제지출 개호비", None, "적용 개호비"])
        if c.past_caregiving_days:
            for i, v in enumerate([
                float(c.past_caregiving_days), None, int(c.past_caregiving_price), None,
                float(r.prior_contribution), None, None, None,
                int(c.past_caregiving_actual or 0), None, int(r.past_caregiving_total),
            ]):
                if v is not None:
                    _put(ws, row, 2 + i, v, box=True, fmt=NUM if i in (2, 8, 10) else None)
            row += 1
        _put(ws, row + 1, 2, "기왕 개호비 합계 : ", font=HEAD_FONT)
        _put(ws, row + 1, 11, int(r.past_caregiving_total), font=HEAD_FONT, fmt=MONEY)
        row += 3
        _put(ws, row, 2, "[향후 개호비]", font=BLOCK_FONT)
        if r.caregiving_basis:
            _put(ws, row, 4, f"개호비 단가 기준 : {r.caregiving_basis}")
        row += 1
        row = _headers(ws, row, ["순번", "기간초일", "기간말일", "개호비 단가", "인원",
                                 "기왕증(%)", "월비용", "M1", "호프만1", "M2", "호프만2",
                                 "M1-M2", "적용호프만", "기간 개호비"])
        from .hoffman import cumulative

        rate = float(c.legal_rate) / 100
        for cr in r.caregiving_rows:
            for i, v in enumerate([
                cr.step, _d(cr.start), _d(cr.end), int(cr.unit_price), float(cr.headcount),
                float(cr.prior_ratio), int(cr.monthly_cost), cr.m1,
                float(cumulative(cr.m1, rate)), cr.m2, float(cumulative(cr.m2, rate)),
                cr.m1 - cr.m2, float(cr.factor), int(cr.amount),
            ]):
                _put(ws, row, 2 + i, v, box=True,
                     fmt=NUM if i in (3, 6, 13) else (RATE if i in (8, 10, 12) else None),
                     fill=CAP_FILL if (i == 12 and cr.capped) else None)
            row += 1
        _put(ws, row + 1, 2, "향후 개호비 합계 : ", font=HEAD_FONT)
        _put(ws, row + 1, 11, int(r.future_caregiving_total), font=HEAD_FONT, fmt=MONEY)

    _detail_sheet(wb, r, "개호비", build)


def _orthosis_sheet(wb, r):
    def build(ws, row):
        _put(ws, row, 2, "[기왕 보조구]", font=BLOCK_FONT)
        row += 1
        row = _headers(ws, row, ["종류", None, "단가", None, "기왕증(%)", None, "비용 총액"])
        _put(ws, row + 1, 2, "기왕 보조구 합계 : ", font=HEAD_FONT)
        _put(ws, row + 1, 11, 0, font=HEAD_FONT, fmt=MONEY)
        row += 3
        _put(ws, row, 2, "[향후 보조구]", font=BLOCK_FONT)
        row += 1
        row = _headers(ws, row, ["종류", None, "단가", "최초 필요일", "필요 최종일",
                                 "수명(월)", "기왕증(%)", "수치 합계", None, "비용 총액"])
        for src, cost in zip(r.case.orthoses, r.orthosis_rows):
            for i, v in enumerate([
                src.name, None, int(src.cost), _d(src.first), _d(src.last),
                src.duration_month, float(src.prior), float(cost.factor_sum()), None,
                int(cost.total()),
            ]):
                if v is None:
                    continue
                _put(ws, row, 2 + i, v, box=True,
                     fmt=NUM if i in (2, 9) else (RATE if i == 7 else None),
                     fill=CAP_FILL if (i == 7 and cost.is_capped) else None)
            row += 1
        _put(ws, row + 1, 2, "향후 보조구 합계 : ", font=HEAD_FONT)
        _put(ws, row + 1, 11, int(r.orthosis_total), font=HEAD_FONT, fmt=MONEY)

    _detail_sheet(wb, r, "보조구", build)


def _severance_sheet(wb, r):
    def build(ws, row):
        _put(ws, row, 2, "[일실 퇴직금 산출]", font=BLOCK_FONT)
        row += 1
        row = _headers(ws, row, [None, None, "산출방식 A", None, "산출방식 B",
                                 None, "산출방식 C", None, "산출방식 D"])
        _put(ws, row + 1, 2, "일실 퇴직금(일반) 합계 : ", font=HEAD_FONT)
        _put(ws, row + 1, 11, int(r.case.severance), font=HEAD_FONT, fmt=MONEY)
        row += 3
        _put(ws, row, 2, "[산출 방식]", font=BLOCK_FONT)
        row += 1
        for name, desc in [
            ("산출 방식 A", "(정년퇴직금현가 - 기수령퇴직금) * 상실율"),
            ("산출 방식 B", "(정년시퇴직금현가 - 사고시의 계산상퇴직금) * 상실율"),
            ("산출 방식 C", "(정년시의 퇴직금현가 - 실제퇴직시 계산상퇴직금의 현가) * 상실율"),
            ("산출 방식 D", "(정년시의 퇴직금현가 - 실제퇴직시 받은 퇴직금의 현가) * 상실율"),
        ]:
            _put(ws, row, 2, name, font=HEAD_FONT)
            _put(ws, row, 4, desc)
            row += 1

    _detail_sheet(wb, r, "퇴직금(일반)", build)


def _warning_sheet(wb: Workbook, r: Result) -> None:
    """사람이 확인할 점. 노동 금액 계산표의 '경고' 시트와 같은 모양이다."""
    ws = wb.create_sheet("경고")
    _put(ws, 1, 1, "사람이 확인할 점", font=TITLE_FONT)
    for i, text in enumerate(r.warnings, 3):
        c = _put(ws, i, 1, text, fill=WARN_FILL)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 140
    ws.sheet_view.showGridLines = False


def _widths(ws):
    widths = {1: 2, 2: 16, 3: 12, 4: 13, 5: 13, 6: 13, 7: 12, 8: 12, 9: 11,
              10: 12, 11: 11, 12: 12, 13: 11, 14: 13, 15: 16, 16: 3, 17: 20,
              18: 3, 19: 10, 20: 18}
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.sheet_view.showGridLines = False


def write_workbook(result: Result, path: str | Path) -> Path:
    """손해배상액 계산표를 xlsx 로 저장한다."""
    wb = Workbook()
    wb.remove(wb.active)
    _summary_sheet(wb, result)
    _treatment_sheet(wb, result)
    _caregiving_sheet(wb, result)
    _orthosis_sheet(wb, result)
    _severance_sheet(wb, result)
    if result.warnings:
        _warning_sheet(wb, result)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
