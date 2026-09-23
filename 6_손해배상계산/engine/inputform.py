"""엑셀 입력서 — 빈 양식 생성과 채워진 양식 읽기.

YAML 대신 엑셀로 사건을 입력한다. 라벨은 B열, 값은 D열에 둔다.
표 블록(장해, 치료비, 보조구, 노임단가)은 머리글 아래로 행을 이어 쓴다.
표는 머리글 아래부터 다음 표 제목('[...]') 바로 앞까지 읽는다. 사이의 빈 행은 건너뛴다.
"""

from __future__ import annotations

import re
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .case import Case, ImpairmentInput, OrthosisInput, TreatmentInput, _d

LABEL = Font(name="맑은 고딕", size=10, bold=True)
BODY = Font(name="맑은 고딕", size=10)
BLOCK = Font(name="맑은 고딕", size=11, bold=True, color="1F4E79")
HINT = Font(name="맑은 고딕", size=9, color="808080", italic=True)
INPUT_FILL = PatternFill("solid", fgColor="FFF9E6")
HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# 라벨 -> Case 필드. 값은 D열에서 읽는다.
FIELDS = [
    ("[기초사항]", None, None),
    ("사건번호", "case_no", "예: 2024가단1234"),
    ("성명", "name", None),
    ("성별", "sex", "남 / 여"),
    ("사건종류", "case_type", "손해배상(자) / 손해배상(산)"),
    ("사건유형", "injury_type", "부상 / 사망"),
    ("판결/조정", "judgment_type", "판결 / 조정"),
    ("생년월일", "birth", "20240101 또는 2024-01-01"),
    ("사고일자", "accident", None),
    ("입원치료 종료일", "cure_end", "부상만. 없으면 비움"),
    ("가동연한(년)", "work_limit_years", "기본 65"),
    ("변론 종결일", "argument_end", None),
    ("법정이율(%)", "legal_rate", "기본 5"),
    ("기대여명(년)", "life_expectancy", "표시용"),
    ("여명 종료일", "life_end", "표시용"),
    ("[일실수입]", None, None),
    ("직종", "occupation", "예: 보통인부"),
    ("농촌노임 사용", "rural", "예 / 아니오"),
    ("[기타손해]", None, None),
    ("기왕 치료비", "past_treatment", None),
    ("기왕 개호비 총일수", "past_caregiving_days", None),
    ("기왕 개호비 단가", "past_caregiving_price", None),
    ("기왕 개호비 실제지출", "past_caregiving_actual", "실제 지출이 더 적으면 이 금액 적용"),
    ("향후 개호 시작일", "caregiving_start", None),
    ("향후 개호 종료일", "caregiving_end", None),
    ("향후 개호 인원", "caregiving_headcount", "예: 1 또는 0.5"),
    ("향후 개호 기간분할", "caregiving_month_mode", "월 단위 / 분기반기 단위 (기본 월 단위)"),
    ("일실 퇴직금", "severance", "직접 계산한 금액"),
    ("장례비", "funeral_cost", "사망이면 기본 5,000,000"),
    ("[과실상계 · 공제]", None, None),
    ("원고측 과실비율(%)", "fault_rate", None),
    ("과실상계 전 공제액", "pre_offset_deduction", "산재 휴업급여·장해급여 등"),
    ("지급 치료비", "paid_cure", None),
    ("손해배상 선급금", "advance", None),
    ("기타공제 - 비율공제", "ratio_deduction", None),
    ("기타공제 - 전액공제", "full_deduction", None),
    ("[위자료]", None, None),
    ("적용 위자료", "solatium", None),
]

DATE_FIELDS = {"birth", "accident", "cure_end", "argument_end", "life_end",
               "caregiving_start", "caregiving_end"}
INT_FIELDS = {"work_limit_years"}
DEC_FIELDS = {"legal_rate", "life_expectancy", "past_treatment", "past_caregiving_days",
              "past_caregiving_price", "past_caregiving_actual", "caregiving_headcount",
              "severance", "funeral_cost", "fault_rate", "pre_offset_deduction",
              "paid_cure", "advance", "ratio_deduction", "full_deduction", "solatium"}

TABLES = {
    "노동능력 상실률": ["진료과", "개별수치(%)", "기왕증(%)", "한시장해 년수"],
    "향후 치료비": ["종류", "비용", "최초 필요일", "필요 최종일", "수명(월)", "기왕증(%)"],
    "향후 보조구": ["종류", "단가", "최초 필요일", "필요 최종일", "수명(월)", "기왕증(%)"],
    "노임단가 직접입력": ["연도-반기", "단가"],
}
TABLE_ROWS = 8
WAGE_TABLE = "노임단가 직접입력"
WAGE_KEY_HINT = "예: 2024-1 (연도-반기. 농촌노임이면 연도-분기 1~4)"
# 표 끝을 정하는 제목 행. 표 사이 빈 행이 하나뿐이라 빈 행으로는 끝을 알 수 없다.
SECTION_TITLES = {f"[{t}]" for t in TABLES} | {label for label, f, _ in FIELDS if f is None}

_WAGE_KEY = re.compile(r"(\d{4})\s*-\s*(\d{1,2})")


def parse_wage_key(value, rural: bool = False) -> tuple[int, int]:
    """노임단가 직접입력의 '연도-반기' 값을 (연도, 반기) 로 읽는다. 농촌노임이면 (연도, 분기).

    엑셀은 일반 서식 칸에 적은 '2024-1' 을 날짜 2024-01-01 로 바꿔 저장한다. 그래서 1일인
    날짜로 읽히면 월을 반기(분기) 번호로 본다. 입력서와 사건 파일(wages)이 같이 쓴다.
    """
    top, unit = (4, "분기") if rural else (2, "반기")
    if isinstance(value, date):
        year, index, ok = value.year, value.month, value.day == 1
    else:
        m = _WAGE_KEY.fullmatch(str(value).strip())
        ok = m is not None
        year, index = (int(m.group(1)), int(m.group(2))) if m else (0, 0)
    if not ok or not 1 <= index <= top:
        raise ValueError(
            f"'{value}' 은(는) 연도-{unit} 값이 아닙니다. 2024-1 처럼 연도-{unit}({unit} 1~{top})으로 적으십시오"
        )
    return year, index


def make_template(path: str | Path) -> Path:
    """빈 입력서를 만든다."""
    wb = Workbook()
    ws = wb.active
    ws.title = "입력"
    ws["B2"] = "손해배상 계산 입력서"
    ws["B2"].font = Font(name="맑은 고딕", size=14, bold=True)
    ws["B3"] = "노란 칸만 채우고 01_입력 폴더에 넣으면 계산표가 나옵니다. 비운 칸은 0으로 봅니다."
    ws["B3"].font = HINT

    row = 5
    for label, fieldname, hint in FIELDS:
        if fieldname is None:
            row += 1
            ws.cell(row, 2, label).font = BLOCK
            row += 1
            continue
        ws.cell(row, 2, label).font = LABEL
        c = ws.cell(row, 4)
        c.fill = INPUT_FILL
        c.border = BOX
        c.font = BODY
        if hint:
            ws.cell(row, 6, hint).font = HINT
        row += 1

    for title, headers in TABLES.items():
        row += 1
        ws.cell(row, 2, f"[{title}]").font = BLOCK
        row += 1
        for i, h in enumerate(headers):
            c = ws.cell(row, 2 + i, h)
            c.font = LABEL
            c.fill = HEAD_FILL
            c.border = BOX
            c.alignment = Alignment(horizontal="center")
        if title == WAGE_TABLE:
            ws.cell(row, 6, WAGE_KEY_HINT).font = HINT
        row += 1
        for _ in range(TABLE_ROWS):
            for i in range(len(headers)):
                c = ws.cell(row, 2 + i)
                c.fill = INPUT_FILL
                c.border = BOX
                if title == WAGE_TABLE and i == 0:
                    c.number_format = "@"     # 텍스트 칸. 일반 서식이면 엑셀이 '2024-1' 을 날짜로 바꾼다
            row += 1

    for col, w in {2: 24, 3: 14, 4: 20, 5: 14, 6: 40, 7: 14}.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.sheet_view.showGridLines = False

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _clean(v):
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip().replace(",", "")
        return v or None
    return v


# strict 검사에서 채워져 있어야 하는 칸. 라벨을 그대로 오류 메시지에 쓴다.
REQUIRED_LABELS = [("생년월일", "birth"), ("사고일자", "accident")]
WAGE_BASIS_LABEL = "직종 또는 [노임단가 직접입력]"


def _check_required(filled: dict, case: Case) -> None:
    """빈 양식이 0원짜리 계산표로 나가는 것을 막는다.

    Case 에 기본값(생년월일 1990.1.1., 사고일자 2024.1.1., 직종 보통인부)이 있어
    다 읽힌 뒤에는 빈 칸을 알 수 없다. 그래서 Case 가 아니라 실제로 채워진
    칸(filled)을 본다.
    """
    missing = [label for label, fieldname in REQUIRED_LABELS if fieldname not in filled]
    if not case.impairments and not case.wages and "occupation" not in filled:
        missing.append(f"노동능력 상실률과 월소득·노임({WAGE_BASIS_LABEL})이 모두 비어 있음")
    if missing:
        raise ValueError(
            "입력서에 필수값이 비어 있습니다: "
            + ", ".join(missing)
            + ". 노란 칸을 채워 01_입력 에 다시 넣으십시오."
        )


def read_form(path: str | Path, strict: bool = False) -> Case:
    """채워진 입력서를 Case 로 읽는다.

    strict=True 면 필수값이 비었는지 검사한다. 자동 처리(watch.py)는 strict=True 로
    부른다. 빈 양식을 그대로 읽어야 하는 곳에서는 기본값(False)을 쓴다.
    """
    ws = load_workbook(path, data_only=True).active
    labels = {}
    for r in range(1, ws.max_row + 1):
        key = ws.cell(r, 2).value
        if isinstance(key, str):
            labels.setdefault(key.strip(), r)

    kw: dict = {}
    for label, fieldname, _hint in FIELDS:
        if fieldname is None or label not in labels:
            continue
        raw = _clean(ws.cell(labels[label], 4).value)
        if raw is None:
            continue
        try:
            if fieldname == "caregiving_month_mode":
                kw[fieldname] = "월" in str(raw)
            elif fieldname == "sex":
                kw[fieldname] = "F" if str(raw).startswith("여") else "M"
            elif fieldname == "rural":
                kw[fieldname] = str(raw).startswith("예")
            elif fieldname in DATE_FIELDS:
                kw[fieldname] = raw.date() if isinstance(raw, datetime) else _d(raw)
            elif fieldname in INT_FIELDS:
                kw[fieldname] = int(Decimal(str(raw)))
            elif fieldname in DEC_FIELDS:
                kw[fieldname] = Decimal(str(raw))
            else:
                kw[fieldname] = str(raw)
        except (ValueError, TypeError, ArithmeticError) as exc:
            raise ValueError(
                f"입력서 '{label}' 칸의 값 '{raw}' 을(를) 읽지 못했습니다({exc}). "
                "날짜는 20240101 또는 2024-01-01, 금액·비율은 숫자로 적으십시오."
            ) from exc

    section_rows = sorted(row for key, row in labels.items() if key in SECTION_TITLES)

    def table_rows(title, width):
        """(엑셀 행 번호, 값들). 머리글 아래부터 다음 표 제목 바로 앞까지, 빈 행은 건너뛴다."""
        head = labels.get(f"[{title}]")
        if head is None:
            return []
        end = next((row for row in section_rows if row > head), ws.max_row + 1)
        out = []
        for r in range(head + 2, end):
            vals = [_clean(ws.cell(r, 2 + i).value) for i in range(width)]
            if any(v is not None for v in vals):
                out.append((r, vals))
        return out

    def parse_rows(title, width, parse):
        """표 행을 읽는다. 실패하면 표 이름과 엑셀 행 번호를 붙여 알린다."""
        out = []
        for r, v in table_rows(title, width):
            try:
                item = parse(v)
            except (ValueError, TypeError, ArithmeticError) as exc:
                raise ValueError(
                    f"입력서 [{title}] 표 {r}행을 읽지 못했습니다({exc}). "
                    "날짜는 20240101 또는 2024-01-01, 금액·비율은 숫자로 적으십시오."
                ) from exc
            if item is not None:
                out.append(item)
        return out

    def impairment(v):
        if not v[0] or v[1] is None:
            return None
        return ImpairmentInput(
            dept=str(v[0]),
            rate=Decimal(str(v[1])),
            prior=Decimal(str(v[2] or 0)),
            years=Decimal(str(v[3])) if v[3] else None,
        )

    def cost_row(cls):
        def parse(v):
            if not v[0] or v[1] is None:
                return None
            first = v[2].date() if isinstance(v[2], datetime) else _d(v[2])
            last = v[3].date() if isinstance(v[3], datetime) else (_d(v[3]) if v[3] else first)
            return cls(
                name=str(v[0]),
                cost=Decimal(str(v[1])),
                first=first,
                last=last,
                duration_month=int(Decimal(str(v[4] or 1))),
                prior=Decimal(str(v[5] or 0)),
                repeating=last != first,
            )
        return parse

    rural = kw.get("rural", False)

    def wage(v):
        if not v[0] or v[1] is None:
            return None
        year, index = parse_wage_key(v[0], rural)
        return f"{year}-{index}", int(Decimal(str(v[1])))

    kw["impairments"] = parse_rows("노동능력 상실률", 4, impairment)
    kw["treatments"] = parse_rows("향후 치료비", 6, cost_row(TreatmentInput))
    kw["orthoses"] = parse_rows("향후 보조구", 6, cost_row(OrthosisInput))
    kw["wages"] = dict(parse_rows(WAGE_TABLE, 2, wage))
    case = Case(**kw)
    if strict:
        _check_required(kw, case)
    return case
