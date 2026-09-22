"""사건 YAML -> 계산 -> 엑셀표 전 과정이 골든 케이스 3을 재현하는지 확인한다."""

import sys
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from engine.calculate import calculate
from engine.case import load
from engine.excel import write_workbook
from tests.golden import case_fault_deduction as F


@pytest.fixture(scope="module")
def book(tmp_path_factory):
    case = load(ROOT / "cases" / "sample.yaml")
    table = {tuple(int(x) for x in k.split("-")): v for k, v in case.wages.items()}
    last = table[max(table)]
    result = calculate(
        case,
        wage_of=lambda p: table.get((p.year, p.index), last),
        has_wage=lambda y, i: (y, i) in table,
    )
    path = write_workbook(result, tmp_path_factory.mktemp("out") / "계산표.xlsx")
    return openpyxl.load_workbook(path), result


def test_시트_구성(book):
    wb, _ = book
    assert wb.sheetnames == ["종합", "치료비", "개호비", "보조구", "퇴직금(일반)"]


def test_기초사항_라벨과_값(book):
    wb, _ = book
    ws = wb["종합"]
    assert ws["B2"].value == "손해배상액 계산표"
    assert ws["B6"].value == "사건번호" and ws["H6"].value == "손해배상(자)"
    assert ws["L6"].value == "34세 0개월 0일"
    assert ws["H10"].value == "2054.12.31"     # 가동종료일 = 생년월일 + 65년 - 1일


def test_일실수입_순번이_골든과_같다(book):
    wb, _ = book
    ws = wb["종합"]
    header = next(r for r in range(1, 60) if ws.cell(r, 2).value == "순번")
    for i, g in enumerate(F.INCOME_ROWS):
        row = header + 1 + i
        assert ws.cell(row, 2).value == g[0]
        assert ws.cell(row, 3).value == f"{g[1]:%Y.%m.%d}"
        assert ws.cell(row, 4).value == f"{g[2]:%Y.%m.%d}"
        assert ws.cell(row, 5).value == g[3]       # 노임단가
        assert ws.cell(row, 7).value == g[5]       # 월소득
        assert Decimal(str(ws.cell(row, 8).value)) == Decimal(g[6])    # 상실률
        assert ws.cell(row, 9).value == g[7]       # M1
        assert Decimal(str(ws.cell(row, 10).value)) == Decimal(g[8])   # 호프만1
        assert Decimal(str(ws.cell(row, 14).value)) == Decimal(g[11])  # 적용호프만
        assert ws.cell(row, 15).value == g[12]     # 기간 일실수입


def test_합계_블록(book):
    wb, r = book
    ws = wb["종합"]
    found = {}
    for row in range(1, ws.max_row + 1):
        label = ws.cell(row, 2).value
        if isinstance(label, str) and label.endswith(" : "):
            found[label.strip(" :")] = ws.cell(row, 14).value or ws.cell(row, 20).value
    assert found["재산적 손해(소극손해 + 적극손해)"] == int(F.PROPERTY_DAMAGE)
    assert found["공제액 합계"] == int(F.DEDUCTION_TOTAL)
    assert found["공제 후 재산적 손해"] == int(F.PROPERTY_FINAL)
    assert found["재산상손해 합계"] == int(F.PROPERTY_FINAL)
    assert found["합계"] == int(F.GRAND_TOTAL)


def test_과실상계_공제_표시(book):
    wb, _ = book
    ws = wb["종합"]
    vals = {ws.cell(r, 2).value: ws.cell(r, 5).value for r in range(1, ws.max_row + 1)}
    assert vals["기왕증 과실분"] == int(F.PAID_CURE_DEDUCTION)
    assert vals["손해배상 선급"] == int(F.ADVANCE)
    assert vals["원고측 과실비율액"] == int(F.FAULT_SHARE)


def test_위자료_자동계산(book):
    wb, _ = book
    ws = wb["종합"]
    row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 2).value == 1
               and ws.cell(r, 3).value == "홍길동")
    assert ws.cell(row, 4).value == int(F.SOLATIUM_AUTO)
    assert ws.cell(row, 6).value == int(F.SOLATIUM_APPLIED)
