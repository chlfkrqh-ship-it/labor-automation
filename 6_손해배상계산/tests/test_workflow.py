"""입력서 -> 자동 처리 -> 계산표 전 과정 확인 (골든 케이스 3 재현)."""

import os
import shutil
import subprocess
import sys
from dataclasses import replace
from datetime import date, datetime
from pathlib import Path

import openpyxl
import pytest
import yaml

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import cli
import extract
import watch
from engine.calculate import calculate
from engine.case import load
from engine.inputform import TABLE_ROWS, make_template, read_form
from tests.golden import case_fault_deduction as F

SAMPLE = ROOT / "cases" / "sample.yaml"

FORM_VALUES = {
    "사건번호": "2024가단1", "성명": "홍길동", "성별": "남",
    "사건종류": "손해배상(자)", "사건유형": "부상", "판결/조정": "판결",
    "생년월일": "19900101", "사고일자": "20240101", "입원치료 종료일": "20240630",
    "가동연한(년)": 65, "변론 종결일": "20260308", "법정이율(%)": 5,
    "기대여명(년)": 47.6, "여명 종료일": "20710727", "직종": "보통인부",
    "원고측 과실비율(%)": 30, "지급 치료비": 1000000,
    "손해배상 선급금": 2000000, "적용 위자료": 20000000,
}
IMPAIRMENTS = [("정형외과", 40, 20), ("안과", 20, None)]
WAGES = [("2024-1", 165545), ("2024-2", 167081), ("2025-1", 169804),
         ("2025-2", 171037), ("2026-1", 172068)]


def _index(ws) -> dict:
    return {ws.cell(r, 2).value: r for r in range(1, ws.max_row + 1)
            if isinstance(ws.cell(r, 2).value, str)}


def _totals(xlsx: Path) -> dict:
    ws = openpyxl.load_workbook(xlsx)["종합"]
    totals = {}
    for row in range(1, ws.max_row + 1):
        for label_col, value_col in ((2, 14), (17, 20)):
            label = ws.cell(row, label_col).value
            if isinstance(label, str) and label.endswith(" : "):
                totals[label.strip(" :")] = ws.cell(row, value_col).value
    return totals


def _fill(path: Path) -> Path:
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    idx = _index(ws)
    for k, v in FORM_VALUES.items():
        ws.cell(idx[k], 4, v)
    r = idx["[노동능력 상실률]"] + 2
    for i, (dept, rate, prior) in enumerate(IMPAIRMENTS):
        ws.cell(r + i, 2, dept)
        ws.cell(r + i, 3, rate)
        if prior:
            ws.cell(r + i, 4, prior)
    r = idx["[노임단가 직접입력]"] + 2
    for i, (k, v) in enumerate(WAGES):
        ws.cell(r + i, 2, k)
        ws.cell(r + i, 3, v)
    wb.save(path)
    return path


@pytest.fixture(scope="module")
def workspace(tmp_path_factory):
    root = watch.ensure_root(tmp_path_factory.mktemp("노동사건자동화"))
    filled = shutil.copy(root / "00_양식" / "입력서.xlsx", root / "01_입력" / "사건.xlsx")
    _fill(Path(filled))
    processed = watch.scan_once(root)
    return root, processed


def test_빈_양식이_만들어진다(tmp_path):
    form = make_template(tmp_path / "입력서.xlsx")
    assert form.exists()
    case = read_form(form)          # 빈 양식도 읽혀야 한다
    assert case.injury_type == "부상"
    assert case.work_limit_years == 65


def test_입력서를_읽으면_사건이_된다(tmp_path):
    form = _fill(make_template(tmp_path / "입력서.xlsx"))
    case = read_form(form)
    assert case.case_no == "2024가단1"
    assert case.birth == F.BIRTH and case.accident == F.ACCIDENT
    assert case.cure_end == F.CURE_END
    assert case.resolved_work_end() == F.WORK_END
    assert [i.dept for i in case.impairments] == ["정형외과", "안과"]
    assert case.impairments[0].prior == 20
    assert case.fault_rate == F.FAULT_RATE
    assert len(case.wages) == 5


def test_입력폴더에_넣으면_결과가_나온다(workspace):
    root, processed = workspace
    assert processed == 1
    assert not list((root / "01_입력").glob("*.xlsx"))     # 입력은 비워지고
    assert list((root / "03_보관").glob("*.xlsx"))          # 보관으로 이동
    assert not list((root / "99_오류").iterdir())           # 오류 없음
    results = list((root / "02_결과").glob("*.xlsx"))
    assert len(results) == 1
    assert "2024가단1(홍길동)" in results[0].name


def test_결과가_골든과_일치한다(workspace):
    root, _ = workspace
    totals = _totals(next((root / "02_결과").glob("*.xlsx")))
    assert totals["일실수입 전체합계"] == int(F.INCOME_TOTAL)
    assert totals["재산상손해 합계"] == int(F.PROPERTY_FINAL)
    assert totals["합계"] == int(F.GRAND_TOTAL)


def test_잘못된_입력은_오류폴더로(tmp_path_factory):
    root = watch.ensure_root(tmp_path_factory.mktemp("bad"))
    shutil.copy(root / "00_양식" / "입력서.xlsx", root / "01_입력" / "빈양식.xlsx")
    (root / "01_입력" / "빈양식.xlsx").touch()
    watch.scan_once(root)
    assert list((root / "99_오류").glob("*.txt")), "오류 사유 파일이 남아야 한다"
    assert not list((root / "02_결과").iterdir())


# ── 노임 조회 (cli._wage_lookup — cli·watch 공용) ──────────────────────────


def _calc(case):
    wage_of, has_wage, notes = cli._wage_lookup(case)
    care, care_notes = cli._caregiving_lookup(case)
    return calculate(case, wage_of, has_wage, **care), notes + care_notes


def _first_day_after(key, rural=False):
    """노임표 키 key 다음 구간에 드는 날(그 구간 첫날 + 9일)."""
    y, i = key
    if rural:
        return date(y, 3 * i + 1, 10) if i < 4 else date(y + 1, 1, 10)
    return date(y, 5, 10) if i == 1 else date(y, 9, 10)


def test_노임표_마지막_반기_뒤의_사고와_개호도_마지막_단가로_계산된다():
    if not (ROOT / "data" / "wage_occupation.csv").exists():
        pytest.skip("data/ 노임표가 없다")
    from engine.tables import WageTable

    wt = WageTable.load()
    base = load(SAMPLE)                                   # 보통인부
    data = replace(base, wages={})                        # data/ 노임표를 쓴다
    oid = wt.find_occupation(base.occupation)[0]["id"]
    last = max((y, h) for (o, y, h) in wt.by_half if o == oid)
    last_wage = wt.occupation_wage(oid, *last)
    # 같은 단가를 적은 직접입력(사고일 2024. 1. 1.이 속한 반기부터 마지막 반기까지)
    table = {f"{y}-{h}": wt.occupation_wage(oid, y, h) for (o, y, h) in wt.by_half if o == oid and y >= 2024}

    after = _first_day_after(last)                        # 마지막이 2026년 상반기면 2026. 5. 10.
    later = _first_day_after((after.year, 2) if after.month == 5 else (after.year + 1, 1))
    for accident in (after, later):
        case = replace(data, accident=accident, cure_end=None)
        result, notes = _calc(case)                       # 전에는 KeyError
        assert {r.wage for r in result.income_rows} == {last_wage}
        assert len(notes) == 1 and "마지막 공표 단가" in notes[0]
        direct, _ = _calc(replace(case, wages=table))
        assert result.settlement["합계"] == direct.settlement["합계"]

    for month_mode in (True, False):                      # 향후 개호 시작이 마지막 단가 뒤
        care = replace(data, caregiving_start=later, caregiving_end=date(2071, 7, 27),
                       caregiving_month_mode=month_mode, caregiving_occupation=base.occupation)
        result, notes = _calc(care)
        assert result.caregiving_rows and {r.unit_price for r in result.caregiving_rows} == {last_wage}
        assert any("마지막 공표 단가" in n for n in notes)
        direct, _ = _calc(replace(care, caregiving_occupation="", caregiving_wages=table))
        assert result.future_caregiving_total == direct.future_caregiving_total
    with pytest.raises(ValueError, match="caregiving_occupation"):     # 개호 단가 기준이 없으면 멈춘다
        _calc(replace(data, caregiving_start=later, caregiving_end=date(2071, 7, 27)))
    with pytest.raises(cli.InputError, match="caregiving_wages"):      # 농촌 개호 단가는 표로
        _calc(replace(data, caregiving_start=later, caregiving_end=date(2071, 7, 27),
                      caregiving_occupation=base.occupation, caregiving_rural=True))

    q_last = max((int(r["year"]), int(r["quarter"])) for r in wt.quarterly)
    rural, notes = _calc(replace(data, rural=True, accident=_first_day_after(q_last, rural=True), cure_end=None))
    assert {r.wage for r in rural.income_rows} == {wt.rural_wage(*q_last, "M")}
    assert "분기부터" in notes[0]

    # 공표가 끊긴 직종이면 '새 노임을 기다리라'가 아니라 직종을 확인하라고 알린다
    for occ in wt.occupations:
        own = [(y, h) for (o, y, h) in wt.by_half if o == occ["id"]]
        if own and max(own) < last and wt.find_occupation(occ["name"])[0]["id"] == occ["id"]:
            stale = max(own)
            result, notes = _calc(replace(data, occupation=occ["name"], accident=_first_day_after(stale),
                                          cure_end=None))
            assert {r.wage for r in result.income_rows} == {wt.occupation_wage(occ["id"], *stale)}
            assert "공표되지 않았습니다" in notes[0]
            break



def test_농촌_사건은_직종을_찾지_않고_직종명이_다르면_후보를_보여_준다(tmp_path, capsys):
    if not (ROOT / "data" / "wage_occupation.csv").exists():
        pytest.skip("data/ 노임표가 없다")
    base = replace(load(SAMPLE), wages={}, cure_end=None)
    rural, _ = _calc(replace(base, rural=True, occupation="농업"))      # 노임표에 '농업' 직종은 없다
    assert rural.income_rows and rural.wage_basis.startswith("농촌 일용노임")
    with pytest.raises(cli.InputError, match="후보: .*배관공"):
        cli._wage_lookup(replace(base, occupation="배관"))
    # 고른 직종은 계산표에 적히고, 엔진 경고는 요약의 '확인:' 줄로 나온다
    cli.run_injury(replace(base, injury_type="사망", funeral_cost=5000000), tmp_path / "사망.xlsx")
    out = capsys.readouterr().out
    assert "확인:" in out and "장례비" in out and "자동계산" not in out

def test_노임표_중간에_빠진_반기는_알아볼_수_있는_오류로_멈춘다():
    if not (ROOT / "data" / "wage_occupation.csv").exists():
        pytest.skip("data/ 노임표가 없다")
    case = replace(load(SAMPLE), wages={}, birth=date(1960, 1, 1), accident=date(1994, 10, 1),
                   cure_end=None)                          # 1995년 상반기는 노임표에 없다
    with pytest.raises(cli.InputError, match="1995년 상반기 단가가 노임표에 없습니다"):
        _calc(case)


def test_직접입력_노임에_빠진_반기가_있으면_멈춘다():
    base = load(SAMPLE)
    full, notes = _calc(base)
    assert int(full.settlement["합계"]) == int(F.GRAND_TOTAL) and notes == []

    partial = replace(base, wages={k: v for k, v in base.wages.items() if not k.startswith("2024")})
    with pytest.raises(cli.InputError, match=r"2024-1, 2024-2 이\(가\) 없습니다"):
        _calc(partial)                                    # 전에는 2024년에도 2026년 단가를 조용히 썼다
    gap = replace(base, wages={k: v for k, v in base.wages.items() if k != "2025-1"})
    with pytest.raises(cli.InputError, match=r"2025-1 이\(가\) 없습니다"):
        _calc(gap)                                        # 전에는 2024년 하반기 단가를 끝까지 이었다
    single, _ = _calc(replace(base, wages={"2024-1": 172068}))   # 한 단가를 전 기간에
    assert {r.wage for r in single.income_rows} == {172068}
    with pytest.raises(cli.InputError, match="연도-반기"):
        _calc(replace(base, wages={"2024-3": 172068}))


# ── 사건 파일 필수값 ─────────────────────────────────────────────────────


def test_생년월일이나_사고일자가_없는_사건_파일은_멈춘다(tmp_path, monkeypatch):
    import yaml

    raw = yaml.safe_load(SAMPLE.read_text(encoding="utf-8"))
    del raw["birth"]
    raw["accident"] = None
    path = tmp_path / "사건.yaml"
    path.write_text(yaml.safe_dump(raw, allow_unicode=True), encoding="utf-8")

    with pytest.raises(cli.InputError, match=r"생년월일\(birth\), 사고일자\(accident\)"):
        cli.load_injury(path)
    monkeypatch.setattr(cli, "utf8_console", lambda: None)
    monkeypatch.setattr(sys, "argv", ["cli.py", str(path), "-o", str(tmp_path / "x.xlsx")])
    with pytest.raises(SystemExit, match="입력 오류: .*생년월일"):
        cli.main()
    assert not (tmp_path / "x.xlsx").exists()

    root = watch.ensure_root(tmp_path / "감시")          # 전에는 1990. 1. 1.·2024. 1. 1.로 계산했다
    shutil.copy(path, root / "01_입력" / "사건.yaml")
    assert watch.scan_once(root) == 0
    assert "생년월일" in next((root / "99_오류").glob("*.txt")).read_text(encoding="utf-8")
    assert not list((root / "02_결과").iterdir())


# ── 입력서 읽기 ─────────────────────────────────────────────────────────



def test_kind_가_빠진_노동_사건_파일은_무엇을_적을지_알린다(tmp_path):
    raw = yaml.safe_load((ROOT / "cases" / "labor_sample.yaml").read_text(encoding="utf-8"))
    raw.pop("kind")
    path = tmp_path / "사건.yaml"
    path.write_text(yaml.safe_dump(raw, allow_unicode=True), encoding="utf-8")
    with pytest.raises(cli.InputError, match="kind: labor"):
        cli.load_injury(path)

def test_노임_키가_엑셀에서_날짜로_바뀌어도_읽힌다(tmp_path):
    blank = openpyxl.load_workbook(make_template(tmp_path / "빈양식.xlsx")).active
    r0 = _index(blank)["[노임단가 직접입력]"] + 2
    assert all(blank.cell(r0 + i, 2).number_format == "@" for i in range(TABLE_ROWS))

    form = _fill(make_template(tmp_path / "입력서.xlsx"))
    wb = openpyxl.load_workbook(form)
    ws = wb.active
    r = _index(ws)["[노임단가 직접입력]"] + 2
    for i, (key, _) in enumerate(WAGES):
        y, h = map(int, key.split("-"))
        ws.cell(r + i, 2, datetime(y, h, 1)).number_format = "yyyy-mm"   # 엑셀이 '2024-1' 을 바꾼 모양
    wb.save(form)
    assert list(read_form(form, strict=True).wages) == [k for k, _ in WAGES]

    root = watch.ensure_root(tmp_path / "감시")
    shutil.copy(form, root / "01_입력" / "사건.xlsx")
    assert watch.scan_once(root) == 1                     # 전에는 int('01 00:00:00') 에서 멈췄다
    assert _totals(next((root / "02_결과").glob("*.xlsx")))["합계"] == int(F.GRAND_TOTAL)

    ws.cell(r, 2, 2024.1)                                 # 숫자로 들어간 키는 알아볼 수 있게 멈춘다
    wb.save(form)
    with pytest.raises(ValueError, match=r"\[노임단가 직접입력\] 표 \d+행.*연도-반기"):
        read_form(form, strict=True)


def test_표_8행을_다_채워도_다음_표를_읽지_않는다(tmp_path):
    form = _fill(make_template(tmp_path / "입력서.xlsx"))
    wb = openpyxl.load_workbook(form)
    ws = wb.active
    idx = _index(ws)
    r = idx["[노동능력 상실률]"] + 2
    for i in range(TABLE_ROWS):
        ws.cell(r + i, 2, f"진료과{i}")
        ws.cell(r + i, 3, 5)
    for title in ("[향후 치료비]", "[향후 보조구]"):
        r = idx[title] + 2
        for i in range(TABLE_ROWS):
            ws.cell(r + i, 2, f"항목{i}")
            ws.cell(r + i, 3, 100000)
            ws.cell(r + i, 4, "20260101")
    wb.save(form)
    case = read_form(form, strict=True)                   # 전에는 다음 표 머리글을 읽다 멈췄다
    assert (len(case.impairments), len(case.treatments), len(case.orthoses)) == (8, 8, 8)

    r = idx["[향후 치료비]"] + 2
    for i in (5, 6):                                      # 가운데 빈 행 두 개 뒤의 행도 읽는다
        for c in range(2, 8):
            ws.cell(r + i, c).value = None
    wb.save(form)
    assert len(read_form(form, strict=True).treatments) == 6

    ws.cell(r, 4, "날짜아님")
    wb.save(form)
    with pytest.raises(ValueError, match=rf"\[향후 치료비\] 표 {r}행"):
        read_form(form, strict=True)


# ── 감시 폴더 ──────────────────────────────────────────────────────────


def _labor_input(root, sample="labor_sample.yaml", name="사건.yaml"):
    shutil.copy(ROOT / "cases" / sample, root / "01_입력" / name)


def test_감시_보관_이동이_실패해도_오류로_적지_않고_계속한다(tmp_path, monkeypatch, capsys):
    root = watch.ensure_root(tmp_path / "감시")
    _labor_input(root)
    real_move = shutil.move

    def archive_fails(src, dst, *a, **k):
        if "03_보관" in str(dst):
            raise PermissionError("다른 프로그램이 사용 중")
        return real_move(src, dst, *a, **k)

    monkeypatch.setattr(watch.shutil, "move", archive_fails)
    assert watch.scan_once(root) == 1                     # 전에는 99_오류 이동에서 예외로 끝났다
    assert len(list((root / "02_결과").glob("*.xlsx"))) == 1
    assert not list((root / "99_오류").iterdir())
    assert "[경고]" in capsys.readouterr().out
    assert watch.scan_once(root) == 0                     # 처리중 파일을 다시 계산하지 않는다
    assert len(list((root / "02_결과").glob("*.xlsx"))) == 1


def test_감시_처리_도중_입력이_사라지거나_오류_이동이_실패해도_멈추지_않는다(tmp_path, monkeypatch):
    root = watch.ensure_root(tmp_path / "감시")
    _labor_input(root)
    real_process = watch.process

    def process_then_vanish(path, root_, stamp=None):
        out = real_process(path, root_, stamp)
        path.unlink()                                     # 다른 PC 의 감시가 먼저 옮긴 상황
        return out

    monkeypatch.setattr(watch, "process", process_then_vanish)
    assert watch.scan_once(root) == 1
    assert not list((root / "99_오류").iterdir())

    monkeypatch.setattr(watch, "process", real_process)
    (root / "01_입력" / "빈.yaml").write_text("case_no: 1\n", encoding="utf-8")
    real_move = shutil.move

    def error_move_fails(src, dst, *a, **k):
        if "99_오류" in str(dst):
            raise PermissionError("다른 프로그램이 사용 중")
        return real_move(src, dst, *a, **k)

    monkeypatch.setattr(watch.shutil, "move", error_move_fails)
    assert watch.scan_once(root) == 0
    assert list((root / "99_오류").glob("빈_*.txt"))


def test_감시_열려_있는_입력은_건너뛰고_닫히면_처리한다(tmp_path, monkeypatch):
    root = watch.ensure_root(tmp_path / "감시")
    _labor_input(root)
    real_rename = os.rename

    def locked(src, dst):
        raise PermissionError(13, "다른 프로세스가 파일을 사용 중입니다")

    monkeypatch.setattr(watch.os, "rename", locked)
    assert watch.scan_once(root) == 0                     # 계산도 오류 기록도 하지 않는다
    assert (root / "01_입력" / "사건.yaml").exists()
    assert not list((root / "02_결과").iterdir()) and not list((root / "99_오류").iterdir())
    monkeypatch.setattr(watch.os, "rename", real_rename)
    assert watch.scan_once(root) == 1


def test_감시_루프는_예상_밖_오류에도_계속되고_남은_처리중_파일을_되살린다(tmp_path, monkeypatch, capsys):
    root = watch.ensure_root(tmp_path / "감시")
    (root / "01_입력" / watch.CLAIM).mkdir()
    _labor_input(root, name=f"{watch.CLAIM}/사건.yaml")     # 지난번 감시가 처리 도중 멈춘 파일
    monkeypatch.setattr(watch, "utf8_console", lambda: None)
    monkeypatch.setattr(sys, "argv", ["watch.py", "--root", str(root)])
    watch.main()
    assert list((root / "02_결과").glob("*노동금액계산표*.xlsx"))
    assert not list((root / "01_입력" / watch.CLAIM).iterdir())

    calls = []

    def flaky(root_):
        calls.append(root_)
        if len(calls) == 1:
            raise RuntimeError("예상 밖")
        raise KeyboardInterrupt

    monkeypatch.setattr(watch, "scan_once", flaky)
    monkeypatch.setattr(watch.time, "sleep", lambda s: None)
    monkeypatch.setattr(sys, "argv", ["watch.py", "--root", str(root), "--loop"])
    watch.main()
    assert len(calls) == 2 and "종료합니다" in capsys.readouterr().out


def test_같은_이름의_입력도_보관에_모두_남고_결과에_사건번호가_붙는다(tmp_path):
    root = watch.ensure_root(tmp_path / "감시")
    for sample in ("labor_dismissal_sample.yaml", "labor_sample.yaml"):
        _labor_input(root, sample)
        assert watch.scan_once(root) == 1
    assert len(list((root / "03_보관").glob("사건_*.yaml"))) == 2       # 앞 사건 입력을 덮어쓰지 않는다
    names = sorted(p.name for p in (root / "02_결과").glob("*.xlsx"))
    assert [n.split("_노동금액계산표_")[0] for n in names] == ["2024가합0000(김가상)", "2025가단0000(홍길동)"]
    assert watch._label("2024가합1/2025가합2", "홍길동", "사건") == "2024가합1_2025가합2(홍길동)"


# ── 자료 추출·콘솔 출력 ──────────────────────────────────────────────────


def test_추출은_이전_계산표를_자료로_읽지_않는다(tmp_path):
    folder = tmp_path / "사건"
    (folder / "계산").mkdir(parents=True)
    wb = openpyxl.Workbook()
    wb.active["A1"] = "이전 계산 퇴직금 7,913,075 원"
    wb.save(folder / "계산" / "2025가단0000(홍길동)_노동금액계산표.xlsx")
    wb.save(folder / "2025가단0000(홍길동)_노동금액계산표.xlsx")          # 계산 폴더 = 자료 폴더인 경우
    wb = openpyxl.Workbook()
    wb.active["A1"] = "을2 별지 계산표"
    wb.save(folder / "을2_별지계산표.xlsx")                              # 상대방 계산표는 자료다
    (folder / "급여명세서.txt").write_text("기본급 3,000,000원", encoding="utf-8")

    extract.main(str(folder))
    text = (folder / "_추출.txt").read_text(encoding="utf-8")
    head, _, body = text.partition("=" * 70)
    assert "[계산 결과물 — 자료가 아니므로 읽지 않음]" in head and "노동금액계산표" in head
    assert "이전 계산" not in body and "노동금액계산표" not in body
    assert "기본급 3,000,000원" in body and "을2 별지 계산표" in body


def test_cp949_파이프에서도_요약이_끝까지_나온다(tmp_path):
    env = {**os.environ, "PYTHONIOENCODING": "cp949", "PYTHONDONTWRITEBYTECODE": "1"}
    for sample in ("labor_sample.yaml", "sample.yaml"):   # 노동 요약의 '—' 에서 멈추던 것
        out = tmp_path / f"{sample}.xlsx"
        proc = subprocess.run([sys.executable, "-B", "cli.py", str(ROOT / "cases" / sample), "-o", str(out)],
                              cwd=ROOT, env=env, capture_output=True)
        assert proc.returncode == 0, proc.stderr.decode("utf-8", "replace")
        text = proc.stdout.decode("utf-8")                # 한글 항목명이 UTF-8 로 나온다
        assert "합계" in text and f"저장: {out}" in text and out.exists()

    folder = tmp_path / "자료"
    folder.mkdir()
    (folder / "진술서—보충.txt").write_text("내용", encoding="utf-8")
    proc = subprocess.run([sys.executable, "-B", "extract.py", str(folder)], cwd=ROOT, env=env, capture_output=True)
    assert proc.returncode == 0, proc.stderr.decode("utf-8", "replace")
    assert "진술서—보충.txt" in proc.stdout.decode("utf-8")
