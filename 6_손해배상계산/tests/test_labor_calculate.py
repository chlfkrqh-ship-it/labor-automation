"""노동 금액 조립 확인 — cases/labor_sample.yaml 을 cli·감시 폴더와 같은 경로로 끝까지 돌린다.

개별 금액의 정답은 모듈별 테스트(판결 숫자)가 맡는다. 여기서는 모듈 사이 값 전달과 출력 형식을 본다.
"""

import shutil
import sys
from datetime import date
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest
import yaml

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import cli
import watch
from engine.labor import average, retirement
from engine.labor.calculate import calculate_labor, load_labor_case, summary_lines
from engine.labor.common import LaborError, dec
from engine.labor.excel import WARN_FILL, write_labor_workbook

SAMPLE = ROOT / "cases" / "labor_sample.yaml"
DISMISSAL = ROOT / "cases" / "labor_dismissal_sample.yaml"


def _raw(path=SAMPLE) -> dict:
    return yaml.safe_load(Path(path).read_text(encoding="utf-8"))


def _sheet(res, tmp_path, name):
    return openpyxl.load_workbook(write_labor_workbook(res, tmp_path / "노동.xlsx"))[name]


@pytest.fixture(scope="module")
def result():
    return calculate_labor(load_labor_case(SAMPLE))


def test_있는_절이_모두_계산된다(result):
    assert list(result.parts) == ["ordinary", "overtime", "leave", "average_wage", "retirement", "interest"]


def test_원금은_모듈_합계와_같다(result):
    parts = result.parts
    expected = sum((parts[s].total for s in ("overtime", "leave", "retirement")), Decimal(0))
    assert result.principal_total == expected
    assert {c.category for c in result.claims} >= {"퇴직금 차액"}


def test_평균임금에_시간외수당_증가분이_넘어간다(result):
    extra = result.parts["overtime"].extra_wages
    assert extra and all(v >= 0 for v in extra.values())


def test_퇴직금은_평균임금_결과를_쓴다(result):
    avg = result.parts["average_wage"].average_daily_wage
    assert any(getattr(r, "average_daily_wage", None) == avg for r in result.parts["retirement"].rows)


def test_연차수당_1일분은_통상임금_모듈에서_온다(result):
    ordinary = result.parts["ordinary"]
    rows = [r for r in result.parts["leave"].rows if r.daily_wage]
    assert rows
    for r in rows:
        assert r.daily_wage == ordinary.daily_ordinary_of(r.wage_ref_date, allowance="annual_leave")


def test_지연손해금은_모든_원금을_받는다(result):
    it = result.parts["interest"]
    assert len(it.schedules) == len([c for c in result.claims if c.amount > 0])
    settle = [s for s in it.schedules if s.claim.settlement]
    # 퇴직금 차액: 마지막 근무일 2025. 6. 30. + 15일부터 20%, 선고일(2026. 10. 15.)까지는 배제되어 6%
    assert settle and settle[0].segments[0].start == date(2025, 7, 15)
    assert [s.rate for s in settle[0].segments] == [Decimal(6), Decimal(20)]
    assert "청구 최대(배제 없음)" in it.alternatives


def test_계산표_시트(result, tmp_path):
    out = write_labor_workbook(result, tmp_path / "노동.xlsx")
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames[:2] == ["요약", "통상임금"]
    for name in ("시간외수당", "연차휴가수당", "평균임금", "퇴직금", "지연손해금", "소멸시효", "경고", "근거"):
        assert name in wb.sheetnames
    texts = [c.value for row in wb["요약"].iter_rows() for c in row if isinstance(c.value, str)]
    assert any("다 갚는 날까지 연 20%" in t for t in texts)


def test_cli_와_감시폴더(tmp_path, capsys):
    case = tmp_path / "사건.yaml"
    shutil.copy(SAMPLE, case)
    assert cli.is_labor_yaml(case)
    out = cli.run_labor(case)
    assert out.exists() and "노동금액계산표" in out.name
    root = watch.ensure_root(tmp_path / "감시")
    shutil.copy(SAMPLE, root / "01_입력" / "노동.yaml")
    assert watch.scan_once(root) == 1
    assert list((root / "02_결과").glob("*노동금액계산표*.xlsx"))


def test_해고기간_임금_샘플():
    res = calculate_labor(load_labor_case(ROOT / "cases" / "labor_dismissal_sample.yaml"))
    part = res.parts["dismissal"]
    # 월 4,000,000원(기본급 3,000,000 + 상여 400% ÷ 12) × 12개월
    # − 6개월 × min(중간수입 2,000,000, 4,000,000 − floor(4,000,000 × 0.7))
    assert part.total == Decimal(4_000_000) * 12 - 6 * min(Decimal(2_000_000), Decimal(4_000_000) - Decimal(2_800_000))
    it = res.parts["interest"]
    # 복직(구법 도래분): 20% 없음 — 지급일 다음 날부터 6%, 선고 다음 날부터 소송촉진법 12%
    assert {tuple(s.rate for s in sch.segments) for sch in it.schedules} == {(Decimal(6), Decimal(12))}


def test_앞_절이_없으면_무엇이_필요한지_알린다():
    import yaml

    raw = yaml.safe_load(SAMPLE.read_text(encoding="utf-8"))
    del raw["average_wage"]
    with pytest.raises(LaborError, match="average_wage"):
        calculate_labor(load_labor_case(raw))
    raw = yaml.safe_load(SAMPLE.read_text(encoding="utf-8"))
    raw["options"]["없는옵션"] = 1
    with pytest.raises(LaborError, match="알 수 없는 옵션"):
        load_labor_case(raw)


# ---------------------------------------------------------------- 수당별 병행 여부(OW-03a)
def _night_case(claims, period="2024-10"):
    """정기상여금만 병행 항목. 한 달에 야간 10시간·휴일 8시간만 있는 시간외수당 사건."""
    raw = _raw()
    for it in raw["ordinary"]["items"]:
        if it["name"] == "정기상여금":
            it["parallel"] = True
    raw["ordinary"]["parallel_claims"] = claims
    raw["overtime"] = {"size_band": "5-29", "months": [
        {"period": period, "night_hours": 10, "holiday_le8_hours": 8, "paid": {"night": 0, "holiday": 0}}]}
    for k in ("leave", "average_wage", "retirement", "interest"):
        raw.pop(k)
    return raw


def test_야간_휴일_병행_여부가_연장과_다르면_멈춘다():
    # 연장수당만 병행, 야간·휴일수당은 나중에 추가(병행 아님) — 시간외수당 모듈은 수당별 시급을 받지 못한다
    raw = _night_case([{"allowance": "overtime", "parallel": True},
                       {"allowance": "night", "parallel": False},
                       {"allowance": "holiday_work", "parallel": False}])
    with pytest.raises(LaborError, match="수당별 시급 적용은 아직 지원하지 않습니다") as e:
        calculate_labor(load_labor_case(raw))
    assert "night" in str(e.value) and "holiday_work" in str(e.value)
    # 반대 조합: 연장은 병행 아님, 야간은 병행
    raw = _night_case([{"allowance": "overtime", "parallel": False}, {"allowance": "night", "parallel": True}])
    with pytest.raises(LaborError, match="night"):
        calculate_labor(load_labor_case(raw))


def test_야간_휴일_병행_여부가_같거나_경계_뒤면_계산한다():
    raw = _night_case([{"allowance": a, "parallel": False} for a in ("overtime", "night", "holiday_work")])
    res = calculate_labor(load_labor_case(raw))
    lines = res.parts["overtime"].rows[0].lines
    assert {ln.item for ln in lines} == {"night", "holiday"}
    night = res.parts["ordinary"].hourly_of(date(2024, 10, 1), allowance="night")
    assert all(ln.hourly == night for ln in lines)
    # 2024. 12. 19. 이후 제공분은 수당과 관계없이 신 법리 — 병행 여부가 달라도 시급이 같다
    raw = _night_case([{"allowance": "overtime", "parallel": True}, {"allowance": "night", "parallel": False},
                       {"allowance": "holiday_work", "parallel": False}], period="2025-01")
    assert calculate_labor(load_labor_case(raw)).parts["overtime"].total > 0


def test_청구묶음을_적으면_방법을_알리고_멈춘다():
    raw = _raw()
    raw["ordinary"]["parallel_claims"] = [
        {"allowance": "overtime", "bundle": "최초", "served": "2025-09-15", "parallel": True}]
    with pytest.raises(LaborError, match="bundle 을 비운 한 줄"):
        calculate_labor(load_labor_case(raw))
    # 조립 경로가 쓰지 않는 수당의 청구묶음은 계산을 막지 않는다
    raw["ordinary"]["parallel_claims"] = [{"allowance": "public_holiday", "bundle": "2025확장", "parallel": False}]
    assert calculate_labor(load_labor_case(raw)).parts["overtime"].total > 0


# ---------------------------------------------------------------- 값 전달
def test_통상임금_모듈이_나중에_낸_경고도_경고_시트에_실린다(tmp_path):
    res = calculate_labor(load_labor_case(SAMPLE))
    part = res.parts["ordinary"].warnings
    # 연차·평균임금 모듈이 daily_ordinary_of 를 부를 때 통상임금 모듈이 더한 OW-M5 경고
    assert any(w.startswith("OW-M5") and "연차휴가수당" in w for w in part)
    sheet = [c.value for c in _sheet(res, tmp_path, "경고")["A"]]
    for w in part:
        assert f"[통상임금] {w}" in res.warnings
        assert f"[통상임금] {w}" in sheet


def test_연차_절의_5인_미만_표시가_조립_경로에서도_산다():
    # worker.small_business_periods 가 비어 있으면 deps 를 넘기지 않아 leave.small_business 가 쓰인다
    raw = _raw()
    raw["leave"]["small_business"] = True
    res = calculate_labor(load_labor_case(raw))
    assert res.parts["leave"].total == 0
    assert "연차휴가수당" not in {c.category for c in res.claims}


# ---------------------------------------------------------------- 요약·계산표
def test_미복직_장래분이_요약과_해고기간_임금_시트에_나온다(tmp_path):
    raw = _raw(DISMISSAL)
    d = raw["dismissal"]
    del d["reinstatement_date"]
    d["closing_date"] = "2025-05-20"
    d["past_until"] = "2025-04-30"
    res = calculate_labor(load_labor_case(raw))
    part = res.parts["dismissal"]
    assert part.future_monthly_amount == Decimal(4_000_000)
    assert any("2025. 5. 1.부터 복직시까지 월 4,000,000 원" in x for x in summary_lines(res))
    # 장래분은 원금 합계에 넣지 않는다
    assert res.principal_total == part.total
    ws = _sheet(res, tmp_path, "해고기간 임금")
    kinds = [ws.cell(r, 1).value for r in range(4, ws.max_row + 1)]
    assert kinds[-1] == "future" and kinds.count("future") == 1


def test_부당이득_반환과_금전보상은_요약_이름이_다르다():
    raw = _raw(DISMISSAL)
    raw["dismissal"]["mode"] = "refund"
    raw.pop("interest")
    lines = summary_lines(calculate_labor(load_labor_case(raw)))
    assert any(x.startswith("부당이득 반환액") and "원금 합계·지연손해금 제외" in x for x in lines)
    assert not any(x.startswith("해고기간 임금") for x in lines)
    raw = _raw(DISMISSAL)
    raw["dismissal"]["mode"] = "labor_commission_award"
    raw["dismissal"]["labor_commission"] = {"monthly_wage": 3000000, "decision_date": "2024-09-30", "service_days": 30}
    raw.pop("interest")
    lines = summary_lines(calculate_labor(load_labor_case(raw)))
    assert any(x.startswith("금전보상 임금상당액") and "원금 합계·지연손해금 제외" in x for x in lines)
    assert not any(x.startswith("해고기간 임금") for x in lines)


def test_시간외수당_충당_방식별_대안이_요약에_나온다():
    # 10월에 초과 지급 — 충당 방식(OT-18)에 따라 11월 청구액이 달라진다(사건은 contractual/month 를 골랐다)
    raw = _raw()
    raw["overtime"] = {"size_band": "5-29", "months": [
        {"period": "2024-10", "overtime_hours": 10, "paid": {"overtime": 1000000}},
        {"period": "2024-11", "overtime_hours": 20, "paid": {"overtime": 0}}]}
    for k in ("leave", "average_wage", "retirement", "interest"):
        raw.pop(k)
    res = calculate_labor(load_labor_case(raw))
    part = res.parts["overtime"]
    others = {m: v for m, v in part.alternatives.items() if v != part.alternatives[part.method]}
    assert part.method == "contractual/month" and others
    text = "\n".join(summary_lines(res))
    for m, v in others.items():
        assert f"{m} {int(v):,} 원" in text


def test_근거_시트에는_계산한_절의_옵션만_싣는다(tmp_path):
    ws = _sheet(calculate_labor(load_labor_case(DISMISSAL)), tmp_path, "근거")
    keys = []
    for r in range(4, ws.max_row + 1):
        if not ws.cell(r, 1).value:
            break
        keys.append(ws.cell(r, 1).value)
    assert keys and {k.split("_")[0] for k in keys} == {"dw", "ii", "di"}


def test_근거_시트_규칙_상태는_옵션의_모듈에서_찾는다(result, tmp_path, monkeypatch):
    # 규칙 ID 는 모듈마다 붙인다(RS-04 는 평균임금·퇴직금이 따로 쓴다). 한쪽 상태가 바뀌어도 다른 쪽 옵션 행이 따라 바뀌면 안 된다
    monkeypatch.setitem(average.RULES, average.OPTIONS["aw_round_avg_daily"].rule, ("1일 평균임금 끝수", "불명확"))
    monkeypatch.setitem(retirement.RULES, retirement.OPTIONS["rs_round_severance"].rule, ("지급률·퇴직금 끝수", "판례확립"))
    ws = _sheet(result, tmp_path, "근거")
    fills = {}
    for r in range(4, ws.max_row + 1):
        if not ws.cell(r, 1).value:
            break
        fills[ws.cell(r, 1).value] = ws.cell(r, 1).fill.fgColor.rgb
    assert fills["aw_round_avg_daily"] == WARN_FILL.fgColor.rgb
    assert fills["rs_round_severance"] != WARN_FILL.fgColor.rgb


def test_해고기간_임금_시트는_한글_머리글과_금액_서식(tmp_path):
    ws = _sheet(calculate_labor(load_labor_case(DISMISSAL)), tmp_path, "해고기간 임금")
    heads = [ws.cell(3, j).value for j in range(1, ws.max_column + 1)]
    english = {"period_days", "covered_days", "proration", "wage", "deduction_base", "small_business", "allowance",
               "limit", "income", "income_uncapped", "deduction", "pay", "pay_amount", "pay_date"}
    assert not english & set(heads)
    col = {h: j for j, h in enumerate(heads, 1)}
    assert ws.cell(4, col["임금"]).number_format == "#,##0"
    assert ws.cell(4, col["청구 금액"]).number_format == "#,##0"


def test_끝수처리_전_중간값은_소수까지_보인다(result, tmp_path):
    ws = _sheet(result, tmp_path, "연차휴가수당")
    heads = {ws.cell(3, j).value: j for j in range(1, ws.max_column + 1)}
    cells = [ws.cell(r, heads["1일 통상임금"]) for r in range(4, ws.max_row + 1)]
    cells = [c for c in cells if c.value]
    assert cells and all(c.number_format == "#,##0.00##" for c in cells)


# ---------------------------------------------------------------- 옵션 입력 형식
def test_YAML_이_바꿔_읽은_선택지는_되돌린다():
    raw = _raw()
    raw["options"].update(yaml.safe_load("ot_part_time_premium: off\now_weeks_per_year: 52.14\n"))
    case = load_labor_case(raw)
    assert case.options["ot_part_time_premium"] == "off"
    assert case.options["ow_weeks_per_year"] == "52.14"


@pytest.mark.parametrize("text, match", [
    ('di_ended_old_law_20: "false"', "true 또는 false"),     # 문자열 'false' 가 조용히 참이 되지 않게
    ("di_ended_old_law_20:", "비어 있습니다"),                # 빈 값이 조용히 거짓이 되지 않게
    ('di_new_law_cutoff: "2025-13-45"', "날짜"),
    ('di_claimed_rate_cap: "12%"', "% 없이"),
    ("ot_part_time_premium: yes", "따옴표"),
])
def test_옵션_값_형식이_맞지_않으면_알린다(text, match):
    raw = _raw()
    raw["options"].update(yaml.safe_load(text))
    with pytest.raises(LaborError, match=match):
        load_labor_case(raw)


def test_숫자가_아닌_값은_LaborError():
    with pytest.raises(LaborError, match="숫자로 읽을 수 없는"):
        dec("12%")


def test_청구_이율_상한을_소수로_적으면_경고한다():
    raw = _raw(DISMISSAL)
    raw["options"]["di_claimed_rate_cap"] = 0.12
    res = calculate_labor(load_labor_case(raw))
    assert any("di_claimed_rate_cap 은 연 % 단위" in w for w in res.warnings)
    raw["options"]["di_claimed_rate_cap"] = 12
    res = calculate_labor(load_labor_case(raw))
    assert not any("di_claimed_rate_cap" in w for w in res.warnings)
