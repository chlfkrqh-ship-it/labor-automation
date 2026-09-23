"""대법원 프로그램 엑셀표저장 결과와의 대조 (골든 테스트).

프로그램이 내놓은 값을 정답으로 두고 엔진이 같은 숫자를 내는지 확인한다.
"""

import sys
from datetime import date
from decimal import Decimal
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from engine.cost import RecurringCost
from engine.hoffman import cumulative
from engine.schedule import build_income_rows, total_income
from engine.wage_period import occupation_key
from tests.golden.case_hoffman_cap import (
    ACCIDENT,
    CURE_END,
    INCOME_ROWS,
    INCOME_TOTAL,
    LOSS_RATE,
    TREATMENT_ROWS,
    TREATMENT_TOTAL,
    WAGES,
    WORK_END,
)


def _rows():
    return build_income_rows(
        accident_date=ACCIDENT,
        start=ACCIDENT,
        end=WORK_END,
        wage_of=lambda p: WAGES.get((p.year, p.index), 172068),
        loss_rate_of=lambda p: Decimal(100) if p.end <= CURE_END else LOSS_RATE,
        has_wage=lambda y, i: (y, i) in WAGES,
        boundaries=[CURE_END.replace(day=CURE_END.day + 1)],
    )


def test_순번_구성이_같다():
    got = _rows()
    assert len(got) == len(INCOME_ROWS)
    for r, g in zip(got, INCOME_ROWS):
        assert (r.start, r.end) == (g[1], g[2]), f"순번{g[0]} 기간"
        assert r.wage == g[3], f"순번{g[0]} 노임단가"
        assert r.days == g[4], f"순번{g[0]} 일수"
        assert r.salary == g[5], f"순번{g[0]} 월소득"
        assert r.loss_rate == g[6], f"순번{g[0]} 상실률"


def test_호프만_수치가_같다():
    got = _rows()
    for r, g in zip(got, INCOME_ROWS):
        step, m1, h1, m2, h2, adj = g[0], g[7], g[8], g[9], g[10], g[11]
        assert r.m1 == m1, f"순번{step} M1"
        assert r.m2 == m2, f"순번{step} M2"
        assert cumulative(r.m1) == Decimal(h1), f"순번{step} 호프만1"
        assert cumulative(r.m2) == Decimal(h2), f"순번{step} 호프만2"
        assert r.factor == Decimal(adj), f"순번{step} 적용호프만"


def test_240_상한이_순번3에_걸린다():
    got = _rows()
    # 호프만1 254.5131 - 호프만2 6.8857 = 247.6274 인데, 240 - 6.8857 = 233.1143 로 잘린다
    assert got[2].raw_factor == Decimal("247.6274")
    assert got[2].factor == Decimal("233.1143")
    assert got[2].capped
    assert sum(r.factor for r in got) == Decimal("240")


def test_기간일실수입과_안분액():
    got = _rows()
    for r, g in zip(got, INCOME_ROWS):
        assert r.amount == g[12], f"순번{g[0]} 기간일실수입"
        assert r.prorated_days == g[13], f"순번{g[0]} 안분일수"
        assert r.prorated_amount == g[14], f"순번{g[0]} 안분액"


def test_일실수입_전체합계():
    assert total_income(_rows()) == INCOME_TOTAL


def test_노임_구간판정():
    # 2025.09.01. 부터 다음 연도 상반기 단가를 쓴다는 규칙의 실증
    assert occupation_key(ACCIDENT) == (2025, 2)
    assert WAGES[occupation_key(ACCIDENT)] == 171037
    from datetime import date

    assert occupation_key(date(2025, 9, 1)) == (2026, 1)
    assert WAGES[occupation_key(date(2025, 9, 1))] == 172068


def test_향후치료비():
    for name, _kind, cost, first, last, life, prior, factor, total in TREATMENT_ROWS:
        c = RecurringCost(ACCIDENT, first, last, Decimal(cost), life, Decimal(prior))
        assert c.factor_sum() == Decimal(factor), f"{name} 수치합계"
        assert c.total() == Decimal(total), f"{name} 비용총액"
    assert TREATMENT_TOTAL == sum(Decimal(r[8]) for r in TREATMENT_ROWS)


# ================================================================ 골든 케이스 2 (사망)
from engine.dateutil import age_at, diff_ymd, months_from
from engine.fraction import fraction_to_decimal, remove_fraction
from engine.hoffman import MONTHLY_CAP
from engine.income import period_income
from engine.severance import method_b, severance_amount
from tests.golden import case_death as D


def test_사망_사고시연령():
    # 연령은 GetAgeString 을 쓰므로 초일산입이 없다
    y = age_at(D.BIRTH, D.ACCIDENT)
    assert (y.year, y.month, y.day) == D.AGE_AT_ACCIDENT
    # 기간 계산(GetDiffDaysString)은 초일산입이라 하루 더 나온다
    assert diff_ymd(D.BIRTH, D.ACCIDENT).day == D.AGE_AT_ACCIDENT[2] + 1


def test_사망_호프만():
    m1 = months_from(D.ACCIDENT, D.WORK_END)
    assert m1 == D.M1
    assert cumulative(m1) == Decimal(D.HOFFMAN1)
    # 254.1673 > 240 이므로 240 으로 잘린다
    assert min(cumulative(m1), MONTHLY_CAP) == Decimal(D.ADJUST_HOFFMAN)


def test_사망_월소득과_생계비():
    assert D.WAGE * D.DAYS == D.SALARY
    assert occupation_key(D.ACCIDENT) == (2024, 2)
    got = period_income(
        D.ADJUST_HOFFMAN, D.SALARY, 100, fraction_to_decimal(D.LIVING_COST)
    )
    assert got == D.INCOME_TOTAL


def test_사망_합계():
    # 구 프로그램 출력은 장례비 5,000,000원을 [기타손해]에 적고도 재산상 손해를 일실수입과 같게 두었다
    assert D.PROPERTY_DAMAGE == D.INCOME_TOTAL
    assert D.PROPERTY_DAMAGE + D.SOLATIUM == D.GRAND_TOTAL


def _death_result():
    from engine.calculate import calculate
    from engine.case import Case

    # 구 프로그램은 사고일~가동종료일을 한 순번으로 잡는다. 다음 반기 단가가 없으면 엔진도 한 순번이다.
    case = Case(name="가상", injury_type="사망", birth=D.BIRTH, accident=D.ACCIDENT,
                wages={"2024-2": D.WAGE}, funeral_cost=D.FUNERAL_COST, solatium=D.SOLATIUM)
    return calculate(case, lambda p: D.WAGE, lambda y, i: (y, i) == occupation_key(D.ACCIDENT))


def test_사망_사건을_끝까지_계산하면_프로그램_합계와_같다():
    r = _death_result()
    assert r.case.resolved_work_end() == D.WORK_END
    assert [(x.m1, x.m2, x.factor) for x in r.income_rows] == [(D.M1, D.M2, Decimal(D.ADJUST_HOFFMAN))]
    assert r.income_total == D.INCOME_TOTAL
    assert r.active_total == 0                          # 장례비는 적극손해에 넣지 않는다
    assert r.property_damage == D.PROPERTY_DAMAGE
    assert r.settlement["합계"] == D.GRAND_TOTAL
    assert any("장례비 5,000,000원" in w for w in r.warnings)


def test_사망_계산표의_적극손해_항목과_합계가_맞물린다(tmp_path):
    import openpyxl

    from engine.excel import write_workbook

    r = _death_result()
    wb = openpyxl.load_workbook(write_workbook(r, tmp_path / "사망.xlsx"))
    assert wb.sheetnames[-1] == "경고"
    ws = wb["종합"]

    def first(text):
        return next(row for row in range(1, ws.max_row + 1) if ws.cell(row, 2).value == text)

    start, end = first("[적극손해]"), first("적극손해 합계 : ")
    assert sum(ws.cell(row, 5).value or 0 for row in range(start + 1, end)) == ws.cell(end, 14).value == 0
    # 장례비는 [적극손해] 밖의 별도 블록에 적는다
    assert first("장례비") > end and ws.cell(first("장례비"), 5).value == int(D.FUNERAL_COST)
    head = first("순번")                                  # [일실수입] 머리글
    assert ws.cell(head, 8).value == "생계비" and ws.cell(head + 1, 8).value == D.LIVING_COST
    row = next(rr for rr in range(1, ws.max_row + 1)
               if ws.cell(rr, 2).value == 1 and ws.cell(rr, 3).value == "가상")
    assert ws.cell(row, 4).value is None                  # 사망 사건은 위자료 자동계산을 비운다
    assert ws.cell(row, 8).value == int(D.PROPERTY_DAMAGE)
    assert ws.cell(row, 10).value == int(D.GRAND_TOTAL)


def test_일실퇴직금_산출방식B():
    _, month_factor, day_factor = D.SERVICE_AT_ACCIDENT
    assert severance_amount(D.MONTHLY_SALARY, month_factor) == D.SEVERANCE_AT_ACCIDENT_MONTH
    assert severance_amount(D.MONTHLY_SALARY, day_factor) == D.SEVERANCE_AT_ACCIDENT_DAY
    assert method_b(D.SEVERANCE_RETIREMENT_PV, D.SEVERANCE_AT_ACCIDENT_MONTH) == D.METHOD_B_MONTH
    assert method_b(D.SEVERANCE_RETIREMENT_PV, D.SEVERANCE_AT_ACCIDENT_DAY) == D.METHOD_B_DAY


def test_재직기간_월할_일할():
    r = diff_ymd(D.HIRE, D.ACCIDENT)
    assert (r.year, r.month, r.day) == (0, 1, 22)
    assert round(r.month_calc, 13) == round(float(D.SERVICE_AT_ACCIDENT[1]), 13)
    assert round(r.day_calc, 13) == round(float(D.SERVICE_AT_ACCIDENT[2]), 13)


# ================================================ 골든 케이스 3 (과실상계·공제, 신 프로그램)
from engine.disability import (
    Impairment,
    combined_with_prior,
    combined_without_prior,
    prior_contribution,
    truncate2,
)
from engine.settlement import (
    after_fault_offset,
    fault_share,
    paid_cure_deduction,
    settle,
    solatium_auto,
)
from tests.golden import case_fault_deduction as F


def _f_rows():
    return build_income_rows(
        accident_date=F.ACCIDENT,
        start=F.ACCIDENT,
        end=F.WORK_END,
        wage_of=lambda p: F.WAGES.get((p.year, p.index), 172068),
        loss_rate_of=lambda p: Decimal(100) if p.end <= F.CURE_END else F.COMBINED_RATE,
        has_wage=lambda y, i: (y, i) in F.WAGES,
        boundaries=[date(2024, 7, 1)],
    )


def test_중복장해율_프로그램출력과_일치():
    items = [Impairment(40, prior=20, dept="정형외과"), Impairment(20, dept="안과")]
    assert truncate2(combined_with_prior(items)) == F.COMBINED_RATE
    assert truncate2(prior_contribution(items)) == F.PRIOR_CONTRIBUTION


def test_장해표_행별_표시값이_프로그램출력과_일치():
    # [노동능력 상실률] 표는 행마다 그 행까지의 중복장해·단순중복장해·기왕증 기여도를 보인다.
    # float 로 계산하면 1행 기왕증 기여도가 19.999999999999982 -> 19.99 로 깎였다(프로그램 표시 20).
    items = []
    for dept, rate, prior, combined, simple, contribution in F.IMPAIRMENTS:
        items.append(Impairment(rate, prior=prior or 0, dept=dept))
        assert truncate2(combined_with_prior(items)) == Decimal(combined), f"{dept} 중복장해"
        assert truncate2(combined_without_prior(items)) == Decimal(simple), f"{dept} 단순중복장해"
        assert truncate2(prior_contribution(items)) == Decimal(contribution), f"{dept} 기왕증 기여도"


def test_과실_일실수입_전순번():
    got = _f_rows()
    assert len(got) == len(F.INCOME_ROWS)
    for r, g in zip(got, F.INCOME_ROWS):
        assert (r.start, r.end) == (g[1], g[2]), f"순번{g[0]} 기간"
        assert r.wage == g[3] and r.days == g[4] and r.salary == g[5]
        assert r.loss_rate == Decimal(g[6]), f"순번{g[0]} 상실률"
        assert r.m1 == g[7] and r.m2 == g[9], f"순번{g[0]} M"
        assert cumulative(r.m1) == Decimal(g[8]), f"순번{g[0]} 호프만1"
        assert cumulative(r.m2) == Decimal(g[10]), f"순번{g[0]} 호프만2"
        assert r.factor == Decimal(g[11]), f"순번{g[0]} 적용호프만"
        assert r.amount == g[12], f"순번{g[0]} 기간일실수입"
    assert total_income(got) == F.INCOME_TOTAL


def test_과실비율액():
    assert fault_share(F.PROPERTY_DAMAGE, F.FAULT_RATE) == F.FAULT_SHARE
    assert after_fault_offset(F.PROPERTY_DAMAGE, F.FAULT_RATE) == F.AFTER_FAULT_DISPLAY


def test_지급치료비_기왕증_과실분():
    got = paid_cure_deduction(F.PAID_CURE, F.PRIOR_CONTRIBUTION, F.FAULT_RATE)
    assert got == F.PAID_CURE_DEDUCTION


def test_최종합계는_과실비율액을_빼는_경로():
    # 표시값 경로로 하면 1원이 어긋난다
    assert F.AFTER_FAULT_DISPLAY - F.DEDUCTION_TOTAL == F.PROPERTY_FINAL - 1
    r = settle(
        F.PROPERTY_DAMAGE,
        F.FAULT_RATE,
        paid_cure=F.PAID_CURE,
        prior_ratio=F.PRIOR_CONTRIBUTION,
        advance=F.ADVANCE,
        solatium=F.SOLATIUM_APPLIED,
    )
    assert r["원고측_과실비율액"] == F.FAULT_SHARE
    assert r["공제액_합계"] == F.DEDUCTION_TOTAL
    assert r["재산상손해_합계"] == F.PROPERTY_FINAL
    assert r["합계"] == F.GRAND_TOTAL


def test_위자료_자동계산():
    assert solatium_auto(F.COMBINED_RATE, F.FAULT_RATE) == F.SOLATIUM_AUTO


def test_사고시연령_34세():
    y = age_at(F.BIRTH, F.ACCIDENT)
    assert (y.year, y.month, y.day) == F.AGE_AT_ACCIDENT
