#!/usr/bin/env python3
"""사건 폴더의 자료에서 텍스트를 뽑아 한 파일에 모은다.

    python extract.py 사건/삼양식품

PDF·docx·xlsx·csv·txt·md 를 읽어 `<폴더>/_추출.txt` 로 쓴다.
한글(.hwp)은 지원하지 않는다. PDF 나 docx 로 저장해서 넣어야 한다.

엔진이 만든 계산 결과물(계산 폴더 `계산/` 아래 파일, `*_노동금액계산표*.xlsx`)은 자료가
아니므로 읽지 않고 머리에 이름만 적는다. 다시 계산할 때 이전 계산값이 자료로 섞이지 않게 한다.
상대방·법원의 별지 계산표는 자료이므로 이름에 '계산표'가 있다는 것만으로 거르지 않는다.

추출 결과는 사건 자료이므로 저장소에 커밋하지 않는다(.gitignore 로 막혀 있다).
"""

from __future__ import annotations

import sys
from pathlib import Path

SUPPORTED = {".pdf", ".docx", ".xlsx", ".xlsm", ".csv", ".txt", ".md"}
SKIP_NAMES = {"_추출.txt", "사건.yaml"}
OUTPUT_DIRS = {"계산"}                 # 손배계산.md 0절의 계산 폴더
OUTPUT_MARK = "_노동금액계산표"         # cli.py·watch.py 의 노동 금액 계산표 이름


def is_engine_output(rel: Path) -> bool:
    """엔진이 만든 계산 결과물인지. rel 은 사건 폴더 기준 상대경로."""
    if any(part in OUTPUT_DIRS for part in rel.parts[:-1]):
        return True
    return rel.suffix.lower() in (".xlsx", ".xlsm") and OUTPUT_MARK in rel.stem


def from_pdf(path: Path) -> str:
    from pypdf import PdfReader

    out = []
    for i, page in enumerate(PdfReader(str(path)).pages, 1):
        t = (page.extract_text() or "").strip()
        if t:
            out.append(f"--- p.{i} ---\n{t}")
    return "\n".join(out) or "(텍스트를 뽑지 못했다. 스캔 PDF 일 수 있다)"


def from_docx(path: Path) -> str:
    import docx

    d = docx.Document(str(path))
    out = [p.text.strip() for p in d.paragraphs if p.text.strip()]
    for ti, t in enumerate(d.tables, 1):
        out.append(f"[표 {ti}]")
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                out.append(" | ".join(cells))
    return "\n".join(out)


def from_xlsx(path: Path) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"[시트 {ws.title}]")
        for row in ws.iter_rows():
            cells = [f"{c.coordinate}={c.value}" for c in row if c.value is not None]
            if cells:
                out.append(" | ".join(cells))
    return "\n".join(out)


def from_text(path: Path) -> str:
    for enc in ("utf-8", "cp949", "utf-8-sig"):
        try:
            return path.read_text(encoding=enc)
        except UnicodeDecodeError:
            continue
    return "(인코딩을 읽지 못했다)"


READERS = {".pdf": from_pdf, ".docx": from_docx, ".xlsx": from_xlsx,
           ".xlsm": from_xlsx, ".csv": from_text, ".txt": from_text, ".md": from_text}


def main(folder: str) -> None:
    root = Path(folder)
    if not root.is_dir():
        sys.exit(f"폴더가 없습니다: {root}")

    parts, skipped, outputs = [], [], []
    for path in sorted(root.rglob("*")):
        if not path.is_file() or path.name in SKIP_NAMES or path.name.startswith("~$"):
            continue
        if is_engine_output(path.relative_to(root)):
            outputs.append(path.relative_to(root))
            continue
        if path.suffix.lower() not in SUPPORTED:
            skipped.append(path.relative_to(root))
            continue
        try:
            body = READERS[path.suffix.lower()](path)
        except Exception as exc:
            body = f"(읽기 실패: {type(exc).__name__} {exc})"
        parts.append(f"\n{'='*70}\n■ {path.relative_to(root)}\n{'='*70}\n{body}")
        print(f"  읽음  {path.relative_to(root)}")

    out = root / "_추출.txt"
    header = f"사건 폴더: {root}\n파일 {len(parts)}건"
    if outputs:
        header += "\n\n[계산 결과물 — 자료가 아니므로 읽지 않음]\n"
        header += "\n".join(f"  {s}" for s in outputs)
        for s in outputs:
            print(f"  제외   {s}")
    if skipped:
        header += "\n\n[지원하지 않는 형식 — PDF 나 docx 로 저장해서 넣으세요]\n"
        header += "\n".join(f"  {s}" for s in skipped)
        for s in skipped:
            print(f"  건너뜀 {s}")
    out.write_text(header + "\n".join(parts), encoding="utf-8")
    print(f"\n-> {out}  ({out.stat().st_size:,} bytes)")
    if skipped:
        print("   .hwp 는 텍스트를 뽑지 못합니다. 한글에서 PDF 로 저장해 주세요.")


if __name__ == "__main__":
    # Windows 파이프(cp949)에서 파일 이름·안내 출력이 UnicodeEncodeError 로 멈추지 않게 한다.
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8")
    if len(sys.argv) > 1 and sys.argv[1] in ('-h', '--help'):
        print(__doc__); sys.exit(0)
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
