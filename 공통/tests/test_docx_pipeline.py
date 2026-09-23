"""서면 docx 파이프라인(docx_merge → docx_restyle → docx_footnotes → docx_normalize)을 확인한다.

프레임은 추적되는 4_서면작성/templates/표준양식.docx 를 임시 폴더에서 고쳐 만든다. 입증방법 항목은
실제 사건 프레임처럼 '(가)번호매기기_내용' 스타일로 바꾸고, 목차가 필요한 시험은 '다음' 아래에
'1.번호매기기' 대제목을 넣는다. 사건 자료나 OneDrive 폴더에는 아무것도 쓰지 않는다.
"""
import contextlib
import copy
import hashlib
import io
import os
import subprocess
import sys
import tempfile
import types
import unittest
from datetime import datetime
from pathlib import Path
from unittest import mock

import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.text.paragraph import Paragraph
from PIL import Image

ROOT = Path(__file__).parents[2]
SCRIPTS = ROOT / '4_서면작성' / 'scripts'
TEMPLATE = ROOT / '4_서면작성' / 'templates' / '표준양식.docx'
sys.path.insert(0, str(SCRIPTS))
import docx_footnotes  # noqa: E402
import docx_merge  # noqa: E402
import docx_normalize  # noqa: E402
import docx_restyle  # noqa: E402

TOC = ('이 사건의 경위', '징계사유의 존재', '결론')
CP949 = dict(os.environ, PYTHONIOENCODING='cp949')
EVIDENCE = '\n입 증 방 법\n\n    을나 제1호증  징계위원회 회의록\n    을나 제2호증  회의록\n\n---\n① [확인 필요 사항]\n- 없음\n'


def make_frame(path, toc=(), title_style=None):
    """표준양식.docx 로 서면_프레임.docx 를 만든다. title_style 을 주면 본문 첫 서면명 제목의 스타일을 바꾼다."""
    d = docx.Document(str(TEMPLATE))
    paras = d.paragraphs
    i = next(i for i, p in enumerate(paras) if p.text.replace(' ', '') == '다음')
    cur = paras[i + 1]._p
    blank = copy.deepcopy(cur)
    for title in toc:
        head = copy.deepcopy(blank); cur.addnext(head); cur = head
        par = Paragraph(head, paras[0]._parent)
        par.style = d.styles['1.번호매기기']
        par.add_run(title)
        gap = copy.deepcopy(blank); cur.addnext(gap); cur = gap
    for p in d.paragraphs:
        if p.text.startswith('을나 제23호증'):
            p.style = d.styles['(가)번호매기기_내용']
            ppr = p._p.find(qn('w:pPr'))
            for el in ppr.findall(qn('w:numPr')) + ppr.findall(qn('w:ind')):
                ppr.remove(el)
        if title_style and p.style.name == '본문_대제목' and p.text.strip():
            p.style = d.styles[title_style]
    d.save(str(path))


def png(path, width, height, color):
    Image.new('RGB', (width, height), color).save(str(path))


def sha(data):
    return hashlib.sha1(data).hexdigest()


def rows(path):
    """본문 문단: (스타일, 글자, keepNext, keepLines, 그림 여부)"""
    out = []
    for p in docx.Document(str(path)).paragraphs:
        ppr = p._p.find(qn('w:pPr'))
        out.append((p.style.name, p.text,
                    ppr is not None and ppr.find(qn('w:keepNext')) is not None,
                    ppr is not None and ppr.find(qn('w:keepLines')) is not None,
                    p._p.find('.//' + qn('w:drawing')) is not None))
    return out


def captions_and_pictures(path):
    """[(캡션 글자, 바로 뒤 그림 문단들의 그림 sha1 목록)]"""
    d = docx.Document(str(path))
    ps = d.paragraphs
    found = []
    for i, p in enumerate(ps):
        if p.style.name != '표의첫행제목':
            continue
        pics = []
        for q in ps[i + 1:]:
            blip = q._p.find('.//' + qn('a:blip'))
            if blip is None:
                break
            pics.append(sha(d.part.related_parts[blip.get(qn('r:embed'))].blob))
        found.append((p.text, pics))
    return found


def run(func, *args, **kwargs):
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        result = func(*args, **kwargs)
    return result, buf.getvalue()


class CaseDir(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.round = Path(self.tmp.name) / '라운드1'
        self.outdir = self.round / '출력'
        self.excerpts = self.round / '발췌'
        self.outdir.mkdir(parents=True)
        self.excerpts.mkdir()
        self.frame = self.outdir / '서면_프레임.docx'
        self.md = self.outdir / '서면초안.md'
        self.out = self.outdir / '(테스트) 준비서면_초안.docx'

    def tearDown(self):
        self.tmp.cleanup()

    def merge(self, draft, toc=(), encoding='utf-8', title_style=None):
        make_frame(self.frame, toc, title_style)
        self.md.write_text(draft, encoding=encoding)
        return run(docx_merge.merge, str(self.frame), str(self.md), str(self.out))

    def leftovers(self):
        return sorted(p.name for p in self.outdir.iterdir() if p.name.startswith('.docx_merge_'))

    def backups(self):
        return sorted(p for p in self.outdir.iterdir() if '.bak_merge_' in p.name)


class HeadingTests(CaseDir):
    """대제목이 조용히 본문으로 떨어지지 않는다(프레임 목차와 다른 제목, 목차 없는 프레임, BOM)."""

    DRAFT = ('1. 이 사건의 경위\n\n가. 징계의 경위\n\n피고는 원고에게 정직 3월의 징계를 하였습니다(을나 제1호증).\n\n'
             '2. 징계사유는 존재합니다\n\n가. 판단 기준\n\n원고는 3년 이상 근무하였습니다.\n\n'
             '3. 결론\n\n원고의 청구는 기각되어야 합니다.\n' + EVIDENCE)
    MATCHING = DRAFT.replace('징계사유는 존재합니다', '징계사유의 존재')   # 프레임 목차(TOC)와 같은 항목명

    def heads(self):
        return [text for style, text, *_ in rows(self.out) if style == '1.번호매기기']

    def test_heading_that_differs_from_frame_toc_stays_a_heading(self):
        code, log = self.merge(self.DRAFT, toc=TOC)
        self.assertEqual(code, 1, log)          # 프레임 목차와 어긋나면 알리고 종료 코드 1
        self.assertEqual(self.heads(), ['이 사건의 경위', '징계사유는 존재합니다', '결론'])
        self.assertIn('프레임 목차에 없는 대제목', log)
        self.assertIn('초안에 없어 빠진 프레임 목차: 징계사유의 존재', log)
        self.assertIn('계층 스타일 비정합 0건', log)
        self.assertNotIn('2. 징계사유는 존재합니다', [text for _, text, *_ in rows(self.out)])

    def test_matching_toc_is_clean(self):
        code, log = self.merge(self.MATCHING.replace('징계사유의 존재', '징계사유의  존재'), toc=TOC)
        self.assertEqual(code, 0, log)          # 공백 차이는 같은 항목으로 본다
        self.assertEqual(self.heads(), list(TOC))
        self.assertNotIn('경고', log)

    def test_frame_without_toc_gets_numbered_headings(self):
        code, log = self.merge(self.DRAFT, toc=())
        self.assertEqual(code, 0, log)
        self.assertEqual(self.heads(), ['이 사건의 경위', '징계사유는 존재합니다', '결론'])
        subs = [text for style, text, *_ in rows(self.out) if style == '가.번호매기기']
        self.assertEqual(subs, ['징계의 경위', '판단 기준'])
        self.assertIn('프레임에 목차가 없어', log)

    def test_draft_saved_with_bom(self):
        code, log = self.merge(self.MATCHING, toc=TOC, encoding='utf-8-sig')
        self.assertEqual(code, 0, log)
        self.assertEqual(self.heads()[0], '이 사건의 경위')
        self.assertFalse(any('\ufeff' in text for _, text, *_ in rows(self.out)))

    def test_out_of_sequence_number_is_reported(self):
        draft = self.DRAFT.replace('2. 징계사유는 존재합니다', '5. 징계사유는 존재합니다')
        code, log = self.merge(draft, toc=TOC)
        self.assertEqual(code, 1, log)          # 병합 뒤 점검이 '대제목 누락 의심'으로 잡는다
        self.assertIn('본문으로 둔 줄: 5. 징계사유는 존재합니다', log)
        self.assertIn('대제목 누락 의심 1건', log)

    def test_title_in_evidence_title_style_before_daum(self):
        """서면명 제목에 별지제목을 쓰는 프레임에서도 restyle 범위가 '다음' 뒤 입증방법까지다."""
        code, log = self.merge(self.MATCHING, toc=TOC, title_style='별지제목')
        self.assertEqual(code, 0, log)
        self.assertNotIn('대상 문단 0건', log)
        self.assertEqual([s for s, text, *_ in rows(self.out) if text == '피고는 원고에게 정직 3월의 징계를 하였습니다(을나 제1호증).'],
                         ['가.번호매기기_내용'])

    def test_heading_candidate_rules(self):
        self.assertEqual(docx_restyle.h1_candidate('2. 징계사유는 존재합니다'), (2, '징계사유는 존재합니다'))
        self.assertEqual(docx_restyle.h1_candidate('\ufeff1. 이 사건의 경위'), (1, '이 사건의 경위'))
        self.assertEqual(docx_restyle.h1_candidate('3. 파면과 해임은 관련이 없습니다.'), (3, '파면과 해임은 관련이 없습니다.'))
        self.assertIsNone(docx_restyle.h1_candidate('2025. 3. 1. 참가인은 원고를 해고하였습니다.'))
        self.assertIsNone(docx_restyle.h1_candidate('1. 피고는 원고에게 10,000,000원을 지급하라.'))
        self.assertIsNone(docx_restyle.h1_candidate('2. 소송비용은 피고가 부담한다.'))


class RestyleCheckTests(CaseDir):
    """--check 가 구조 붕괴를 0건으로 넘기지 않는다."""

    def merged(self):
        draft = ('1. 이 사건의 경위\n\n피고는 원고를 징계하였습니다.\n\n[을나 제1호증 회의록]\n\n'
                 '2. 결론\n\n원고의 청구는 기각되어야 합니다.\n' + EVIDENCE)
        png(self.excerpts / '을나1.png', 600, 300, 'white')
        code, log = self.merge(draft, toc=('이 사건의 경위', '결론'))
        self.assertEqual(code, 0, log)
        return docx.Document(str(self.out))

    def check(self, d):
        d.save(str(self.out))
        return run(docx_restyle.main, str(self.out), check=True)

    def test_clean_merge_passes(self):
        code, log = self.check(self.merged())
        self.assertEqual(code, 0, log)

    def test_heading_left_as_body_is_flagged(self):
        d = self.merged()
        head = next(p for p in d.paragraphs if p.text == '결론')
        head.style = d.styles['1.번호매기기_내용']
        head.runs[0].text = '2. 결론'
        code, log = self.check(d)
        self.assertEqual(code, 1)
        self.assertIn('대제목 누락 의심 1건', log)

    def test_no_heading_at_all_is_flagged(self):
        d = self.merged()
        for p in d.paragraphs:
            if p.style.name == '1.번호매기기':
                p.style = d.styles['Normal']
        code, log = self.check(d)
        self.assertEqual(code, 1)
        self.assertIn("대제목('1.번호매기기') 문단이 하나도 없습니다", log)

    def test_caption_without_keep_next_is_flagged(self):
        d = self.merged()
        cap = next(p for p in d.paragraphs if p.style.name == '표의첫행제목')
        cap.paragraph_format.keep_with_next = None
        code, log = self.check(d)
        self.assertEqual(code, 1)
        self.assertIn('캡션 붙임(keepNext) 없음 1건', log)


class ExcerptImageTests(CaseDir):
    """발췌 그림 이름: 가지번호('의')와 같은 호증의 두 번째 발췌('-2')가 겹치지 않는다."""

    DRAFT = ('1. 이 사건의 경위\n\n첫 문단입니다.\n\n[을나 제2호증 회의록]\n\n두 번째 문단입니다.\n\n'
             '[을나 제2호증 회의록]\n\n세 번째 문단입니다.\n\n[을나 제2호증의 2 회의록]\n' + EVIDENCE)

    def files(self, **names):
        colors = iter(['red', 'green', 'blue', 'yellow', 'gray'])
        made = {}
        for name, size in names.items():
            path = self.excerpts / (name + '.png')
            png(path, size[0], size[1], next(colors))
            made[name] = sha(path.read_bytes())
        return made

    def test_branch_number_and_second_excerpt_get_their_own_images(self):
        made = self.files(**{'을나2': (600, 300), '을나2-2': (600, 700), '을나2의2': (600, 200)})
        code, log = self.merge(self.DRAFT)
        self.assertEqual(code, 0, log)
        got = captions_and_pictures(self.out)
        self.assertEqual(got[0], ('[을나 제2호증 회의록]', [made['을나2']]))
        self.assertEqual(got[1], ('[을나 제2호증 회의록]', [made['을나2-2']] * 2))   # 세로 14cm 초과 → 두 장 분할
        self.assertEqual(got[2], ('[을나 제2호증의 2 회의록]', [made['을나2의2']]))

    def test_old_name_that_could_be_either_stops(self):
        self.files(**{'을나2': (600, 300), '을나2-2': (600, 300)})
        with self.assertRaises(SystemExit) as cm:
            self.merge(self.DRAFT)
        self.assertIn('을나2의2.png', str(cm.exception.code))
        self.assertFalse(self.out.exists())
        self.assertEqual(self.leftovers(), [])

    def test_old_branch_name_is_read_when_unambiguous(self):
        made = self.files(**{'을나2': (600, 300), '을나2-2': (600, 200)})
        draft = self.DRAFT.replace('[을나 제2호증 회의록]\n\n세 번째', '세 번째')
        code, log = self.merge(draft)
        self.assertEqual(code, 0, log)
        self.assertIn('옛 이름', log)
        self.assertEqual(captions_and_pictures(self.out)[-1], ('[을나 제2호증의 2 회의록]', [made['을나2-2']]))

    def test_explicit_image_counts_toward_sequence(self):
        made = self.files(**{'갑4': (600, 300), '갑4-2': (600, 200)})
        draft = ('1. 이 사건의 경위\n\n첫 문단입니다.\n\n[갑 제4호증 회의록]\n![](갑4.png)\n\n'
                 '두 번째 문단입니다.\n\n[갑 제4호증 회의록]\n' + EVIDENCE)
        code, log = self.merge(draft)
        self.assertEqual(code, 0, log)
        self.assertEqual([pics for _, pics in captions_and_pictures(self.out)], [[made['갑4']], [made['갑4-2']]])

    def test_same_image_under_two_captions_stops(self):
        self.files(**{'갑4': (600, 300)})
        draft = ('1. 이 사건의 경위\n\n[갑 제4호증 회의록]\n![](갑4.png)\n\n문단입니다.\n\n'
                 '[갑 제5호증 확인서]\n![](갑4.png)\n' + EVIDENCE)
        with self.assertRaises(SystemExit) as cm:
            self.merge(draft)
        self.assertIn('두 캡션', str(cm.exception.code))
        self.assertFalse(self.out.exists())

    def test_caption_and_pictures_keep_together(self):
        self.files(**{'을나2': (600, 300), '을나2-2': (600, 700), '을나2의2': (600, 200)})
        code, log = self.merge(self.DRAFT)
        self.assertEqual(code, 0, log)
        table = rows(self.out)
        caps = [r for r in table if r[0] == '표의첫행제목']
        pics = [r for r in table if r[4]]
        self.assertEqual(len(caps), 3)
        self.assertEqual(len(pics), 4)
        self.assertTrue(all(r[2] and r[3] for r in caps + pics), table)


class SafetyTests(CaseDir):
    """보관본을 덮어쓰지 않고, 실패하면 기존 출력을 건드리지 않는다."""

    DRAFT = ('1. 이 사건의 경위\n\n피고는 원고를 징계하였습니다[[각주: 대법원 판결 설시.]].\n\n'
             '2. 결론\n\n원고의 청구는 기각되어야 합니다.\n' + EVIDENCE)

    def test_backup_in_same_second_does_not_overwrite(self):
        fixed = types.SimpleNamespace(datetime=types.SimpleNamespace(now=lambda: datetime(2026, 9, 23, 0, 43, 5)))
        with mock.patch.object(docx_merge, 'datetime', fixed):
            self.assertEqual(self.merge(self.DRAFT)[0], 0)
            d = docx.Document(str(self.out))
            d.add_paragraph('HUMAN EDIT MARKER')
            d.save(str(self.out))
            self.assertEqual(self.merge(self.DRAFT)[0], 0)
            self.assertEqual(self.merge(self.DRAFT)[0], 0)
        backups = self.backups()
        self.assertEqual([p.name for p in backups],
                         ['(테스트) 준비서면_초안.bak_merge_20260923_004305.docx',
                          '(테스트) 준비서면_초안.bak_merge_20260923_004305_2.docx'])
        self.assertIn('HUMAN EDIT MARKER', [p.text for p in docx.Document(str(backups[0])).paragraphs])

    def test_draft_without_evidence_section_changes_nothing(self):
        self.assertEqual(self.merge(self.DRAFT)[0], 0)
        before = self.out.read_bytes()
        with self.assertRaises(SystemExit) as cm:
            self.merge(self.DRAFT.replace('입 증 방 법', ''))
        self.assertIn("'입 증 방 법' 줄이 없습니다", str(cm.exception.code))
        self.assertEqual(self.out.read_bytes(), before)
        self.assertEqual(self.backups(), [])
        self.assertEqual(self.leftovers(), [])

    def test_frame_without_evidence_title_stops(self):
        make_frame(self.frame)
        d = docx.Document(str(self.frame))
        for p in d.paragraphs:
            if p.style.name == '별지제목':
                p.style = d.styles['Normal']
        d.save(str(self.frame))
        self.md.write_text(self.DRAFT, encoding='utf-8')
        with self.assertRaises(SystemExit) as cm:
            run(docx_merge.merge, str(self.frame), str(self.md), str(self.out))
        self.assertIn('별지제목', str(cm.exception.code))
        self.assertFalse(self.out.exists())

    def test_failure_after_merge_keeps_previous_output(self):
        self.assertEqual(self.merge(self.DRAFT)[0], 0)
        before = self.out.read_bytes()
        with mock.patch.object(docx_merge.docx_footnotes, 'convert', side_effect=RuntimeError('boom')):
            with self.assertRaises(RuntimeError):
                self.merge(self.DRAFT)
        self.assertEqual(self.out.read_bytes(), before)
        self.assertEqual(self.backups(), [])
        self.assertEqual(self.leftovers(), [])

    def test_clean_merge_reports_three_checks(self):
        code, log = self.merge(self.DRAFT)
        self.assertEqual(code, 0, log)
        self.assertIn('계층 스타일 비정합 0건', log)
        self.assertIn('남은 [[각주:]] 표시 0건 / 현재 각주 1개', log)
        self.assertIn('원문자 글꼴 비정합 0건', log)

    def test_folder_form_still_writes_old_name(self):
        make_frame(self.frame)
        self.md.write_text(self.DRAFT, encoding='utf-8')
        # 한국어 Windows 의 파이프 출력(cp949)에서도 UTF-8 로 내보내는지 함께 본다
        result = subprocess.run([sys.executable, '-B', str(SCRIPTS / 'docx_merge.py'), str(self.outdir)],
                                capture_output=True, text=True, encoding='utf-8', env=CP949)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertTrue((self.outdir / '최종서면.docx').exists())
        self.assertIn('호환용', result.stdout)


class DraftBoundaryTests(CaseDir):
    """본문과 하단 목록의 경계('---' 대시 3개 이상)와 병합 전 점검."""

    BODY = '1. 이 사건의 경위\n\n피고는 원고를 징계하였습니다.\n\n2. 결론\n\n원고의 청구는 기각되어야 합니다.\n'
    BOTTOM = ('① [확인 필요 사항]\n- 없음\n② ⚠️ [취약 논리 및 보완 방향]\n- 징계양정 비교 사례 부족\n'
              '※ 본 서면 초안은 AI 보조 작성물로, 최종 검토 및 법적 판단은 담당 변호사·노무사의 확인이 필요합니다.\n')
    EV = '\n입 증 방 법\n\n    을나 제1호증  징계위원회 회의록\n'

    def texts(self):
        return [text for _, text, *_ in rows(self.out)]

    def test_long_rule_line_ends_body(self):
        code, log = self.merge(self.BODY + self.EV + '\n-----\n' + self.BOTTOM)
        self.assertEqual(code, 0, log)
        joined = '\n'.join(self.texts())
        self.assertNotIn('확인 필요 사항', joined)
        self.assertNotIn('AI 보조 작성물', joined)
        self.assertNotIn('-----', joined)

    def test_bottom_lists_without_rule_line_stop(self):
        with self.assertRaises(SystemExit) as cm:
            self.merge(self.BODY + self.EV + '\n' + self.BOTTOM)
        self.assertIn("'---' 한 줄", str(cm.exception.code))
        self.assertFalse(self.out.exists())

    def test_rule_line_matches_style_check(self):
        import style_check
        rule = getattr(style_check, 'RULE_LINE', None)
        if rule is None:
            self.skipTest('style_check.RULE_LINE 없음')
        for line in ('---', '-----', '  ---  ', '--', '- 목록', '|---|---|'):
            self.assertEqual(bool(docx_merge.RULE_LINE.match(line)), bool(rule.match(line)), line)

    def test_markdown_table_stops(self):
        table = '\n| 연도 | 처분 |\n|---|---|\n| 2021 | 정직 3월 |\n'
        with self.assertRaises(SystemExit) as cm:
            self.merge(self.BODY + table + self.EV)
        self.assertIn('md 표', str(cm.exception.code))
        self.assertFalse(self.out.exists())

    def test_unbalanced_footnote_marker_stops(self):
        body = self.BODY.replace('징계하였습니다.', '징계하였습니다[[각주: 근속기간은 갑 제1호증 참조).].')
        with self.assertRaises(SystemExit) as cm:
            self.merge(body + self.EV)
        self.assertIn("'[[각주:' 와 ']]' 짝", str(cm.exception.code))
        self.assertFalse(self.out.exists())


def visible(p):
    """문단 글자(탭 '\\t', 줄바꿈 '\\n' 포함). 스크립트의 함수를 쓰지 않고 따로 센다."""
    out = []
    for el in p._p.iter(qn('w:t'), qn('w:tab'), qn('w:br')):
        if el.getparent().tag != qn('w:r'):
            continue
        if el.tag == qn('w:t'):
            out.append(el.text or '')
        else:
            out.append('\t' if el.tag == qn('w:tab') else '\n')
    return ''.join(out)


class FootnoteTests(unittest.TestCase):
    """각주 변환은 표시 글자만 떼어 내고 탭·줄바꿈·변경 추적 문구를 제자리에 둔다."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = Path(self.tmp.name) / '수정안.docx'

    def tearDown(self):
        self.tmp.cleanup()

    def notes(self):
        return run(docx_footnotes.convert, str(self.path), list_only=True)[1]

    def test_tab_break_and_tracked_insert_are_kept(self):
        d = docx.Document()
        p = d.add_paragraph()
        p.add_run('가. 원고는\t피고 회사의 근로자입니다')
        p.add_run('(갑 제1호증).')
        p.add_run().add_break()
        p.add_run('대법원 2019도10516 판결[[각주: 판결 원문 설시.]]은 이를 밝혔습니다.')
        q = d.add_paragraph()
        q._p.append(parse_xml('<w:r %s><w:tab/><w:t>들여쓴 문장[[각주: 각주 A.]]입니다.</w:t></w:r>' % nsdecls('w')))
        q._p.append(parse_xml('<w:ins %s w:id="1" w:author="변호사" w:date="2026-09-20T00:00:00Z">'
                              '<w:r><w:t>(변호사 추가 문장)</w:t></w:r></w:ins>' % nsdecls('w')))
        q._p.append(parse_xml('<w:r %s><w:t xml:space="preserve"> 끝.</w:t></w:r>' % nsdecls('w')))
        d.save(str(self.path))
        code, log = run(docx_footnotes.convert, str(self.path))
        self.assertEqual(code, 0, log)
        a, b = docx.Document(str(self.path)).paragraphs[:2]
        self.assertEqual(visible(a), '가. 원고는\t피고 회사의 근로자입니다(갑 제1호증).\n대법원 2019도10516 판결은 이를 밝혔습니다.')
        self.assertEqual(visible(b), '\t들여쓴 문장입니다.(변호사 추가 문장) 끝.')
        ins = b._p.find(qn('w:ins'))
        self.assertEqual(''.join(t.text for t in ins.iter(qn('w:t'))), '(변호사 추가 문장)')
        # 각주번호는 표시가 있던 자리(어구 바로 뒤)에 들어간다
        kids = [k for k in a._p if k.tag == qn('w:r')]
        ref = next(i for i, k in enumerate(kids) if k.find(qn('w:footnoteReference')) is not None)
        self.assertTrue(''.join(t.text for t in kids[ref - 1].iter(qn('w:t'))).endswith('판결'))
        self.assertIn('[1] 판결 원문 설시.', self.notes())
        self.assertIn('[2] 각주 A.', self.notes())

    def test_marker_split_across_runs(self):
        d = docx.Document()
        p = d.add_paragraph()
        p.add_run('판결[[각')
        p.add_run('주: 설시').bold = True
        p.add_run(' 내용.]]은 ')
        p.add_run('그러합니다.')
        d.save(str(self.path))
        self.assertEqual(run(docx_footnotes.convert, str(self.path))[0], 0)
        self.assertEqual(visible(docx.Document(str(self.path)).paragraphs[0]), '판결은 그러합니다.')
        self.assertIn('[1] 설시 내용.', self.notes())

    def test_broken_marker_is_counted_and_not_swallowed(self):
        d = docx.Document()
        d.add_paragraph('원고는 3년 이상 근무하였습니다[[각주: 근속기간은 갑 제1호증 참조).]. '
                        '또한 피고는[[각주: 정상 표시.]] 다투지 않았습니다.')
        d.save(str(self.path))
        code, log = run(docx_footnotes.convert, str(self.path), check_only=True)
        self.assertEqual(code, 1)
        self.assertIn('남은 [[각주:]] 표시 2건', log)
        self.assertIn("맞지 않는 표시 1건", log)
        code, log = run(docx_footnotes.convert, str(self.path))
        self.assertEqual(code, 1, log)          # 잘못 닫은 표시가 남았다고 알린다
        text = visible(docx.Document(str(self.path)).paragraphs[0])
        self.assertEqual(text, '원고는 3년 이상 근무하였습니다[[각주: 근속기간은 갑 제1호증 참조).]. 또한 피고는 다투지 않았습니다.')
        self.assertIn('[1] 정상 표시.', self.notes())
        code, log = run(docx_footnotes.convert, str(self.path), check_only=True)
        self.assertEqual(code, 1)
        self.assertIn('남은 [[각주:]] 표시 1건', log)


class NormalizeTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = Path(self.tmp.name) / '수정안.docx'

    def tearDown(self):
        self.tmp.cleanup()

    def test_tabs_and_breaks_are_not_copied(self):
        d = docx.Document()
        d.add_paragraph().add_run('① 원고는\t피고의 근로자이고, ② 피고는\n원고를 해고하였습니다.')
        d.save(str(self.path))
        font, bad = docx_normalize.normalize(str(self.path))
        self.assertEqual(bad, 1)
        p = docx.Document(str(self.path)).paragraphs[0]
        self.assertEqual(visible(p), '① 원고는\t피고의 근로자이고, ② 피고는\n원고를 해고하였습니다.')
        self.assertEqual(len(p._p.findall('.//' + qn('w:tab'))), 1)
        self.assertEqual(len(p._p.findall('.//' + qn('w:br'))), 1)
        self.assertEqual(docx_normalize.normalize(str(self.path), check_only=True), (font, 0))

    def test_words_option_points_to_style_check(self):
        d = docx.Document()
        d.add_paragraph('① 원고는 한자(漢字)를 병기하였습니다.')
        d.save(str(self.path))
        before = self.path.read_bytes()
        result = subprocess.run([sys.executable, '-B', str(SCRIPTS / 'docx_normalize.py'), str(self.path), '--words'],
                                capture_output=True, text=True, encoding='utf-8', env=CP949)
        self.assertEqual(result.returncode, 2)
        self.assertIn('금지 낱말 검사는 style_check.py', result.stdout)
        self.assertEqual(self.path.read_bytes(), before)


if __name__ == '__main__':
    unittest.main()
