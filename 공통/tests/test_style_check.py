"""제출처별 대리인 표기 검사(4_서면작성/CLAUDE.md 2-2절)를 확인한다.

노동위원회에 내는 서면에 '소송대리인'이 남아 있으면 위반으로 잡혀야 한다.
표지·서명 블록은 표 칸과 누름틀(내용 컨트롤) 안에 있어 python-docx 의
paragraph.text 로는 보이지 않으므로, 그 자리까지 읽는지도 함께 본다.
"""
import importlib.util
from pathlib import Path
import tempfile
import unittest

import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

spec = importlib.util.spec_from_file_location(
    'style_check', Path(__file__).parents[2] / '4_서면작성' / 'scripts' / 'style_check.py')
style_check = importlib.util.module_from_spec(spec)
spec.loader.exec_module(style_check)

RULE = '노동위 서면에 소송대리인'


class VenueAgentTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def make(self, name, body=(), cells=(), stamped=None):
        """body 는 본문 문단, cells 는 표 칸, stamped 는 누름틀 안에 넣을 글자."""
        path = self.root / name
        document = docx.Document()
        for line in body:
            document.add_paragraph(line)
        if stamped is not None:
            paragraph = document.add_paragraph()
            paragraph._p.append(parse_xml(
                '<w:sdt %s><w:sdtPr/><w:sdtContent><w:r><w:t>%s</w:t></w:r>'
                '</w:sdtContent></w:sdt>' % (nsdecls('w'), stamped)))
            paragraph.add_run(' 귀중')
        if cells:
            table = document.add_table(rows=len(cells), cols=1)
            for row, line in enumerate(cells):
                table.cell(row, 0).text = line
        document.save(path)
        return path

    def rules(self, path):
        return [row['rule'] for row in style_check.check(path)['violations']]

    def test_labor_venue_with_litigation_agent_is_violation(self):
        path = self.make('노동위.docx',
                         body=['위 사건에 관하여 피신청인의 소송대리인은 다음과 같이 답변서를 제출합니다.',
                               '전남지방노동위원회 귀중'],
                         cells=['소송대리인', '법무법인 평안'])
        result = style_check.check(path)
        self.assertIn(RULE, [row['rule'] for row in result['violations']])
        self.assertFalse(result['ok'])
        # 본문 문단과 표 칸을 각각 잡는다
        self.assertEqual(sum(1 for row in result['violations'] if row['rule'] == RULE), 2)

    def test_venue_inside_content_control_is_read(self):
        path = self.make('누름틀.docx',
                         body=['위 사건에 관하여 피신청인의 소송대리인은 다음과 같이 답변서를 제출합니다.'],
                         stamped='중앙노동위원회')
        self.assertIn(RULE, self.rules(path))

    def test_court_venue_keeps_litigation_agent(self):
        path = self.make('법원.docx',
                         body=['위 사건에 관하여 피고보조참가인의 소송대리인은 다음과 같이 변론을 준비합니다.',
                               '서울행정법원 제1부 귀중'],
                         cells=['피고보조참가인 소송대리인'])
        self.assertNotIn(RULE, self.rules(path))

    def test_labor_venue_with_plain_agent_is_clean(self):
        path = self.make('노동위_정상.docx',
                         body=['위 사건에 관하여 피신청인의 대리인은 다음과 같이 답변서를 제출합니다.',
                               '전남지방노동위원회 귀중'],
                         cells=['대리인', '담당노무사 최 락 보'])
        self.assertNotIn(RULE, self.rules(path))

    def test_cited_decision_is_not_a_venue(self):
        """본문에서 노동위원회 판정을 인용한 법원 서면은 제출처가 법원이다."""
        path = self.make('인용.docx',
                         body=['중앙노동위원회 2026부해109 판정의 취지는 다음과 같습니다.',
                               '위 사건에 관하여 원고의 소송대리인은 다음과 같이 변론을 준비합니다.',
                               '서울행정법원 제1부 귀중'])
        self.assertNotIn(RULE, self.rules(path))

    def test_markdown_draft_has_no_cover(self):
        path = self.root / '서면초안.md'
        path.write_text('피신청인은 다음과 같이 답변합니다.\n', encoding='utf-8')
        self.assertEqual(style_check.cover_texts(path), [])
        self.assertNotIn(RULE, self.rules(path))


class StyleRuleTests(unittest.TestCase):
    """SKILL.md '저장 전 자가 점검'의 규칙이 걸려야 할 문장과 걸리면 안 되는 문장."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def check(self, *paragraphs):
        path = self.root / '서면초안.md'
        path.write_text('\n\n'.join(paragraphs) + '\n', encoding='utf-8')
        return style_check.check(path)

    def found(self, result, rule, kind='violations'):
        return [row['detail'] for row in result[kind] if row['rule'] == rule]

    def test_captions_and_bracket_citations_are_not_memos(self):
        result = self.check('[을나 제5호증 사실확인서]',
                            '[을 제2호증 경력확인서]',
                            '[증 제11호증 이체확인증]',
                            '[갑 제7호증 이체확인증(2026. 7. 5.)]',
                            '[서울중앙지방법원 2020. 1. 1. 선고 2019가합12345 해고무효확인 판결 中]',
                            '피고는 보충교섭을 거쳤습니다[갑 제3호증 보충교섭협약서(2022. 12. 14.자) 제10조].')
        self.assertEqual(self.found(result, '대괄호 메모 잔존'), [])

    def test_memos_are_violations(self):
        for memo in ('[반박 보완 필요]', '[확인 필요]', '【판례 확인 필요】', '[갑 제3호증 원본 확인 필요]',
                     '[원고 입사일 확인]', '[금액: 계산 확인 전]', '[입사일 확인 증빙]', '[송금 확인증빙]'):
            with self.subTest(memo=memo):
                self.assertEqual(len(self.found(self.check(memo), '대괄호 메모 잔존')), 1)

    def test_footnote_holds_judgment_text_verbatim(self):
        result = self.check('피고는 징계사유가 있다고 주장합니다.[[각주: 대법원은 징계사유가 존재하는지를 확인할 필요가 있고, '
                            '살피건대 이는 사용자의 재량에 속한다고 판단하였다(대법원 1992. 8. 14. 선고 91다29811 판결 참조).]]')
        self.assertEqual(result['violations'], [])
        self.assertEqual(self.found(result, '해라체 종결', 'notes'), [])
        self.assertEqual(result['sentences'], 1)

    def test_household_terms_are_not_conjunctions(self):
        rule = '접속부사 뒤 쉼표 없음'
        for text in ('가사 사용인에게는 근로기준법이 적용되지 않습니다.',
                     '가사사용인에 대하여는 근로기준법이 적용되지 않습니다.',
                     '가사근로자법은 2022. 6. 16. 시행되었습니다.',
                     '가사(家事) 사용인에 대하여는 적용하지 아니한다고 정하고 있습니다.',
                     '가사, 원고의 주장이 사실이라 하더라도 결론은 같습니다.',
                     '그러나, 위 문언은 적용 범위를 정한 것에 불과합니다.',
                     '한편으로는 원고의 사정도 고려할 수 있습니다.'):
            with self.subTest(text=text):
                self.assertEqual(self.found(self.check(text), rule), [])
        for text, word in (('가사 원고의 주장이 사실이라 하더라도 결론은 같습니다.', '가사'),
                           ('가사 이를 받아들인다고 해도 결론은 같습니다.', '가사'),
                           ('그러나 위 문언은 적용 범위를 정한 것에 불과합니다.', '그러나'),
                           ('따라서 이 사건 해고는 위법합니다.', '따라서')):
            with self.subTest(text=text):
                self.assertEqual(self.found(self.check(text), rule), [word])

    def test_ordinals_are_banned_only_as_enumeration(self):
        rule = '금지 낱말'
        for text in ('원고는 둘째 자녀 출산 후 육아휴직을 신청하였습니다.',
                     '피고의 취업규칙은 매월 둘째 주와 넷째 주 토요일을 휴무일로 정하고 있습니다(을 제3호증).',
                     '원고는 2023. 5. 셋째 주에 연장근로를 하였다고 주장합니다.'):
            with self.subTest(text=text):
                self.assertEqual(self.found(self.check(text), rule), [])
        for text in ('둘째, 원고는 징계위원회에 출석하지 않았습니다.', '첫째로 원고는 징계위원회에 출석하지 않았습니다.'):
            with self.subTest(text=text):
                self.assertEqual(len(self.found(self.check(text), rule)), 1)

    def test_sentences_ending_with_citation_are_counted(self):
        for text in ('관련 등기가 존재하지 않습니다(갑 제1호증 등기사항전부증명서).',
                     '해고는 서면으로 통지하여야 합니다(대법원 1992. 8. 14. 선고 91다29811 판결).',
                     '피고는 보충교섭을 거쳤습니다[갑 제3호증 보충교섭협약서(2022. 12. 14.자) 제10조].',
                     '원고는 2020. 3. 1. 입사하였습니다 (갑 제1호증).'):
            with self.subTest(text=text):
                self.assertEqual(self.check(text)['sentences'], 1)
        result = self.check('원고는 2020. 3. 1. 입사하였습니다(갑 제1호증 근로계약서). '
                            '원고는 영업부에서 근무하였습니다(갑 제2호증 인사기록카드). '
                            '피고는 이 사건 해고를 하였다(갑 제4호증 해고통지서).')
        self.assertEqual(result['sentences'], 3)
        self.assertEqual(self.found(result, '3문장 이상 문단', 'notes'), ['3문장'])
        self.assertEqual(self.found(result, '해라체 종결', 'notes'), ['다(갑 제4호증 해고통지서).'])

    def test_outline_marker_is_not_a_sentence(self):
        result = self.check('다. 소결', '가. 원고는 해고되었다.')
        self.assertEqual(result['sentences'], 1)
        self.assertEqual(len(self.found(result, '해라체 종결', 'notes')), 1)

    def test_list_citation_ends_a_fact_paragraph(self):
        rule = '증거 인용 뒤 평가 문단에 지시 접속어 없음'
        result = self.check('원고는 2020. 3. 1. 입사하였습니다(갑 제2, 3호증).', '원고는 성실히 근무하였습니다.')
        self.assertEqual(len(self.found(result, rule, 'notes')), 1)
        # 발췌 캡션 한 줄은 사실 문단이 아니다
        result = self.check('[갑 제8호증 징계처분사유설명서]', '원고는 성실히 근무하였습니다.')
        self.assertEqual(self.found(result, rule, 'notes'), [])
        # 앞 문단과의 관계를 드러내는 연결어로 시작하면 잡지 않는다(SKILL.md '문단·문장 이음')
        for start in ('그런데도 피고는 해명하지 않았습니다.', '실제로 원고는 성실히 근무하였습니다.', '참고로, 원고는 상을 받았습니다.'):
            result = self.check('원고는 2020. 3. 1. 입사하였습니다(갑 제2호증).', start)
            self.assertEqual(self.found(result, rule, 'notes'), [], start)
        # 증거를 인용한 문단 바로 뒤의 캡션·그림 묶음은 평가 문단이 아니다(2026. 10. 3. 한영대 재심청구서)
        result = self.check('신고인은 "다 빼고 찾아 뵐게요"라고 답하였습니다(증 제7호증).',
                            '[증 제7호증 신고인과 주고받은 메시지]\n![](증7-2.png)')
        self.assertEqual(self.found(result, rule, 'notes'), [])
        # 인용 문단 뒤의 (n) 소제목·'(n) 소결'·'가.' 제목은 평가 문단이 아니다
        result = self.check('원고는 입사하였습니다(갑 제2호증).', '(2) 원고는 성실히 근무하였습니다.')
        self.assertEqual(self.found(result, rule, 'notes'), [])
        result = self.check('원고는 입사하였습니다(갑 제2호증).', '(3) 소결')
        self.assertEqual(self.found(result, rule, 'notes'), [])
        # 인용 문단 → 캡션 → 다음 문단: 캡션도, 캡션 뒤 문단도 세지 않는다
        result = self.check('원고는 입사하였습니다(갑 제2호증).', '[갑 제2호증 근로계약서]', '원고는 성실히 근무하였습니다.')
        self.assertEqual(self.found(result, rule, 'notes'), [])

    def test_write_verb_is_noted(self):
        """자료에 쓰여 있다는 뜻의 '적다'는 '기재하다'로 쓴다(2026. 10. 16. OES 답변서(1) 담당 노무사).
        참고 항목이고, '적용·적극·실적·목적·적어도·적시' 같은 다른 낱말은 잡지 않는다."""
        rule = "'적다' → '기재'"
        for text in ("채용공고에 '최대 10년까지 재계약 가능'이라고 적혀 있던 사안입니다.",
                     "입사 추천서에 관계를 '전 직장 선배'로 적었을 뿐입니다(사 제2호증의4).",
                     "인사발령 문서에 갱신 거절의 사유가 적혀 있지 않습니다.",
                     "'계약기간 만료'라고만 적은 문서로 통보한 사안입니다.",
                     '계약 종료를 문서에 따로 적지 않았습니다.',
                     '다른 직원의 성명은 비실명으로 적고 생년월일은 적지 않았습니다.'):
            with self.subTest(text=text):
                result = self.check(text)
                self.assertTrue(self.found(result, rule, 'notes'))
                self.assertEqual(result['violations'], [])
        for text in ('근로기준법 제27조가 적용되지 않습니다.',
                     '신청인은 적극적으로 협조하지 않았습니다.',
                     '업무 실적과 사업의 목적에 비추어 판단하였습니다.',
                     '적어도 3년 이상 근무하였습니다.',
                     '최종적인 의사결정이 적시에 이루어지지 못하였습니다.',
                     '적법한 절차를 거쳤습니다.',
                     '연장 여부는 계약직 심의 결과에 기재된 바와 같습니다(사 제13호증).'):
            with self.subTest(text=text):
                self.assertEqual(self.found(self.check(text), rule, 'notes'), [])

    def test_appendix_below_rule_is_not_checked(self):
        result = self.check('원고는 입사하였습니다.', '---', '① [확인 필요 사항]', '[반박 보완 필요] 둘째, 근태기록')
        self.assertEqual(result['violations'], [])
        self.assertEqual(result['appendix_paragraphs'], 2)

    def test_model_ratios_are_measured(self):
        """본보기와 견주는 두 값(2026. 10. 21.): 연결어로 연 문단의 비율, 200자 넘는 문장의 비율. 제목 줄은 세지 않는다."""
        long = '원고는 ' + '가' * 210 + '라고 주장합니다.'
        result = self.check('(1) 원고의 주장은 이유 없습니다.', '그러나, 원고의 주장은 이유 없습니다.', '원고는 입사하였습니다.', long)
        m = result['measured']
        self.assertEqual(m['connector_open_ratio'], 33.3)
        self.assertEqual(m['over200_ratio'], 33.3)
        rules = [n['rule'] for n in result['notes']]
        self.assertIn('문단 첫머리 연결어가 본보기보다 적음', rules)
        self.assertIn('200자 넘는 문장이 본보기보다 많음', rules)
        result = self.check('그러나, 원고의 주장은 이유 없습니다.', '이처럼 원고는 입사하였습니다.')
        self.assertEqual(result['measured']['connector_open_ratio'], 100.0)
        self.assertNotIn('문단 첫머리 연결어가 본보기보다 적음', [n['rule'] for n in result['notes']])

    def test_quotes_in_headings_are_violations(self):
        """소제목에는 따옴표를 쓰지 않는다(SKILL.md 금지 사항 16, 2026. 10. 21. 박준선 항소이유서 담당자 지시)."""
        rule = '소제목 따옴표'
        for heading in ("(4) '유실물 관리체계의 형해화'는 근거가 없습니다.",
                        "나) 원고들이 관여한 행위는 상시 계속하는 '직무'에 해당하지 않습니다.",
                        '가. 원고가 말한 "가장 엄중한 처분"은 해임입니다.',
                        '1. ‘파면’ 기준의 부당성'):
            with self.subTest(heading=heading):
                self.assertEqual(len(self.found(self.check(heading), rule)), 1)
        for text in ('(4) 이 사건 비위로 유실물 관리체계가 무너질 위험이 생겼다고 볼 근거가 없습니다.',
                     "징계위원회는 '가장 엄중한 처분이 필요하다'는 이유로 해임을 의결하였습니다(갑 제4호증 10면).",
                     '[갑 제4호증 징계처분사유설명서]'):
            with self.subTest(text=text):
                self.assertEqual(self.found(self.check(text), rule), [])
        # 보완 메모(구분선 아래)의 번호 줄은 보지 않는다
        result = self.check('원고는 입사하였습니다.', '---', "1. '파면' 표기 확인")
        self.assertEqual(self.found(result, rule), [])


class DocxHeadingTests(unittest.TestCase):
    """docx 의 번호 매기기 제목은 실측에서 빠지되 앞 문단의 증거 인용 맥락을 끊는다(md 의 제목 줄과 같게).
    2026. 10. 3. 한영대 재심청구서 병합본에서 '(3) …' 소제목 너머의 다음 목 첫 문단이 참고로 잡혔다."""

    RULE = '증거 인용 뒤 평가 문단에 지시 접속어 없음'

    def check(self, with_heading):
        from docx.enum.style import WD_STYLE_TYPE
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / '제목.docx'
            document = docx.Document()
            heading = document.styles.add_style('(1) 번호매기기', WD_STYLE_TYPE.PARAGRAPH)
            document.add_paragraph('신고인도 먼저 연락하여 서로 연락을 주고받았습니다(증 제7호증).')
            if with_heading:
                document.add_paragraph('산책으로 사생활을 침해한 사실도 없습니다.', style=heading)
            document.add_paragraph('청구인은 점심 산책 때 신고인을 따라다닌 사실이 없습니다.')
            document.save(path)
            return style_check.check(path)

    def test_numbered_heading_breaks_citation_context(self):
        result = self.check(with_heading=True)
        self.assertEqual([n for n in result['notes'] if n['rule'] == self.RULE], [])
        self.assertEqual(result['appendix_paragraphs'], 0)   # 제목은 보완 메모로 세지 않는다
        self.assertEqual(result['paragraphs'], 2)            # 제목은 본문 실측에서 빠진다

    def test_without_heading_the_note_remains(self):
        result = self.check(with_heading=False)
        self.assertEqual(len([n for n in result['notes'] if n['rule'] == self.RULE]), 1)

    def test_quote_in_heading_style_is_violation(self):
        """번호 매기기 제목 스타일의 따옴표만 잡고, 본문과 '기타번호매기기'(입증방법 등)의 따옴표는 그대로 둔다(2026. 10. 21.)."""
        from docx.enum.style import WD_STYLE_TYPE
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / '따옴표.docx'
            document = docx.Document()
            heading = document.styles.add_style('(1) 번호매기기', WD_STYLE_TYPE.PARAGRAPH)
            other = document.styles.add_style('기타번호매기기', WD_STYLE_TYPE.PARAGRAPH)
            document.add_paragraph("원고 징계위원회는 '파면'을 상정한 사실이 없습니다.", style=heading)
            document.add_paragraph("을나 제14호증 '횡령' 관련 언론 보도", style=other)
            document.add_paragraph("징계위원회는 '가장 엄중한 처분'이라고 하였습니다(갑 제4호증).")
            document.save(path)
            result = style_check.check(path)
        self.assertEqual([v['detail'] for v in result['violations'] if v['rule'] == '소제목 따옴표'], ["'"])


if __name__ == '__main__':
    unittest.main()
