"""기존 원본을 변경하지 않는 일회성 실행·대조 도구. 계산은 기존 engine만 호출."""
from pathlib import Path
import sys, json, hashlib, importlib.util, inspect, traceback
from datetime import date
from decimal import Decimal
import openpyxl

ROOT = Path(__file__).resolve().parents[3]
OUT = Path(__file__).resolve().parent

def sha(p):
    h = hashlib.sha256()
    with p.open('rb') as f:
        for b in iter(lambda:f.read(1024*1024), b''): h.update(b)
    return h.hexdigest()

def dump(name, value):
    (OUT/name).write_text(json.dumps(value, ensure_ascii=False, indent=2, default=str), encoding='utf-8')

def calc():
    base = ROOT/'6_손해배상계산'
    sys.path.insert(0, str(base))
    from engine.case import Case, ImpairmentInput
    from engine.calculate import calculate
    from engine.excel import write_workbook
    from tests.golden import case_fault_deduction as F
    from tests import test_golden
    tests = []
    for n, f in inspect.getmembers(test_golden, inspect.isfunction):
        if n.startswith('test_'):
            try: f(); tests.append({'test':n,'status':'passed'})
            except Exception: tests.append({'test':n,'status':'failed','error':traceback.format_exc()})
    # sample.yaml의 명시 입력을 그대로 재현. YAML 로더는 이 실행에 포함되지 않음.
    c = Case(case_no='2024가단1',name='홍길동',birth=F.BIRTH,accident=F.ACCIDENT,
        cure_end=F.CURE_END,argument_end=date(2026,3,8),life_expectancy=F.LIFE_EXPECTANCY,
        life_end=F.LIFE_END,impairments=[ImpairmentInput(x[0],Decimal(x[1]),Decimal(x[2] or 0)) for x in F.IMPAIRMENTS],
        wages={f'{y}-{i}':v for (y,i),v in F.WAGES.items()},fault_rate=F.FAULT_RATE,
        paid_cure=F.PAID_CURE,advance=F.ADVANCE,solatium=F.SOLATIUM_APPLIED)
    result = calculate(c, lambda p:F.WAGES.get((p.year,p.index),F.WAGES[max(F.WAGES)]),lambda y,i:(y,i) in F.WAGES)
    old = base/'02_결과/테스트_계산표.xlsx'
    before = sha(old)
    target = OUT/'6_손해배상계산_재산출.xlsx'
    if target.exists(): raise FileExistsError(target)
    write_workbook(result,target)
    a,b = openpyxl.load_workbook(old),openpyxl.load_workbook(target)
    diffs=[]; nonempty=0
    for s in a.sheetnames:
        if s not in b.sheetnames: diffs.append({'sheet':s,'missing':True});continue
        x,y=a[s],b[s]
        for row in range(1,max(x.max_row,y.max_row)+1):
            for col in range(1,max(x.max_column,y.max_column)+1):
                u,v=x.cell(row,col),y.cell(row,col)
                if u.value is not None or v.value is not None: nonempty+=1
                if u.value!=v.value: diffs.append({'sheet':s,'cell':u.coordinate,'old':u.value,'new':v.value})
    report={'scope':'가상사건 engine 재계산 및 XLSX 생성. CLI YAML 로더·폴더 감시·실사건 값 추출·대법원 프로그램 신규 실행 제외',
        'source':str(old.relative_to(ROOT)),'original_sha256':before,'original_unchanged':sha(old)==before,
        'new_sha256':sha(target),'tests':tests,'sheets':b.sheetnames,'nonempty_cells_compared':nonempty,
        'value_differences':diffs,'income_total':result.income_total,'settlement':result.settlement,
        'visual_qa':'미실시'}
    dump('6_계산대조.json',report)
    print(json.dumps(report,ensure_ascii=False,default=str,indent=2))

if __name__=='__main__': calc()
